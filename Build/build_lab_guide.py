#!/usr/bin/env python3
# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""Build Docs/Master_Class_Lab_Guide.docx from the Master Class lab guide template.

The template (Templates/Master_Class_LabGuide_Template.docx) supplies the
cover, the information page, the Contents field, the revision table, the lab
environment and login sections, Exercise 1, and the paragraph and table
formatting. This script keeps everything up to the end of the template's
Exercise 1, so attendees still have the login details. It drops the
template's example exercises after that and writes the Workbench lab as
Exercise 2 onward, with the screenshots under Docs/Lab_Guide_Assets.
Screenshots that cannot be captured from the repo, such as console views,
are marked placeholders for the instructor to replace. The Workbench
screenshots come from the instructor's own lab run.

    uv run python Utilities/capture_lab_guide_task_screenshots.py
    uv run python Build/build_lab_guide.py

The Contents field is refreshed by Word afterwards (Build/update_lab_guide_contents.ps1),
which this script runs when PowerShell is available.
"""

from __future__ import annotations

import copy
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = PROJECT_ROOT / "Templates" / "Master_Class_LabGuide_Template.docx"
ASSETS = PROJECT_ROOT / "Docs" / "Lab_Guide_Assets"
OUTPUT = PROJECT_ROOT / "Docs" / "Master_Class_Lab_Guide.docx"
CONTENTS_SCRIPT = PROJECT_ROOT / "Build" / "update_lab_guide_contents.ps1"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W = f'xmlns:w="{W}"'
SITE = "Performance and Config Analysis"
GUIDE_VERSION = "1.0"
GUIDE_DATE = "October 2, 2026"
# Paths on the lab's BigFix Server virtual machine. The Output Folder default
# is the logged-in user's Documents folder, which is Administrator on the lab.
WORKBENCH_DIR = (
    r"C:\Users\Administrator\Desktop\BigFixSrc\Master Class Diagnostics Workbench"
)
SAMPLE_DIR = WORKBENCH_DIR + r"\SampleData\BigFix Master Class 2026\Results"
OUTPUT_DIR = r"C:\Users\Administrator\Documents\Master Class Diagnostics"
NFR_TITLE = "BigFix Non-Functional Requirements: A Checklist Approach v11.x.4"
NFR_URL = (
    "https://help.hcl-software.com/bigfix/landing/Technical+Documents/"
    "Landing_page_shared/Technical_Documents.html"
)
LAB_RAM_GB = "12"
TASK_DEPLOYMENT = "BigFix Master Class -- Collect Deployment Information (2.0)"
TASK_CLIENT_PERF = (
    "BigFix Master Class -- Collect Client Performance Counter Analysis Data (2.0)"
)
ANALYSIS_CLIENT_PERF = "BigFix Master Class -- Client Performance Counters - Basic (2.0)"
TASK_PERFMON = "BigFix Master Class -- Setup and Run Perfmon Counters (2.0)"
TASK_FILLDB_ENABLE = "BigFix Master Class -- Enable FillDB Performance Logging (2.0)"
TASK_FILLDB_DISABLE = (
    "WARNING: BigFix Master Class -- FillDB Performance Logging is Enabled (2.0)"
)
IMAGE_WIDTH = Inches(6.6)
TASK_IMAGE_WIDTH = Inches(6.0)


# ---------------------------------------------------------------------------
# Template access
# ---------------------------------------------------------------------------


def body_blocks(document):
    return [
        child for child in document.element.body if child.tag in (qn("w:p"), qn("w:tbl"))
    ]


def block_text(block) -> str:
    return "".join(t.text or "" for t in block.iter(qn("w:t"))).strip()


def find_block(document, prefix: str, tag: str = "w:p"):
    for block in body_blocks(document):
        if block.tag == qn(tag) and block_text(block).startswith(prefix):
            return block
    raise LookupError(f"template block not found: {prefix!r}")


def strip_bookmarks(element) -> None:
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for node in list(element.iter(qn(tag))):
            node.getparent().remove(node)


def set_paragraph_text(paragraph, text: str) -> None:
    """Keep the first run's formatting, drop the rest, set the text."""
    runs = list(paragraph.iter(qn("w:r")))
    keep = None
    for run in runs:
        if run.find(qn("w:t")) is not None:
            keep = run
            break
    if keep is None:
        keep = runs[0]
    for run in runs:
        if run is not keep:
            run.getparent().remove(run)
    for child in list(keep):
        if child.tag != qn("w:rPr"):
            keep.remove(child)
    t = parse_xml(f'<w:t {_W} xml:space="preserve"></w:t>')
    t.text = text
    keep.append(t)


class Exemplars:
    """Paragraphs and tables cloned from the template for their formatting."""

    def __init__(self, document) -> None:
        self.h1 = find_block(document, "Student exercises")
        self.h2 = find_block(document, "Overview")
        self.body = find_block(document, "In this exercise, you will install")
        self.note = find_block(document, "NOTE:")
        self.important = find_block(document, "Important:")
        self.done = find_block(document, "You have now successfully completed Exercise 1")
        self.table = find_block(document, "Host Name", "w:tbl")
        for block in body_blocks(document):
            if block.tag != qn("w:p") or block_text(block):
                continue
            # The cover picture paragraph carries a section break in its
            # properties. Cloning it would start a new section per screenshot.
            if next(block.iter(qn("w:sectPr")), None) is not None:
                continue
            # The drawing sits inside a run, so search the subtree.
            if next(block.iter(qn("w:drawing")), None) is not None:
                self.picture = block
                break
        else:
            raise LookupError("template has no screenshot paragraph")


# ---------------------------------------------------------------------------
# Numbering
# ---------------------------------------------------------------------------


class Numbering:
    """Fresh list instances so every step list restarts at 1."""

    def __init__(self, document) -> None:
        self.part = document.part.numbering_part.element
        self.next_num = max(
            (int(n.get(qn("w:numId"))) for n in self.part.iter(qn("w:num"))),
            default=0,
        )
        self.steps_abstract = self._abstract(
            900, '<w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>'
        )
        self.bullet_abstract = self._abstract(
            901,
            '<w:numFmt w:val="bullet"/><w:lvlText w:val="&#8226;"/>',
            '<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/></w:rPr>',
        )

    def _abstract(self, ident: int, fmt: str, rpr: str = "<w:rPr/>") -> int:
        xml = (
            f'<w:abstractNum {_W} w:abstractNumId="{ident}">'
            '<w:multiLevelType w:val="singleLevel"/>'
            f'<w:lvl w:ilvl="0"><w:start w:val="1"/>{fmt}<w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="1080" w:hanging="360"/></w:pPr>'
            f"{rpr}</w:lvl></w:abstractNum>"
        )
        first_num = self.part.find(qn("w:num"))
        node = parse_xml(xml)
        if first_num is not None:
            first_num.addprevious(node)
        else:
            self.part.append(node)
        return ident

    def new_list(self, abstract: int) -> int:
        self.next_num += 1
        self.part.append(
            parse_xml(
                f'<w:num {_W} w:numId="{self.next_num}">'
                f'<w:abstractNumId w:val="{abstract}"/>'
                '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>'
                "</w:num>"
            )
        )
        return self.next_num


# ---------------------------------------------------------------------------
# Writer
# ---------------------------------------------------------------------------


class Guide:
    def __init__(self, document) -> None:
        self.document = document
        self.ex = Exemplars(document)
        self.numbering = Numbering(document)
        self.body = document.element.body
        self.sect_pr = self.body.find(qn("w:sectPr"))

    def _append(self, element) -> None:
        self.sect_pr.addprevious(element)
        return element

    def heading1(self, text: str, *, new_page: bool = False) -> None:
        p = copy.deepcopy(self.ex.h1)
        strip_bookmarks(p)
        set_paragraph_text(p, text)
        if new_page:
            p.find(qn("w:pPr")).append(parse_xml(f"<w:pageBreakBefore {_W}/>"))
        self._append(p)

    def heading2(self, text: str) -> None:
        p = copy.deepcopy(self.ex.h2)
        strip_bookmarks(p)
        set_paragraph_text(p, text)
        self._append(p)

    def para(
        self, text: str, *, bold_prefix: str | None = None, indent: int | None = None
    ):
        p = copy.deepcopy(self.ex.body)
        set_paragraph_text(p, text)
        if indent is not None:
            ind = p.find(qn("w:pPr")).find(qn("w:ind"))
            ind.set(qn("w:left"), str(indent))
            ind.set(qn("w:firstLine"), "0")
        if bold_prefix:
            run = next(r for r in p.iter(qn("w:r")) if r.find(qn("w:t")) is not None)
            label = copy.deepcopy(run)
            label.find(qn("w:t")).text = bold_prefix
            rpr = label.find(qn("w:rPr"))
            for tag in ("w:b", "w:bCs"):
                node = rpr.find(qn(tag))
                if node is None:
                    node = parse_xml(f"<{tag} {_W}/>")
                    rpr.insert(1, node)
                node.attrib.pop(qn("w:val"), None)
            run.addprevious(label)
        self._append(p)
        return p

    def label(self, text: str) -> None:
        p = self.para(text)
        run = next(r for r in p.iter(qn("w:r")) if r.find(qn("w:t")) is not None)
        rpr = run.find(qn("w:rPr"))
        for tag in ("w:b", "w:bCs"):
            node = rpr.find(qn(tag))
            if node is None:
                node = parse_xml(f"<{tag} {_W}/>")
                rpr.insert(1, node)
            node.attrib.pop(qn("w:val"), None)

    def note(self, text: str) -> None:
        self.para(text, bold_prefix="NOTE: ", indent=720)

    def important(self, text: str) -> None:
        p = copy.deepcopy(self.ex.important)
        runs = [r for r in p.iter(qn("w:r")) if r.find(qn("w:t")) is not None]
        runs[0].find(qn("w:t")).text = "Important"
        runs[1].find(qn("w:t")).text = ": " + text
        for extra in runs[2:]:
            extra.getparent().remove(extra)
        self._append(p)

    def done(self, number: int) -> None:
        p = copy.deepcopy(self.ex.done)
        set_paragraph_text(p, f"You have now successfully completed Exercise {number}.")
        self._append(p)

    def _list(self, items: list[str], abstract: int, num_id: int | None = None) -> None:
        if num_id is None:
            num_id = self.numbering.new_list(abstract)
        for item in items:
            p = copy.deepcopy(self.ex.body)
            set_paragraph_text(p, item)
            ppr = p.find(qn("w:pPr"))
            ind = ppr.find(qn("w:ind"))
            ind.set(qn("w:left"), "1080")
            ind.set(qn("w:hanging"), "360")
            ind.attrib.pop(qn("w:firstLine"), None)
            num_pr = parse_xml(
                f'<w:numPr {_W}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
            )
            ppr.insert(0, num_pr)
            self._append(p)

    def steps(self, items: list[str], *, restart: bool = False) -> None:
        """Numbered steps. A screenshot between calls keeps the numbering
        going, so only the first call of a step list passes restart=True."""
        if restart or getattr(self, "_steps_num", None) is None:
            self._steps_num = self.numbering.new_list(self.numbering.steps_abstract)
        self._list(items, self.numbering.steps_abstract, self._steps_num)

    def bullets(self, items: list[str]) -> None:
        self._list(items, self.numbering.bullet_abstract)

    def picture(self, name: str, width=IMAGE_WIDTH) -> None:
        path = ASSETS / name
        if not path.is_file():
            raise FileNotFoundError(path)
        p = copy.deepcopy(self.ex.picture)
        for run in list(p.iter(qn("w:r"))):
            run.getparent().remove(run)
        self._append(p)
        from docx.text.paragraph import Paragraph

        Paragraph(p, self.document._body).add_run().add_picture(str(path), width=width)

    def placeholder(self, what: str) -> None:
        """A highlighted marker where the instructor adds a screenshot."""
        p = self.para(f"[SCREENSHOT: {what}]", indent=720)
        run = next(r for r in p.iter(qn("w:r")) if r.find(qn("w:t")) is not None)
        rpr = run.find(qn("w:rPr"))
        rpr.append(parse_xml(f'<w:highlight {_W} w:val="yellow"/>'))
        for tag in ("w:b", "w:bCs"):
            node = rpr.find(qn(tag))
            if node is None:
                rpr.insert(1, parse_xml(f"<{tag} {_W}/>"))

    def shot(self, name: str, what: str, width=IMAGE_WIDTH) -> None:
        """An instructor screenshot when the file exists, else its placeholder."""
        if (ASSETS / name).is_file():
            self.picture(name, width)
        else:
            self.placeholder(f"{what}. Save it as Docs/Lab_Guide_Assets/{name}")

    def link(self, lead: str, text: str, url: str) -> None:
        """A body paragraph that ends in a real hyperlink."""
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        p = self.para(lead)
        run = next(r for r in p.iter(qn("w:r")) if r.find(qn("w:t")) is not None)
        rel_id = self.document.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link_run = copy.deepcopy(run)
        link_run.find(qn("w:t")).text = text
        rpr = link_run.find(qn("w:rPr"))
        rpr.append(parse_xml(f'<w:color {_W} w:val="0563C1"/>'))
        rpr.append(parse_xml(f'<w:u {_W} w:val="single"/>'))
        hyperlink = parse_xml(
            f'<w:hyperlink {_W} xmlns:r="http://schemas.openxmlformats.org/'
            f'officeDocument/2006/relationships" r:id="{rel_id}" w:history="1"/>'
        )
        hyperlink.append(link_run)
        run.addnext(hyperlink)

    def code(self, lines: list[str]) -> None:
        for line in lines:
            p = self.para(line, indent=720)
            run = next(r for r in p.iter(qn("w:r")) if r.find(qn("w:t")) is not None)
            rpr = run.find(qn("w:rPr"))
            fonts = rpr.find(qn("w:rFonts"))
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(attr), "Consolas")
            ppr = p.find(qn("w:pPr"))
            spacing = ppr.find(qn("w:spacing"))
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")

    def table(self, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
        tbl = copy.deepcopy(self.ex.table)
        template_rows = tbl.findall(qn("w:tr"))
        head_row, data_row = template_rows[0], template_rows[1]
        for row in template_rows:
            tbl.remove(row)
        grid = tbl.find(qn("w:tblGrid"))
        for col in list(grid):
            grid.remove(col)
        for width in widths:
            grid.append(parse_xml(f'<w:gridCol {_W} w:w="{width}"/>'))
        tbl.find(qn("w:tblPr")).find(qn("w:tblW")).set(qn("w:w"), str(sum(widths)))

        def build_row(exemplar, values, header: bool):
            cells = exemplar.findall(qn("w:tc"))
            # Cell 1 of the template rows is the narrow numbering column.
            # Cell 2 carries the formatting the guide's tables use.
            proto = cells[1]
            row = copy.deepcopy(exemplar)
            for cell in row.findall(qn("w:tc")):
                row.remove(cell)
            for value, width in zip(values, widths):
                cell = copy.deepcopy(proto)
                tc_pr = cell.find(qn("w:tcPr"))
                tcw = tc_pr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = parse_xml(f'<w:tcW {_W} w:w="{width}" w:type="dxa"/>')
                    tc_pr.append(tcw)
                tcw.set(qn("w:w"), str(width))
                paragraph = cell.find(qn("w:p"))
                set_paragraph_text(paragraph, value)
                jc = paragraph.find(qn("w:pPr")).find(qn("w:jc"))
                if not header and jc is not None:
                    jc.getparent().remove(jc)
                row.append(cell)
            tr_pr = row.find(qn("w:trPr"))
            if tr_pr is not None:
                height = tr_pr.find(qn("w:trHeight"))
                if height is not None:
                    tr_pr.remove(height)
                if header:
                    hdr = tr_pr.find(qn("w:tblHeader"))
                    if hdr is not None:
                        hdr.attrib.pop(qn("w:val"), None)
            return row

        tbl.append(build_row(head_row, headers, True))
        for values in rows:
            tbl.append(build_row(data_row, values, False))
        self._append(tbl)
        self.para("")


# ---------------------------------------------------------------------------
# Front matter edits
# ---------------------------------------------------------------------------


def update_front_matter(document) -> None:
    title = find_block(document, "Student Workbook")
    set_paragraph_text(title, "Student Lab Guide")
    subtitle = copy.deepcopy(title)
    set_paragraph_text(subtitle, "Automated Performance and Configuration Analysis")
    title.addprevious(subtitle)
    date = find_block(document, "June 2024")
    set_paragraph_text(date, GUIDE_DATE)

    revisions = find_block(document, "Release Date", "w:tbl")
    rows = revisions.findall(qn("w:tr"))
    entries = [(GUIDE_DATE, GUIDE_VERSION, "Master Class team", "First version.")]
    for row, values in zip(rows[1:], entries):
        for cell, value in zip(row.findall(qn("w:tc")), values):
            set_paragraph_text(cell.find(qn("w:p")), value)
    for row in rows[1 + len(entries) :]:
        for cell in row.findall(qn("w:tc")):
            set_paragraph_text(cell.find(qn("w:p")), "")


PROTECTED_END = "You have now successfully completed Exercise 1."
EXAMPLE_HEADING = "EXAMPLE: BigFix Foundation"


def drop_template_exercises(document) -> None:
    """Drop the template's example exercises after Exercise 1.

    Everything before the EXAMPLE heading stays: the cover, the Contents, the
    revision table, the Overview, the Lab Environment table with the logins,
    Accessing Lab Environment, and Exercise 1. Attendees need those to log in.
    """
    blocks = body_blocks(document)
    protected = next(
        (i for i, b in enumerate(blocks) if block_text(b) == PROTECTED_END), None
    )
    if protected is None:
        raise LookupError(f"template paragraph not found: {PROTECTED_END!r}")
    start = None
    for block in blocks[protected + 1 :]:
        if block.tag != qn("w:p"):
            continue
        ppr = block.find(qn("w:pPr"))
        style = ppr.find(qn("w:pStyle")) if ppr is not None else None
        if (
            style is not None
            and style.get(qn("w:val")) == "Heading1"
            and block_text(block).startswith(EXAMPLE_HEADING)
        ):
            start = block
            break
    if start is None:
        raise LookupError(f"template Heading 1 {EXAMPLE_HEADING!r} not found")
    body = document.element.body
    removing = False
    for child in list(body):
        if child is start:
            removing = True
        if child.tag == qn("w:sectPr"):
            continue
        if removing:
            body.remove(child)


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------

TIME_BUDGET = [
    ["Before You Start", "2 minutes"],
    ["Exercise 2: Collect the Data", "5 minutes"],
    ["Exercise 3: Run the Workbench", "5 minutes"],
    ["Exercise 4: Read the Client Performance Outputs", "4 minutes"],
    ["Exercise 5: Read the Perfmon Outputs", "4 minutes"],
]

CLIENT_AREAS = [
    ["1", "Client Loop Time", "Average Loop Health", "Health Summary"],
    [
        "2",
        "Expensive Properties (Top 10)",
        "Top 10 Most Expensive Evaluations",
        "Evaluation Offenders",
    ],
    ["3", "UDP Notification Reach", "UDP Notification Coverage", "UDP Coverage"],
    ["4", "Command Polling", "Command Polling Review", "Polling Summary"],
    [
        "5",
        "CPU Usage Settings",
        "Client CPU Idle Settings",
        "Client CPU Idle Settings",
    ],
    ["6", "FastHashVerify", "FastHashVerify Review", "FastHashVerify Summary"],
]

PERFMON_COUNTERS = [
    [
        "PRF-INF-01",
        "Storage Queue Depth",
        "Physical Disk: Current Disk Queue Length",
        "≤ 1 per physical device",
    ],
    [
        "PRF-INF-02",
        "Storage Latency",
        "Physical Disk: Avg. Disk sec/Transfer",
        "≤ 1 ms per physical device",
    ],
    [
        "PRF-INF-03",
        "Storage Operations",
        "Physical Disk: Disk Transfers/sec",
        "No saturation under load. DBMS storage supports 5,000 transfers per second.",
    ],
    [
        "PRF-INF-04",
        "CPU Utilization",
        "Processor: % Processor Time",
        "Virtual: 50% to 80%. Physical: ≤ 80%.",
    ],
    [
        "PRF-INF-05",
        "CPU Latency",
        "System: Processor Queue Length",
        "≤ 4",
    ],
    [
        "PRF-INF-06",
        "CPU Privileged Time",
        "Processor: % Privileged Time",
        "≤ 10% of aggregate CPU",
    ],
    [
        "PRF-INF-07",
        "Network Health",
        "Network Interface: packets discarded and packet errors",
        "All counters approach zero per interface",
    ],
    [
        "PRF-INF-09",
        "Memory Health",
        "Paging File: % Usage Peak",
        "≤ 5% of aggregate physical memory",
    ],
]


def take_action(g: Guide, task: str, inputs: list[str], result: str) -> None:
    """Steps to run one task from the site against the BigFix Server."""
    g.steps([f"Select {task}. The task description opens in the work area."])
    g.steps(inputs)
    g.steps(
        [
            "Click Take Action. On the Target tab select the BigFix Server. Click OK.",
            result,
        ]
    )


def write_content(g: Guide) -> None:
    g.heading1("Diagnostics Workbench Lab")

    g.heading2("Overview")
    g.para(
        "In this lab you run the Master Class Diagnostics Workbench. You cover the two "
        "modules from the session. Client Performance Analysis reviews how the BigFix "
        "clients evaluate content. Perfmon Infrastructure Analysis reviews the Windows "
        "performance of a BigFix server."
    )
    g.para("The lab takes 20 minutes:")
    g.table(["Part", "Time"], TIME_BUDGET, [7000, 3000])

    g.heading2("Before You Start")
    g.para(
        "The collection tasks need hours of client reporting and a 24-hour Perfmon "
        "capture. In this lab you run them only to see how they work. The Workbench "
        "then builds from the sample data that ships with it."
    )
    g.para("Every module follows the same path:")
    g.para(
        "BigFix deployment › Collection task › Captured data › Workbench module › "
        "Mini Report",
        indent=720,
    )
    g.para(
        f"The Workbench is on the BigFix Server virtual machine in {WORKBENCH_DIR}. "
        f"The sample data is in {SAMPLE_DIR}."
    )
    g.note(
        "The Workbench is provided as-is to BigFix Master Class attendees only. It has "
        "no warranty or support. Do not redistribute it without approval."
    )

    # Exercise 2 --------------------------------------------------------
    g.heading2("Exercise 2: Collect the Data")
    g.para(
        "In this exercise, you run the collection content for both modules from the "
        "BigFix Console. You do not wait for the Deployment Information task to finish."
    )
    g.label("Client Performance")
    g.steps(
        restart=True,
        items=[
            "Switch to the BigFix Server virtual machine. Log in to the BigFix Console "
            "with the console credentials in the Lab Environment table.",
            f"Select the All Content domain. Expand Sites and Custom Sites. Select the "
            f"{SITE} site.",
        ],
    )
    g.shot(
        "lg-console-01-site-selected.png",
        f"The {SITE} site selected in the All Content domain",
    )
    g.steps(
        [
            f"Select Analyses. Right-click {ANALYSIS_CLIENT_PERF} and choose Activate. "
            "The status changes to Activated Globally.",
        ]
    )
    g.picture("lg-task-01-agent-performance-analysis.png", TASK_IMAGE_WIDTH)
    g.note(
        "In a real health check the analysis reports for 24 to 48 hours before you "
        "collect. In the lab you collect right away to see how the task works."
    )
    g.steps(["Select Fixlets and Tasks in the site."])
    take_action(
        g,
        TASK_DEPLOYMENT,
        [
            "Enter the console user name in API Username. Enter the console password "
            "in API Password and Confirm Password. Leave Root Server Hostname as it is.",
        ],
        "The action starts. Do not wait for it. Continue with the next step.",
    )
    g.picture("lg-task-05-collect-deployment-info.png", TASK_IMAGE_WIDTH)
    take_action(
        g,
        TASK_CLIENT_PERF,
        [
            "Enter the console user name in API Username. Enter the console password "
            "in API Password and Confirm Password. Leave Root Server Hostname as it is.",
            "Leave Also generate the human readable loop CSV file cleared.",
        ],
        "Wait for the action status to change to Completed.",
    )
    g.picture("lg-task-02-collect-agent-performance.png", TASK_IMAGE_WIDTH)
    g.shot(
        "lg-console-02-client-perf-completed.png",
        "The Collect Client Performance Counter Analysis Data action at Completed",
    )
    g.label("Perfmon")
    take_action(
        g,
        TASK_PERFMON,
        [],
        "Wait for the action status to change to Completed. The BFPerf counter set "
        "is now running.",
    )
    g.picture("lg-task-03-setup-perfmon.png", TASK_IMAGE_WIDTH)
    g.note(
        "The capture takes a sample every 60 seconds and stops on its own after "
        "24 hours. You do not need to stop it. The Workbench uses the sample capture "
        "in Exercise 3."
    )
    g.done(2)

    # Exercise 3 --------------------------------------------------------
    g.heading2("Exercise 3: Run the Workbench")
    g.para(
        "In this exercise, you run the two lab modules against the sample data. You "
        "build once."
    )
    g.label("Step 1: Folders and Run Details")
    g.steps(
        restart=True,
        items=[
            f"On the BigFix Server open {WORKBENCH_DIR}\\Launchers. Double-click Run "
            "Diagnostics Workbench.cmd. The Workbench opens on step 1.",
            f"Click Browse next to Data Folder. Select {SAMPLE_DIR}. The Workbench "
            "reports Data found for 3 of 3 modules.",
            f"Leave Output Folder at {OUTPUT_DIR}. A timestamped run folder is created "
            "there.",
            "Enter a Run Name such as MC Lab Exercise. Any name works. The run folder "
            "and every output file use it.",
        ],
    )
    g.picture("lg-01-step-1-folders.png")
    g.steps(["Click Next. The Choose the Modules to Run page opens."])
    g.label("Step 2: Choose the Modules")
    g.steps(
        [
            "The page opens with every module selected. Click Clear. Select Client "
            "Performance Analysis and Perfmon Infrastructure Analysis. Both rows show "
            "Needs Review until you fill in their details.",
        ]
    )
    g.picture("lg-03-step-2-modules.png")
    g.steps(
        [
            "Click Details on the Client Performance Analysis row. Enter a Report Start "
            "Date and a Report End Date. Any dates work. The report shows this range. "
            "The row changes to Found.",
        ]
    )
    g.picture("lg-04-step-2-loop-details.png")
    g.steps(
        [
            "Click Details on the Perfmon Infrastructure Analysis row. In Capture 1 "
            f"enter {LAB_RAM_GB} in RAM. Select Virtual machine. CPUs and Disks are "
            "read from the capture. The row changes to Found.",
        ]
    )
    g.picture("lg-04-step-2-perfmon-details.png")
    g.note(
        "RAM is not in the Perfmon capture. Without it the Memory Health counter "
        "(PRF-INF-09) is not scored."
    )
    g.steps(["Click Next. The Build the Outputs page opens."])
    g.label("Step 3: Build the Outputs")
    g.steps(["Check the Run Details. Click Build."])
    g.picture("lg-05-step-3-build.png")
    g.steps(
        [
            "Wait for the Build Status of each module to change from Not started to "
            "Running to Done.",
        ]
    )
    g.picture("lg-07-step-3-done.png")
    g.steps(
        [
            "Click Open Output. The run folder opens. It holds a folder for each module "
            "and the run log.",
        ]
    )
    g.picture("lg-08-run-folder.png")
    g.done(3)

    # Exercise 4 --------------------------------------------------------
    g.heading2("Exercise 4: Read the Client Performance Outputs")
    g.para(
        "In this exercise, you read the Client Performance outputs. Each module writes "
        "a Mini Report, an Infographic, and an Excel workbook."
    )
    g.steps(
        restart=True,
        items=["Open the Client_Loop folder."],
    )
    g.picture("lg-09-client-loop-folder.png")
    g.steps(["Open the Mini Report (.docx)."])
    g.shot(
        "lg-11-client-mini-report.png",
        "Page 1 of the Client Performance Mini Report",
        TASK_IMAGE_WIDTH,
    )
    g.steps(
        [
            "Find the six review areas from the session in the Mini Report. The numbers "
            "are for reference only. They are not a priority order.",
        ]
    )
    g.table(
        ["#", "Area", "Mini Report section", "Excel sheet"],
        CLIENT_AREAS,
        [600, 3200, 3400, 2800],
    )
    g.steps(
        [
            "Open the Excel workbook. Each sheet holds the data behind a finding. Filter "
            "or pivot it to answer your own questions.",
        ]
    )
    g.label("Discussion")
    g.bullets(
        [
            "Where is the client loop spending its time?",
            "Which properties show up in the Top 10?",
            "Are notifications reaching the clients?",
        ]
    )
    g.done(4)

    # Exercise 5 --------------------------------------------------------
    g.heading2("Exercise 5: Read the Perfmon Outputs")
    g.para(
        "In this exercise, you read the Perfmon outputs. Each counter is scored "
        "against its target in the NFR guide."
    )
    g.steps(
        restart=True,
        items=["Open the Perfmon folder."],
    )
    g.shot("lg-10-perfmon-folder.png", "The Perfmon folder in the run folder")
    g.steps(["Open the Mini Report (.docx)."])
    g.shot(
        "lg-12-perfmon-mini-report.png",
        "Page 1 of the Perfmon Mini Report",
        TASK_IMAGE_WIDTH,
    )
    g.steps(["Find each counter in the Mini Report. Compare it with its target."])
    g.table(
        ["ID", "Mini Report section", "Counter", "NFR target"],
        PERFMON_COUNTERS,
        [1400, 2200, 3300, 3100],
    )
    g.link(f"The targets come from {NFR_TITLE}. Find it here: ", NFR_URL, NFR_URL)
    g.steps(
        [
            "Open the Excel workbook. Raw Data holds every sample. Findings holds the "
            "result for each counter.",
        ]
    )
    g.label("Discussion")
    g.bullets(
        [
            "Which server resources stand out?",
            "How does the capture window affect what you see?",
            "What would you change first?",
            "How would you measure it?",
        ]
    )
    g.done(5)

    g.heading2("Take Home")
    g.para(
        "Rerun the Workbench on your own schedule. Run it now as a Day 0 baseline. "
        "Rerun it after each change and compare the results."
    )

    # Appendix ------------------------------------------------------------
    g.heading1("Appendix: FillDB Performance (Bonus Round)", new_page=True)
    g.para(
        "FillDB is not part of the live lab. Run it if time permits, or after the "
        "session. The sample data includes a FillDB analyzer CSV, so you can run the "
        "module without collecting logs. "
        "FillDB Performance Analysis reviews how quickly the root server processes "
        "client reports."
    )
    g.label("Collect the logs")
    g.steps(
        restart=True,
        items=[
            f"In the {SITE} site select {TASK_FILLDB_ENABLE}.",
            "Enter the console user name in API Username. Enter the console password "
            "in API Password and Confirm Password. Keep the default duration.",
            "Click Take Action. On the Target tab select the BigFix Server. Click OK.",
        ],
    )
    g.picture("lg-task-04-enable-filldb-logging.png", TASK_IMAGE_WIDTH)
    g.steps(
        [
            "Logging stops on its own after the duration. If you selected Do Not "
            f"Automatically Disable Logging, run {TASK_FILLDB_DISABLE} to stop it.",
            "Copy FillDBPerf.log and its rollover files from the root server. Put them "
            "in a FillDB_Performance folder under your data folder.",
        ]
    )
    g.label("Run the module")
    g.steps(
        [
            "Run the Workbench. On step 2 click Clear and select FillDB Performance "
            "Analysis. Click Details to check the log. The collection dates are "
            "optional.",
        ]
    )
    g.picture("lg-04-step-2-filldb-details.png")
    g.steps(
        [
            "Click Next and click Build.",
            "Open the FillDB_Performance folder in the run folder. Open the Mini Report.",
        ]
    )
    g.shot(
        "lg-13-filldb-mini-report.png",
        "Page 1 of the FillDB Mini Report",
        TASK_IMAGE_WIDTH,
    )
    g.para(
        "The module runs MXFillDBPerf against the logs. The NFR guide (PRF-BIG-30) "
        "expects a health check of Good or Great for these values:"
    )
    g.bullets(
        [
            "Parallel DB Update (Short Batch)",
            "Parallel DB Update",
            "Batch Rate",
            "Parallel Parsing",
        ]
    )
    g.para("The analyzer CSV holds the data. Open it in Excel to filter or pivot it.")
    g.label("Discussion")
    g.bullets(
        [
            "How quickly are client reports being processed?",
            "Does the rate change over the capture window?",
        ]
    )


# ---------------------------------------------------------------------------


def build(output: Path = OUTPUT) -> Path:
    if not TEMPLATE.is_file():
        raise FileNotFoundError(TEMPLATE)
    document = Document(str(TEMPLATE))
    update_front_matter(document)
    # Collect the exemplar paragraphs before the template's exercises go.
    guide = Guide(document)
    drop_template_exercises(document)
    write_content(guide)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def refresh_contents(path: Path) -> None:
    powershell = shutil.which("powershell")
    if powershell is None or not CONTENTS_SCRIPT.is_file():
        print("Contents not refreshed: open the guide in Word and update the field.")
        return
    subprocess.run(
        [
            powershell,
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(CONTENTS_SCRIPT),
            "-DocPath",
            str(path),
        ],
        check=True,
    )


def main() -> int:
    path = build()
    print(f"Wrote {path}")
    refresh_contents(path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
