#!/usr/bin/env python3
# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""
get_filldb_performance.py -- FillDB Performance Parser

Produces a Word document (chart + template-style table) and a PNG chart.

The input may be any of:
  * A raw FillDB performance log (FillDBperf.log) or a folder containing the
    rolling set (FillDBperf.log, FillDBperf.log_0, ...). These are combined in
    chronological order and analyzed with Utilities/MXFillDBPerf.exe, so the
    manual "run the analyzer first" step can't be forgotten.
  * An already-analyzed CSV with columns:
        FillDB Object, Count, Time (ms), Rate/s, Health Check

Usage:
  python get_filldb_performance.py --input C:\\Temp\\...\\FillDBperf.log
  python get_filldb_performance.py --input C:\\Temp\\perf.csv
  python get_filldb_performance.py --input C:\\Temp\\perf.csv --word-output filldb.performance.docx
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import Lib.Core.report_log as log
from Lib.Analysis import filldb_perf_log
from Lib.Core.document_stamp import (
    add_stamp_arguments,
    stamp_docx_document,
    stamp_from_args,
)
from Lib.Core.input_utils import (
    clean_user_path,
    discover_docx_template,
)
from Lib.Core.input_utils import (
    resolve_existing_docx as shared_resolve_existing_docx,
)
from Lib.Core.input_utils import (
    resolve_output_docx as shared_resolve_output_docx,
)
from Lib.Core.input_utils import (
    resolve_output_png as shared_resolve_output_png,
)
from Lib.Themes.word_stub_utils import (
    add_stub_front_matter,
    apply_standalone_narrow_margins,
    create_stub_document_from_template,
    set_picture_alt_text,
)

CSV_COLUMNS = ["FillDB Object", "Count", "Time (ms)", "Rate/s", "Health Check"]
PERFORMANCE_COLUMNS = CSV_COLUMNS
COUNT_COLUMNS = ("Count", "Time (ms)", "Rate/s")
# Analyzer statistics shown in Capture Summary, labeled as the analyzer labels them.
ANALYZER_STATISTICS = (
    "Duration (s)",
    "Throughput (lines/s)",
    "FillDB Duration (hh:mm:ss)",
    "FillDB Parsing Threads",
    "FillDB Database Threads",
)
TABLE_WIDTHS = [3480, 2820, 1160, 873, 2282]
TABLE_WIDTH = sum(TABLE_WIDTHS)
SUMMARY_COLUMNS = ["Statistic", "Value"]
SUMMARY_TABLE_WIDTHS = [3200, TABLE_WIDTH - 3200]
TABLE_STYLE_ID = "BigFixContent"
# Every visual aspect of the data table -- navy header, hairline borders, row
# banding -- comes from the shared BigFixContent style (single source of truth),
# so this module hand-rolls no inline table formatting. We also drop per-row
# health tints: the FillDB chart conveys per-object health, and the Health Check
# cell keeps its colored text.
MUTED = "#595959"
STATUS_COLORS = {
    "great": "#187A43",
    "good": "#B45309",
    "poor": "#C2410C",
}
STATUS_TABLE_STYLES = {
    "great": {"fill": "D5F5E3", "text": "187A43"},
    "good": {"fill": "FFF8E7", "text": "B45309"},
    "poor": {"fill": "FDECEA", "text": "C2410C"},
    "unknown": {"fill": "F0F2F5", "text": "0B1640"},
}
_CAPTURE_STAT_NAMES = {
    "lines processed",
    "filldb start time",
    "filldb end time",
    "start time",
    "end time",
    "duration",
}


def resolve_filldb_input(candidate_text: str) -> Path:
    """Resolve a FillDB input that may be a raw performance log, a folder
    containing the rolling log set, or an already-analyzed CSV.

    Type-specific handling (combine + run MXFillDBPerf.exe vs. use the CSV
    directly) happens later in filldb_perf_log.ensure_analysis_csv; here we only
    confirm the path exists.
    """
    text = clean_user_path(candidate_text)
    if not text:
        raise ValueError("FillDB performance input cannot be blank.")
    candidate = Path(text).expanduser()
    if not candidate.exists():
        display = candidate.resolve(strict=False)
        message = [
            "FillDB performance input was not found.",
            f"Entered value : {text}",
            f"Checked path  : {display}",
        ]
        if not candidate.is_absolute():
            message.append(f"Current folder: {Path.cwd()}")
        raise ValueError("\n".join(message))
    return candidate.resolve()


def resolve_existing_docx(candidate_text: str) -> Path:
    return shared_resolve_existing_docx(candidate_text, "style template")


def discover_template() -> Path | None:
    return discover_docx_template(style_id=TABLE_STYLE_ID, script_file=__file__)


def resolve_output_docx(candidate_text: str, default_name: str) -> Path:
    return shared_resolve_output_docx(candidate_text, default_name)


def resolve_output_png(candidate_text: str, default_path: Path) -> Path:
    return shared_resolve_output_png(candidate_text, default_path)


def require_input(arg_value) -> Path:
    if not clean_user_path(arg_value):
        raise ValueError("FillDB performance input is required. Pass it with --input.")
    return resolve_filldb_input(arg_value)


def resolve_optional_template(arg_value) -> Path | None:
    if arg_value:
        return resolve_existing_docx(arg_value)

    template_path = discover_template()
    if template_path:
        return template_path

    log.warn("No style template was found. Using a blank Word document.")
    return None


def parse_filldb_csv(input_path: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    with input_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError("CSV is empty or missing a header row.")

        missing = [column for column in CSV_COLUMNS if column not in reader.fieldnames]
        if missing:
            raise ValueError(
                "CSV is missing expected column(s): "
                + ", ".join(missing)
                + "\nAvailable columns: "
                + ", ".join(reader.fieldnames)
            )

        for raw_row in reader:
            values = {
                column: (raw_row.get(column) or "").strip() for column in CSV_COLUMNS
            }
            if not any(values.values()):
                continue
            if not values["FillDB Object"]:
                continue
            rows.append(values)

    if not rows:
        raise ValueError(
            "No FillDB performance rows were found after removing blank rows."
        )

    return rows


def parse_number(value: str) -> float:
    cleaned = str(value or "").replace(",", "").strip()
    if not cleaned:
        return 0.0
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def health_status(value: str) -> str:
    text = str(value or "").lower()
    if "poor" in text:
        return "poor"
    if "good" in text:
        return "good"
    if "great" in text:
        return "great"
    return "unknown"


def filldb_assessment_note(rows: list[dict[str, str]]) -> str:
    """Explain a Not Assessed FillDB score when the capture cannot be scored."""
    scoreable = any(
        health_status(row.get("Health Check", "")) in {"great", "good", "poor"}
        for row in object_performance_rows(rows)
    )
    if scoreable:
        return ""
    return (
        "Not Assessed: No FillDB object in this capture had a recognized health "
        "result, so there was not enough data to calculate a health score."
    )


def clean_object_name(value: str) -> str:
    text = str(value or "").strip()
    return text.removesuffix(":")


def format_count(value: float) -> str:
    value = float(value or 0)
    return f"{value:,.0f}"


def format_table_number(value: str) -> str:
    """Add thousands separators to a plain number. Leave any other text alone."""
    text = str(value).strip()
    try:
        number = float(text)
    except ValueError:
        return value
    if text.lstrip("-").isdigit():
        return f"{int(number):,}"
    return f"{number:,}"


def is_capture_stat_row(row: dict[str, str]) -> bool:
    name = clean_object_name(row.get("FillDB Object", "")).strip().lower()
    return name in _CAPTURE_STAT_NAMES


def object_performance_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """Object rows only. The analyzer CSV ends with a Statistic and Value block."""
    objects: list[dict[str, str]] = []
    for row in rows:
        name = clean_object_name(row.get("FillDB Object", "")).strip().lower()
        if name == "statistic":
            break
        if not is_capture_stat_row(row):
            objects.append(row)
    return objects


def titlecase_label(value: str) -> str:
    """Normalize a FillDB object name for display.

    The analyzer's object names mix casing ("Batch Rate" vs. "short property
    results"), which looks sloppy in the chart and table. Capitalize the first
    letter of each word while preserving the rest, so acronyms like "DB" and
    already-capitalized words survive (avoids turning "DB" into "Db").
    """
    text = clean_object_name(value)
    return " ".join(
        word[:1].upper() + word[1:] if word else word for word in text.split(" ")
    )


def shorten_label(value: str, max_chars: int = 30) -> str:
    text = clean_object_name(value)
    return text if len(text) <= max_chars else text[: max_chars - 1].rstrip() + "..."


def find_stat_value(rows: list[dict[str, str]], stat_name: str) -> float:
    wanted = clean_object_name(stat_name).strip().lower()
    for row in rows:
        object_name = clean_object_name(row.get("FillDB Object", "")).strip().lower()
        if object_name == wanted:
            return parse_number(row.get("Count", ""))
    return 0.0


def find_stat_text(rows: list[dict[str, str]], stat_name: str) -> str:
    """Return a statistic row's value as text (the analyzer stores it in the
    Count column), with any trailing timezone offset removed so the date
    parsers downstream can read it."""
    wanted = clean_object_name(stat_name).strip().lower()
    for row in rows:
        object_name = clean_object_name(row.get("FillDB Object", "")).strip().lower()
        if object_name == wanted:
            text = (row.get("Count") or "").strip()
            return re.sub(r"\s*(?:[+-]\d{2}:?\d{2}|Z)$", "", text).strip()
    return ""


def parse_metadata_datetime(value: str) -> datetime | None:
    text = clean_user_path(value)
    if not text:
        return None

    formats = (
        "%m-%d-%Y %H:%M:%S",
        "%m-%d-%Y %H:%M",
        "%m-%d-%Y",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def format_metadata_date(value: str) -> str:
    parsed = parse_metadata_datetime(value)
    if parsed is None:
        text = clean_user_path(value)
        return text if text else "Not specified"
    if parsed.hour == 0 and parsed.minute == 0 and parsed.second == 0:
        return parsed.strftime("%m-%d-%Y")
    return parsed.strftime("%m-%d-%Y %H:%M")


def calculate_capture_duration(start_value: str, end_value: str) -> str:
    start = parse_metadata_datetime(start_value)
    end = parse_metadata_datetime(end_value)
    if start is None or end is None or end < start:
        return ""

    total_seconds = int((end - start).total_seconds())
    days, remainder = divmod(total_seconds, 86400)
    hours, remainder = divmod(remainder, 3600)
    minutes, _ = divmod(remainder, 60)

    parts = []
    if days:
        parts.append(f"{days} day{'s' if days != 1 else ''}")
    if hours:
        parts.append(f"{hours} hour{'s' if hours != 1 else ''}")
    if minutes or not parts:
        parts.append(f"{minutes} minute{'s' if minutes != 1 else ''}")
    return ", ".join(parts[:2])


def build_capture_metadata(start_value: str, end_value: str) -> dict[str, str]:
    duration = calculate_capture_duration(start_value, end_value)
    return {
        "generated": datetime.now().strftime("%m/%d/%Y %H:%M"),
        "start": format_metadata_date(start_value),
        "end": format_metadata_date(end_value),
        "duration": duration if duration else "Not specified",
    }


def generate_filldb_chart(
    rows: list[dict[str, str]],
    output_path: Path,
    metadata: dict[str, str],
) -> None:
    """Render the FillDB Performance infographic from the design HTML.

    Builds the score band, object-health distribution, and health-colored workload
    bars from the analyzed rows and screenshots the design template via
    html_charts.
    """
    from Lib.Design import html_charts

    # Raw FillDB emits Great/Good/Poor/Unknown; the chart displays the report's
    # four standard status words.
    C_GREAT, C_GOOD, C_POOR, C_NODATA = "#187A43", "#B45309", "#C2410C", "#0B1640"
    status_color = {
        "great": C_GREAT,
        "good": C_GOOD,
        "poor": C_POOR,
        "unknown": C_NODATA,
    }

    assessed_rows = [r for r in rows if r.get("Health Check", "").strip()]
    great_count = sum(
        1 for r in assessed_rows if health_status(r["Health Check"]) == "great"
    )
    good_count = sum(
        1 for r in assessed_rows if health_status(r["Health Check"]) == "good"
    )
    review_count = sum(
        1 for r in assessed_rows if health_status(r["Health Check"]) == "poor"
    )
    unknown_count = sum(
        1 for r in assessed_rows if health_status(r["Health Check"]) == "unknown"
    )
    assessed_total = great_count + good_count + review_count + unknown_count

    lines_processed = find_stat_value(rows, "Lines processed")
    total_rows = lines_processed or sum(
        parse_number(r.get("Count", ""))
        for r in rows
        if parse_number(r.get("Time (ms)", "")) > 0
    )

    # Workload bars, longest processing time first.
    ranked = sorted(
        assessed_rows, key=lambda r: parse_number(r.get("Time (ms)", "")), reverse=True
    )
    times_min = [parse_number(r.get("Time (ms)", "")) / 60000 for r in ranked]
    max_min = max(times_min, default=0) or 1
    bars = [
        {
            "label": clean_object_name(r.get("FillDB Object", "")),
            "width": f"{m / max_min * 100:.1f}%",
            "mins": f"{m:.1f}",
            "color": status_color.get(health_status(r.get("Health Check", "")), C_NODATA),
        }
        for r, m in zip(ranked, times_min)
    ]

    # FillDB Health Score: Great x100 + Good x67 + Poor x34.
    scoreable = great_count + good_count + review_count
    if scoreable > 0:
        score = max(
            0,
            min(
                100,
                round(
                    (great_count / scoreable) * 100
                    + (good_count / scoreable) * 67
                    + (review_count / scoreable) * 34
                ),
            ),
        )
        if score >= 85:
            cls, label, color, grad = (
                "st-healthy",
                "Healthy",
                C_GREAT,
                C_GREAT,
            )
        elif score >= 60:
            cls, label, color, grad = (
                "st-attention",
                "Needs Attention",
                C_GOOD,
                C_GOOD,
            )
        else:
            cls, label, color, grad = (
                "st-critical",
                "Critical",
                C_POOR,
                C_POOR,
            )
    else:
        score, cls, label, color, grad = (
            "N/A",
            "st-notassessed",
            "Not Assessed",
            C_NODATA,
            C_NODATA,
        )

    dist_defs = [
        ("Healthy", great_count, C_GREAT),
        ("Needs Attention", good_count, C_GOOD),
        ("Critical", review_count, C_POOR),
        ("Not Assessed", unknown_count, C_NODATA),
    ]
    dist_total = sum(c for _, c, _ in dist_defs) or 1
    dist = [
        {
            "count": c,
            "css": css,
            "label": f"{name} &middot; {c:,}",
            "width": f"{c / dist_total * 100:.1f}%",
        }
        for name, c, css in dist_defs
        if c > 0
    ]

    meta = f"{format_count(total_rows)} lines processed"
    if metadata.get("duration") and metadata["duration"] != "Not specified":
        meta += f" &middot; capture {metadata['duration']}"

    data = {
        "title": "FillDB Health Score",
        "meta": meta,
        "score": score,
        "scoreColor": color,
        "status": {"cls": cls, "label": label},
        "barPct": score if isinstance(score, int) else 0,
        "barColor": grad,
        "distLabel": f"Object health &middot; {assessed_total:,} assessed",
        "dist": dist,
        "barsLabel": "Workload by object &middot; processing time",
        "bars": bars,
        "assessmentNote": filldb_assessment_note(rows),
    }
    html_charts.render_chart("filldb", data, output_path)


def set_table_style(table, style_id: str = TABLE_STYLE_ID) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_style = tbl_pr.find(qn("w:tblStyle"))
    if tbl_style is None:
        tbl_style = OxmlElement("w:tblStyle")
        tbl_pr.insert(0, tbl_style)
    tbl_style.set(qn("w:val"), style_id)


def remove_child(parent, tag: str) -> None:
    child = parent.find(qn(tag))
    if child is not None:
        parent.remove(child)


def set_table_geometry(table, widths: list[int] | None = None) -> None:
    widths = widths or TABLE_WIDTHS
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    # Match the BigFixContent header row's alignment for the whole table.
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_child(tbl_pr, "w:tblLayout")

    grid = tbl.tblGrid
    for grid_col in list(grid):
        grid.remove(grid_col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)


def set_table_look(table) -> None:
    """Enable the conditional formatting BigFixContent defines -- navy first row
    and horizontal row banding -- by setting the table's tblLook. A table built
    programmatically has no tblLook, so Word otherwise ignores the style's
    first-row/banding rules and the table wouldn't match the rest of the report.
    Mirrors the 04A0 look the template's own tables and the XML injectors use."""
    tbl_pr = table._tbl.tblPr
    look = tbl_pr.find(qn("w:tblLook"))
    if look is None:
        look = OxmlElement("w:tblLook")
        tbl_pr.append(look)
    for attr, val in (
        ("w:val", "04A0"),
        ("w:firstRow", "1"),
        ("w:lastRow", "0"),
        ("w:firstColumn", "1"),
        ("w:lastColumn", "0"),
        ("w:noHBand", "0"),
        ("w:noVBand", "1"),
    ):
        look.set(qn(attr), val)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.insert(0, tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1

    run = paragraph.add_run("" if text is None else str(text))
    run.bold = bold
    # Font family and size come from the shared BigFixContent
    # style once merged into the report, so we don't pin a face here.
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def build_word_table(document: DocxDocument, rows: list[dict[str, str]]) -> None:
    rows = object_performance_rows(rows)
    table = document.add_table(rows=1, cols=len(PERFORMANCE_COLUMNS))
    table.autofit = True
    # Borders, navy header fill and row banding all come from the shared
    # BigFixContent style (single source of truth); we only enable its
    # conditional formatting via tblLook. No inline shading/borders here, so
    # this table matches every other table in the report.
    set_table_style(table)
    set_table_geometry(table, TABLE_WIDTHS)
    set_table_look(table)

    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    set_row_cant_split(header_row)

    for idx, column_name in enumerate(PERFORMANCE_COLUMNS):
        cell = header_row.cells[idx]
        set_cell_width(cell, TABLE_WIDTHS[idx])
        # Navy fill + white bold text come from BigFixContent's firstRow.
        set_cell_text(cell, column_name, bold=True)

    for item in rows:
        row = table.add_row()
        set_row_cant_split(row)
        status = health_status(item.get("Health Check", ""))
        status_style = STATUS_TABLE_STYLES.get(status)
        # Row banding comes from the shared BigFixContent style applied to the
        # table, so every data table in the report bands identically. We don't
        # tint rows by health here -- the FillDB chart conveys per-object health,
        # and the Health Check cell keeps its colored text below.
        for idx, column_name in enumerate(PERFORMANCE_COLUMNS):
            cell = row.cells[idx]
            set_cell_width(cell, TABLE_WIDTHS[idx])
            is_status_cell = column_name == "Health Check" and status != "unknown"
            value = item.get(column_name, "")
            if column_name == "FillDB Object":
                value = titlecase_label(value)
            elif column_name in COUNT_COLUMNS:
                value = format_table_number(value)
            set_cell_text(
                cell,
                value,
                bold=is_status_cell,
                color=status_style["text"] if is_status_cell and status_style else None,
            )


def build_capture_summary_table(
    document: DocxDocument,
    rows: list[dict[str, str]],
    metadata: dict[str, str],
) -> None:
    heading = document.add_heading("Capture Summary", level=3)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)

    lines_processed = find_stat_value(rows, "Lines processed")
    summary_rows = [
        ("Report Generated", metadata.get("generated", "Not specified")),
        ("Collection Start", metadata.get("start", "Not specified")),
        ("Collection End", metadata.get("end", "Not specified")),
        ("Capture Duration", metadata.get("duration", "Not specified")),
        (
            "Lines Processed",
            format_count(lines_processed) if lines_processed else "Not available",
        ),
    ]
    for label in ANALYZER_STATISTICS:
        value = find_stat_text(rows, label)
        if value:
            summary_rows.append((label, format_table_number(value)))

    table = document.add_table(rows=1, cols=len(SUMMARY_COLUMNS))
    table.autofit = True
    set_table_style(table)
    set_table_geometry(table, SUMMARY_TABLE_WIDTHS)
    set_table_look(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    set_row_cant_split(header_row)
    for idx, column_name in enumerate(SUMMARY_COLUMNS):
        cell = header_row.cells[idx]
        set_cell_width(cell, SUMMARY_TABLE_WIDTHS[idx])
        set_cell_text(cell, column_name, bold=True)

    for label, value in summary_rows:
        row = table.add_row()
        set_row_cant_split(row)
        for idx, text in enumerate((label, value)):
            cell = row.cells[idx]
            set_cell_width(cell, SUMMARY_TABLE_WIDTHS[idx])
            set_cell_text(cell, text, bold=(idx == 0))


def clear_document_body(document: DocxDocument) -> None:
    body = document._body._element
    sect_pr = None
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            sect_pr = child
            continue
        body.remove(child)
    if sect_pr is not None and sect_pr.getparent() is None:
        body.append(sect_pr)


def usable_page_width(document: DocxDocument):
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    return width if width > Inches(1) else Inches(6.8)


def add_score_method_note(document: DocxDocument) -> None:
    heading = document.add_heading("How This Score Is Calculated", level=3)
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(3)

    intro = document.add_paragraph(
        "The FillDB Health Score is based on the percentage of assessed FillDB "
        "objects in each analyzer health band:"
    )
    intro.paragraph_format.space_after = Pt(3)

    for text in (
        "Great: 100 points",
        "Good: 67 points",
        "Poor: 34 points",
    ):
        bullet = document.add_paragraph(f"- {text}")
        bullet.paragraph_format.space_after = Pt(1)

    for text in (
        "Formula: Great% x 100 + Good% x 67 + Poor% x 34",
        "This score covers the FillDB capture only.",
        "Rows without a recognized Health Check result are shown for context "
        "but are not included in the score.",
        "Score bands: 85-100 Healthy, 60-84 Needs Attention, 0-59 Critical.",
    ):
        para = document.add_paragraph(text)
        para.paragraph_format.space_after = Pt(3)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def write_word_output(
    rows: list[dict[str, str]],
    output_path: Path,
    template_path: Path | None = None,
    chart_path: Path | None = None,
    metadata: dict[str, str] | None = None,
    stamp: str = "",
) -> None:
    document = (
        create_stub_document_from_template(template_path) if template_path else Document()
    )
    apply_standalone_narrow_margins(document)
    add_stub_front_matter(document, "FillDB Performance Analysis")
    if stamp:
        stamp_docx_document(document, stamp)
    if chart_path:
        picture = document.add_picture(str(chart_path), width=usable_page_width(document))
        set_picture_alt_text(
            picture,
            "FillDB performance chart showing health score, status distribution, "
            "and workload by object.",
            "FillDB Performance Chart",
        )
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    assessment_note = filldb_assessment_note(rows)
    if assessment_note:
        note = document.add_paragraph(assessment_note)
        note.paragraph_format.space_before = Pt(3)
        note.paragraph_format.space_after = Pt(6)
    if chart_path:
        add_score_method_note(document)
    build_word_table(document, rows)
    build_capture_summary_table(document, rows, metadata or {})
    document.save(str(output_path))


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Analyze a FillDB performance log (or accept an analyzed CSV) and build a Word table for template insertion."
    )
    parser.add_argument(
        "--input",
        "-i",
        help="Path to a FillDB performance log, a folder of rolling logs, or an analyzed CSV",
    )
    parser.add_argument("--word-output", "-w", help="Output .docx path")
    parser.add_argument(
        "--png-output", "-p", help="Output FillDB visualization .png path"
    )
    parser.add_argument(
        "--collection-start",
        help="FillDB data collection start date/time, for example 05/15/2026 08:00",
    )
    parser.add_argument(
        "--collection-end",
        help="FillDB data collection end date/time, for example 05/16/2026 08:00",
    )
    parser.add_argument(
        "--template",
        "-t",
        help="Optional style template .docx override. If omitted, a nearby template with the required table style is auto-discovered.",
    )
    add_stamp_arguments(parser)

    args = parser.parse_args()

    log.title("FillDB Performance")

    try:
        input_path = require_input(args.input)
        input_stem = input_path.name if input_path.is_dir() else input_path.stem
        default_docx = f"{input_stem}.filldb.performance.docx"
        output_path = resolve_output_docx(args.word_output, default_docx)
        chart_path = resolve_output_png(args.png_output, output_path.with_suffix(".png"))
        template_path = resolve_optional_template(args.template)
        collection_start = clean_user_path(args.collection_start)
        collection_end = clean_user_path(args.collection_end)

        # Accept a raw FillDB perf log (or a folder of rolling logs) and run
        # MXFillDBPerf.exe ourselves, so the manual "run the analyzer first"
        # step can't be skipped. An already-analyzed CSV is used as-is.
        csv_path = filldb_perf_log.ensure_analysis_csv(
            input_path,
            dest_csv=output_path.with_suffix(".source.csv"),
            log=log.step,
        )
        rows = parse_filldb_csv(csv_path)
        # When the collection window wasn't supplied, fall back to the FillDB
        # start/end the analyzer recorded, so the chart header is populated and
        # correct on the raw-log path even without the UI's auto-detection.
        if not clean_user_path(collection_start):
            collection_start = find_stat_text(rows, "FillDB start time")
        if not clean_user_path(collection_end):
            collection_end = find_stat_text(rows, "FillDB end time")
        metadata = build_capture_metadata(collection_start, collection_end)
        generate_filldb_chart(rows, chart_path, metadata)
        write_word_output(
            rows,
            output_path,
            template_path=template_path,
            chart_path=chart_path,
            metadata=metadata,
            stamp=stamp_from_args(args),
        )

        log.ok("FillDB performance report built")
        log.detail("Rows", len(rows))
        if template_path:
            log.detail("Template", template_path)
        log.detail(
            "Collection",
            f"{metadata['start']} - {metadata['end']} ({metadata['duration']})",
        )
        log.result("Chart (PNG)", chart_path)
        log.result("Word document", output_path)
        return 0

    except KeyboardInterrupt:
        log.warn("Canceled.")
        return 1
    except Exception as exc:
        log.error(str(exc))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
