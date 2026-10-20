# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT

"""Known-answer evidence and real standalone report delivery checks."""

from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

import get_loop_health as loop
import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageStat

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = Path(__file__).parent / "fixtures" / "client_delivery.json"
CASE = json.loads(FIXTURE.read_text(encoding="utf-8"))
DELIVERY = pytest.mark.skipif(
    os.environ.get("WORKBENCH_RUN_DELIVERY") != "1",
    reason="Set WORKBENCH_RUN_DELIVERY=1 for real browser/report builds",
)


def write_evidence(folder: Path, *, include_intervals: bool = True) -> dict[str, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    capture = folder / "Client_Loop.csv"
    with capture.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(CASE["loop_headers"])
        writer.writerows(CASE["loop_rows"])
    deployment = folder / "Deployment_Information"
    deployment.mkdir(exist_ok=True)
    lines = ["PROPERTY INFORMATION", "Number of Properties: 7"]
    if include_intervals:
        lines.extend(
            [
                "//===== Active Properties by Evaluation Interval ================",
                "Column Headers: Property Name, Source Analysis Name, Source Analysis Site Name, Site ID, Analysis ID, Property ID, Issued By, Evaluation Interval, Content Source",
                " ",
            ]
        )
        for index, (name, interval, source) in enumerate(CASE["properties"], 1):
            lines.extend(
                [
                    name,
                    "QA Analysis",
                    "QA Site",
                    "1",
                    "2",
                    str(index),
                    "qa",
                    interval,
                    source,
                ]
            )
    (deployment / "Property_Info_QA.log").write_text("\n".join(lines), encoding="utf-8")
    offenders = folder / "top10-query-results.csv"
    with offenders.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "Occurrences",
                "AvgClientProcessingSeconds",
                "TotalClientProcessingSeconds",
                "MaxClientProcessingSeconds",
                "Status",
                "Result",
            ]
        )
        writer.writerow(
            [
                2,
                60,
                120,
                90,
                "OK",
                "Analysis Name=QA Analysis | Display Name=QA Every Report | Name=QA Every Report | Source Name=QA Every Report | Source ID=1",
            ]
        )
        # A site evaluation and a retrieved property, as the lab Top 10 has.
        writer.writerows(CASE["offender_rows"])
    # The UDP BESAPI capture has no header row.
    udp = folder / "UDP_Notifications_QA.xml"
    udp.write_text(
        '<?xml version="1.0" encoding="UTF-8"?><BESAPI><Query><Result>'
        + "".join(
            f'<Answer type="string">{line}</Answer>' for line in CASE["udp_answers"]
        )
        + "</Result></Query></BESAPI>",
        encoding="utf-8",
    )
    settings = [
        "CLIENT SETTINGS INFORMATION",
        "//===== CLIENT SETTINGS ==================================",
        "Column Headers: Setting Name, Setting Value, Number of Machines",
    ]
    for value, count in CASE["minimum_analysis_interval_rows"]:
        settings += ["_BESClient_Report_MinimumAnalysisInterval ", f" {value}", count]
    settings.append("//===== ICMP CONTROLS ==================================")
    (deployment / "Client_Settings_Info_QA.log").write_text(
        "\n".join(settings) + "\n", encoding="utf-8"
    )
    return {
        "capture": capture,
        "deployment": deployment,
        "offenders": offenders,
        "udp": udp,
    }


def calculated_metrics():
    frame = pd.DataFrame(CASE["loop_rows"], columns=CASE["loop_headers"])
    computer, maximum, average = loop.auto_detect_columns(frame)
    rows, missing, errors = loop.process_rows(frame, computer, maximum, average)
    metrics, _ = loop.calculate_metrics(
        rows,
        len(frame),
        missing,
        errors,
        900,
        3600,
        date(2026, 10, 1),
        date(2026, 10, 2),
        datetime(2026, 10, 3, 12),
        good_threshold_seconds=1800,
    )
    return metrics


def test_known_answers_include_threshold_boundaries_and_missing_evidence():
    metrics = calculated_metrics()
    for field, expected in CASE["expected_metrics"].items():
        actual = getattr(metrics, field)
        if isinstance(expected, (float, int)):
            assert actual == pytest.approx(expected), field
        else:
            assert actual == expected, field


def test_known_property_interval_answers(tmp_path):
    evidence = write_evidence(tmp_path)
    intervals = loop.parse_property_intervals(evidence["deployment"])
    assert intervals is not None
    assert len(intervals.properties) == 7
    assert loop.frequent_property_count(intervals) == 3
    assert loop.evaluation_frequency_rows(intervals) == CASE["expected_frequency"]


@pytest.fixture(scope="module")
def delivered(tmp_path_factory):
    base = Path(
        os.environ.get(
            "WORKBENCH_DELIVERY_OUTPUT", str(tmp_path_factory.mktemp("delivery"))
        )
    ).resolve()
    base.mkdir(parents=True, exist_ok=True)
    (base / "delivery-summary.json").unlink(missing_ok=True)
    folder = Path(tempfile.mkdtemp(prefix="run-", dir=base))
    evidence = write_evidence(folder / "evidence")
    output = folder / "report"
    output.mkdir(exist_ok=True)
    command = [
        sys.executable,
        str(ROOT / "Get" / "get_loop_health.py"),
        "--input-file",
        str(evidence["capture"]),
        "--output-folder",
        str(output),
        "--deployment-info-folder",
        str(evidence["deployment"]),
        "--evaluation-times-file",
        str(evidence["offenders"]),
        "--udp-notifications-file",
        str(evidence["udp"]),
        "--report-start-date",
        "10-01-2026",
        "--report-end-date",
        "10-02-2026",
        "--include-automated-findings",
    ]
    with (folder / "build.log").open("w", encoding="utf-8") as log:
        result = subprocess.run(
            command,
            cwd=ROOT,
            env=dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8"),
            stdout=log,
            stderr=subprocess.STDOUT,
            stdin=subprocess.DEVNULL,
            timeout=180,
        )
    assert result.returncode == 0, (folder / "build.log").read_text(encoding="utf-8")
    return output


def artifact(folder, suffix):
    found = list(folder.glob("*" + suffix))
    assert len(found) == 1, found
    return found[0]


@DELIVERY
def test_real_report_workbook_and_chart_agree(delivered):
    workbook = load_workbook(artifact(delivered, ".xlsx"), data_only=True)
    try:
        summary = {
            row[0]: row[1]
            for row in workbook["Health Summary"].iter_rows(values_only=True)
            if row[0] is not None
        }
        assert summary["Devices Scored"] == 6
        assert summary["Devices Reporting"] == 4
        assert summary["Overall Status"] == "CRITICAL"
        assert summary["Average Loop Health Score"] == "56 / 100"
        frequencies = list(
            workbook["Evaluation Frequency"].iter_rows(min_row=2, values_only=True)
        )
        assert [list(row) for row in frequencies] == CASE["expected_frequency"]
        assert "Repeated Eval Cost" not in workbook.sheetnames
        offenders = list(workbook["Evaluation Offenders"].iter_rows(values_only=True))
        interval_col = list(offenders[0]).index("Interval")
        assert CASE["expected_top10"]["interval"] in [
            row[interval_col] for row in offenders[1:]
        ]
    finally:
        workbook.close()
    document = Document(artifact(delivered, ".docx"))
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]
    frequency_table = next(table for table in tables if table[0][0] == "Interval")
    assert frequency_table[1:] == CASE["expected_frequency"]
    assert not any(table[0][0] == "Property" for table in tables)
    top10 = next(table for table in tables if table[0][0] == "Rank")
    assert top10[0][-1] == "Interval"
    assert [
        CASE["expected_top10"]["evaluation"],
        CASE["expected_top10"]["interval"],
    ] in [row[3:5] for row in top10[1:]]
    assert [row[3:5] for row in top10[1:]] == CASE["expected_top10_rows"]
    coverage = next(
        table
        for table in tables
        if table[0][:2] == ["Status", "Devices / Total"]
        and table[1][0] == "Receiving UDP Pings"
    )
    assert coverage[1:] == CASE["expected_udp_coverage"]
    paragraphs = [p.text for p in document.paragraphs]
    assert CASE["expected_minimum_analysis_interval"] in paragraphs
    findings = next(table for table in tables if table[0][0] == "Finding ID")
    assert len(findings) == 2
    assert findings[1][0] == "CLT-01"
    assert "3 active properties are set to every report" in findings[1][1]
    sidecar = json.loads(
        artifact(delivered, ".chartdata.json").read_text(encoding="utf-8")
    )
    data = sidecar
    assert str(data["score"]) == "56"
    assert data["barPct"] == 56
    # Healthy, Warning, Critical over 60 and the hatched no usable data part.
    assert [int(segment["num"]) for segment in data["dist"]] == [1, 2, 1, 2]
    with Image.open(artifact(delivered, ".png")) as image:
        assert image.width == 2250
        assert image.height >= 300
        assert max(ImageStat.Stat(image.convert("RGB")).var) > 100
    report = {
        "schema": 1,
        "case": CASE["case"],
        "metrics": {
            field: getattr(calculated_metrics(), field)
            for field in CASE["expected_metrics"]
        },
        "frequency": [list(row) for row in frequencies],
        "top10": [row[3:5] for row in top10[1:]],
        "chart": {key: data[key] for key in ("score", "barPct", "dist")},
        "finding_ids": sorted(row[0] for row in findings[1:]),
    }
    (delivered.parent / "delivery-summary.json").write_text(
        json.dumps(report, indent=2), encoding="utf-8"
    )
    (delivered.parent.parent / "delivery-summary.json").write_text(
        json.dumps(report, indent=2), encoding="utf-8"
    )


@DELIVERY
def test_real_chart_has_no_script_errors_or_clipped_text(delivered):
    from playwright.sync_api import sync_playwright

    from Lib.Design import html_charts

    payload = json.loads(
        artifact(delivered, ".chartdata.json").read_text(encoding="utf-8")
    )
    html = html_charts._build_html(payload["chart"], payload, card_width=2250)
    (delivered.parent / "chart.html").write_text(html, encoding="utf-8")
    with sync_playwright() as playwright:
        browser = html_charts._launch_chart_browser(playwright)
        try:
            page = browser.new_page(viewport={"width": 2500, "height": 2400})
            errors = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.set_content(html, wait_until="load")
            page.evaluate("document.fonts.ready")
            page.wait_for_function("window.__RENDERED__ === true")
            assert not errors, errors
            problems = page.evaluate("""() => {
                const card = document.querySelector('#mount');
                if (!card) return ['missing chart card'];
                const box = card.getBoundingClientRect();
                const bad = [];
                const walker = document.createTreeWalker(card, NodeFilter.SHOW_TEXT);
                while (walker.nextNode()) {
                    const node = walker.currentNode;
                    if (!node.textContent.trim()) continue;
                    const style = getComputedStyle(node.parentElement);
                    if (style.display === 'none' || style.visibility === 'hidden') continue;
                    const range = document.createRange(); range.selectNodeContents(node);
                    for (const rect of range.getClientRects()) {
                        if (!rect.width || !rect.height) continue;
                        if (rect.left < box.left - 2 || rect.right > box.right + 2 || rect.top < box.top - 2 || rect.bottom > box.bottom + 2) bad.push(node.textContent.trim());
                    }
                }
                return bad;
            }""")
            page.screenshot(
                path=str(delivered.parent / "chart-layout.png"), full_page=True
            )
            assert not problems, problems
        finally:
            browser.close()
