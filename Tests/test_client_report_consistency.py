# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
"""Client navigation and content must survive standalone and full composition."""

from pathlib import Path
from zipfile import ZipFile

import get_loop_health as loop
import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from test_loop_evaluation_frequency import _healthy_metrics, _offender
from test_loop_retrieved_properties import _summary
from test_mini_report_design_fixes import _capture_details


@pytest.mark.parametrize("with_graphics", [False, True])
def test_client_navigation_and_compact_headers_with_or_without_graphics(
    monkeypatch, tmp_path, with_graphics
):
    _capture_details(monkeypatch)
    png = tmp_path / "summary.png"
    if with_graphics:
        Image.new("RGB", (30, 10), "white").save(png)
    output = tmp_path / "client.docx"
    loop.generate_word(
        _healthy_metrics(), png, output,
        loop_offenders=[_offender(9.2, "A site")],
        property_intervals=_summary(tmp_path / "capture"),
        include_automated_findings=True,
    )
    document = Document(output)
    headings = {p.text: p for p in document.paragraphs if p.style.name.startswith("Heading")}
    for label in ("How the Score Is Calculated", "Reading the Infographic", "Loop Health", "Loop Cost"):
        assert headings[label].style.name == "Heading 3"
        assert headings[label].style.font.size.pt == 20
        assert str(headings[label].style.font.color.rgb) == "000032"
    for label in ("Top 10 Most Expensive Evaluations", "Evaluation Frequency"):
        assert headings[label].style.name == "Heading 4"
    guidance = document.element.body.xpath(
        './/w:sdt/w:sdtContent/w:p[w:pPr/w:pStyle]'
    )
    for paragraph in guidance:
        label = "".join(paragraph.xpath(".//w:t/text()"))
        style = paragraph.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val"))
        if label in {"Why Client Loop Speed Matters", "Four Starting Points for Reviewing Loop Speed"}:
            assert style == "Heading3"
        elif label == "Action Latency" or label.startswith(("1. ", "2. ", "3. ", "4. ")):
            assert style == "Heading4"
    for label in ("Top 10 Most Expensive Evaluations", "Evaluation Frequency"):
        style = headings[label].style
        assert style.font.size.pt == 14
        assert style.font.bold
        assert str(style.font.color.rgb) == "007A8A"
        assert style.paragraph_format.keep_with_next
    style = next(s for s in document.styles if s.style_id == "BigFixContent").element
    spacing = next(
        node for node in style.findall(qn("w:tblStylePr"))
        if node.get(qn("w:type")) == "firstRow"
    ).find(qn("w:pPr")).find(qn("w:spacing"))
    assert spacing.get(qn("w:before")) == spacing.get(qn("w:after")) == "0"
    header_style = document.styles["Report Table Header"]
    assert header_style.paragraph_format.space_before.pt == 0
    assert header_style.paragraph_format.space_after.pt == 0
    for table in document.tables:
        if table.cell(0, 0).text in {"Infographic tile", "Finding ID"}:
            assert all(
                p.style.name == "Report Table Header"
                for cell in table.rows[0].cells for p in cell.paragraphs
            )
    text = " ".join(document.element.body.xpath(".//w:t/text()"))
    assert "workbook" not in text.lower()
    assert "spreadsheet" not in text.lower()
    assert "Excel sheet" not in text
    assert "BES Client Setting: Enable Command Polling" in text
    assert "which is available only from the team that produced this report" in text
    assert "even when its evaluation loop is fast" in text


def test_reinjection_preserves_h3_sections_and_one_copy_of_all_client_content(tmp_path):
    root = Path(__file__).resolve().parents[1]
    templates = [p for p in (root / "Templates").glob("*.dotx") if not p.name.startswith("OLD_")]
    if not templates:
        pytest.skip("Full-report composition belongs to the full application")
    import inject_loop_health as injector

    mini = tmp_path / "mini.docx"
    loop.generate_word(_healthy_metrics(), tmp_path / "missing.png", mini)
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    template = max(templates, key=lambda p: p.stat().st_mtime)
    injector.inject_loop_health(mini, template, first, log=lambda _: None)
    injector.inject_loop_health(mini, first, second, log=lambda _: None)
    with ZipFile(second) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": qn("w:p").split("}")[0][1:]}
    paragraphs = xml.xpath("//w:p", namespaces=ns)
    reading = [p for p in paragraphs if "".join(p.xpath(".//w:t/text()", namespaces=ns)) == "Reading the Infographic"]
    assert len(reading) == 1
    assert reading[0].find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "Heading3"
    for label in ("How the Score Is Calculated", "Loop Health", "Why Client Loop Speed Matters", "Four Starting Points for Reviewing Loop Speed"):
        matching = [p for p in paragraphs if "".join(p.xpath(".//w:t/text()", namespaces=ns)) == label]
        assert len(matching) == 1
        assert matching[0].find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "Heading3"
    titles = ["".join(p.xpath(".//w:t/text()", namespaces=ns)) for p in paragraphs]
    assert titles.count("Four Starting Points for Reviewing Loop Speed") == 1
    assert "How to Improve Loop Speed" not in titles
    guidance = etree.parse(str(root / "Lib/Design/client_guidance.xml"))
    for p in guidance.xpath("//w:p", namespaces=ns):
        text = "".join(p.xpath(".//w:t/text()", namespaces=ns))
        assert not text or text in titles
