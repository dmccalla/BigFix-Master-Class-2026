# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
"""Retain CPU Idle coverage for configured and unset populations."""

import get_loop_health as loop
import pytest
from docx import Document
from PIL import Image
from test_loop_evaluation_frequency import _healthy_metrics
from test_mini_report_design_fixes import _capture_details


@pytest.mark.parametrize("configured,unset", [(12, 0), (6, 6), (0, 12)])
def test_cpu_idle_preserves_configured_and_unset_devices(
    monkeypatch, tmp_path, configured, unset
):
    jobs = _capture_details(monkeypatch)
    rows = []
    if configured:
        rows.append(["100", "400", "20.0%", f"{configured} / 12"])
    if unset:
        rows.append(["Not Set", "Not Set", "Not Set", f"{unset} / 12"])
    summary = loop.AgentCpuIdleSummary(total_devices=12, settings_rows=rows)
    image = tmp_path / "summary.png"
    Image.new("RGB", (40, 10), "white").save(image)
    output = tmp_path / "report.docx"
    loop.generate_word(_healthy_metrics(), image, output, agent_cpu_idle_summary=summary)
    duty = [data for _, data, _ in jobs if data.get("kind") == "duty"]
    if configured:
        assert len(duty) == 1
        assert duty[0]["unset"] == unset
        assert sum(row["count"] for row in duty[0]["rows"]) + duty[0]["unset"] == 12
    else:
        assert not duty
        tables = [
            table
            for table in Document(output).tables
            if table.cell(0, 0).text == "Work Idle (ms)"
        ]
        assert len(tables) == 1
        assert "Not Set" in tables[0].cell(1, 0).text
        assert "12 / 12" in tables[0].cell(1, 3).text
