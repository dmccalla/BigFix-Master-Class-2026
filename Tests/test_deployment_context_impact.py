# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT

"""Synthetic downstream coverage for the ten repaired deployment queries.

These are response-shape tests, not a claim of live Session Relevance validation.
MC consumes existing General/Operator/Content/Action context, not the repaired
metric labels, external-site detail tuples, or property-section heading.
"""

from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import get_loop_health as loop
import pytest
from docx import Document
from openpyxl import load_workbook

CONTEXT = {
    "General_Info": ["Actionsite Size: 19,046 KB"],
    "Operator_Info": ["Number of Master Operators: 6"],
    "Action_Info": ["Number of open actions issued by master operators: 42"],
    "Content_Info": [
        "Number of External Sites: 57",
        "Number of Custom Sites: 23",
        "Number of Automatic Computer Groups: 6",
        "Number of Manual Computer Groups: 5",
        "Number of Fixlets: 5,701 / 34,130",
        "Number of Tasks: 1,008 / 2,144",
        "Number of Analyses: 121 / 6,702",
        "Number of Active Analyses: 140",
        "Number of Properties: 10,342",
        "Number of Actived Properties: 877",
        "Number of Custom Properties: 34",
        "Number of Master Operator Fixlets: 15",
        "Number of Master Operator Tasks: 30",
        "Total Number of Master Operator Fixlets and Tasks: 45",
        "Total Number of Master Operator Analyses: 10",
    ],
    "Relay_Server_Info": [],
    "Property_Info": [],
}

# Labels and the five-column tuple follow the repaired collector contracts.
METRICS = {
    "General_Info": [("Using Efficient MIME?", "True")],
    "Relay_Server_Info": [
        ("Maximum Number of Clients per Relay", "125"),
        ("Minimum Number of Clients per Relay", "3"),
        ("Number of Clients with Unknown Relay Distance", "2"),
    ],
    "Operator_Info": [
        ("Average Number of Computers Managed by Non-Master Operators", "24"),
        ("Number of Bes Roles", "7"),
    ],
    "Content_Info": [
        ("Average number of Components per Baselines", "12"),
        ("Maximum Number of Components per Baseline", "150"),
    ],
}
PROPERTY_HEADING = (
    "//===== Custom BES Properties with frequent evaluations (sub 15 minutes) "
    "================"
)


def _responses(empty: bool) -> dict[str, list[str | tuple[str, ...]]]:
    answers: dict[str, list[str | tuple[str, ...]]] = {
        stem: [
            f"{label}: {('False' if value == 'True' else '0') if empty else value}"
            for label, value in metrics
        ]
        for stem, metrics in METRICS.items()
    }
    answers["Content_Info"].append(
        (
            "Synthetic Patches & Updates",
            "17",
            "0" if empty else "123",
            "Not Reported" if empty else "Tue, 29 Sep 2026 12:00:00 +0000",
            "https://example.invalid/cgi-bin/site.xml?one=1&two=2",
        )
    )
    answers["Property_Info"] = [PROPERTY_HEADING]
    assert sum(map(len, answers.values())) == 10
    return answers


def _document(answer: str | tuple[str, ...] | None) -> str:
    root = ET.Element("BESAPI")
    query = ET.SubElement(root, "Query")
    if answer is None:
        ET.SubElement(query, "Error").text = "Synthetic query failure"
    else:
        result = ET.SubElement(query, "Result")
        parent = ET.SubElement(result, "Tuple") if isinstance(answer, tuple) else result
        for value in answer if isinstance(answer, tuple) else (answer,):
            ET.SubElement(parent, "Answer", type="string").text = value
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        + ET.tostring(root, encoding="unicode")
        + "\n"
    )


def _write_capture(
    folder: Path, mode: str, grouped: bool, *, repaired: bool, empty: bool = False
) -> None:
    responses = _responses(empty)
    for stem, context in CONTEXT.items():
        parent = folder / stem.removesuffix("_Info") if grouped else folder
        parent.mkdir(parents=True, exist_ok=True)
        changed = responses.get(stem, [])
        errors = [
            f"ERROR: Query line {100 + index} failed: Synthetic query failure"
            for index in range(len(changed))
        ]
        xml_context, log_context = context, context
        if mode == "mixed":
            # ActionSite size is log-only in current captures. Other groups also
            # exercise fallback to a companion log without duplicating a field.
            xml_context, log_context = context[1::2], context[::2]
        name = parent / f"{stem}_TESTROOT_09-29-2026_1200"
        if mode != "log":
            xml = "".join(_document(answer) for answer in xml_context)
            xml += (
                "".join(_document(answer) for answer in changed)
                if repaired
                else "".join(_document(None) + error + "\n" for error in errors)
            )
            name.with_suffix(".xml").write_text(xml, encoding="utf-8")
        if mode != "xml":
            lines = list(log_context)
            if mode == "log":
                lines += (
                    [", ".join(a) if isinstance(a, tuple) else a for a in changed]
                    if repaired
                    else errors
                )
            name.with_suffix(".log").write_text("\n".join(lines), encoding="utf-8")


def _expected(folder: Path) -> loop.ActionSitePressure:
    return loop.ActionSitePressure(
        deployment_info_folder=folder,
        actionsite_size_kb=19046,
        master_operator_count=6,
        master_open_actions=42,
        external_site_count=57,
        custom_site_count=23,
        automatic_computer_group_count=6,
        manual_computer_group_count=5,
        fixlets_relevant_count=5701,
        fixlets_total_count=34130,
        tasks_relevant_count=1008,
        tasks_total_count=2144,
        analyses_relevant_count=121,
        analyses_total_count=6702,
        active_analyses_count=140,
        properties_count=10342,
        active_properties_count=877,
        custom_properties_count=34,
        master_operator_fixlets=15,
        master_operator_tasks=30,
        master_operator_fixlets_tasks_total=45,
        master_operator_analyses=10,
    )


@pytest.mark.parametrize("mode", ["xml", "log", "mixed"])
@pytest.mark.parametrize("grouped", [False, True], ids=["flat", "grouped"])
@pytest.mark.parametrize("empty", [False, True], ids=["populated", "zero-fallback"])
def test_repaired_responses_preserve_actionsite_context(
    tmp_path: Path, mode: str, grouped: bool, empty: bool
) -> None:
    folder = tmp_path / "Deployment_Information"
    _write_capture(folder, mode, grouped, repaired=False)
    before = loop.parse_actionsite_pressure(folder)
    assert before == _expected(folder)
    _write_capture(folder, mode, grouped, repaired=True, empty=empty)

    # Prove all ten repaired answers reached the sources, rather than comparing
    # two accidentally empty or identical inputs.
    text = "\n".join(
        loop.read_capture_or_log_text(path)
        for path in folder.rglob("*")
        if path.is_file()
    )
    assert "ERROR:" not in text
    for answers in _responses(empty).values():
        for answer in answers:
            for value in answer if isinstance(answer, tuple) else (answer,):
                assert value in text

    after = loop.parse_actionsite_pressure(folder)
    assert after == before
    assert loop.actionsite_pressure_rows(after) == loop.actionsite_pressure_rows(before)
    assert loop.master_operator_custom_content_total(after) == 55


def _workbook_values(path: Path) -> dict[str, list[tuple]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return {
            sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook
        }
    finally:
        workbook.close()


@pytest.mark.parametrize("mode", ["xml", "log", "mixed"])
def test_repaired_context_preserves_loop_workbook_and_word(
    tmp_path: Path, mode: str
) -> None:
    source = tmp_path / "Client_Loop.csv"
    source.write_text(
        "Computer Name,Maximum Cycle Minutes,Average Cycle Minutes\n"
        "TEST-CLIENT-1,42,17\nTEST-CLIENT-2,180,95\n"
        "TEST-CLIENT-3,Not Reported,Not Reported\n",
        encoding="utf-8",
    )
    frame = loop.load_loop_input(source)
    rows, missing, errors = loop.process_rows(frame, *loop.auto_detect_columns(frame))
    metrics, fix_first = loop.calculate_metrics(
        rows,
        len(frame),
        missing,
        errors,
        1800,
        3600,
        date(2026, 9, 28),
        date(2026, 9, 29),
        datetime(2026, 9, 29, 12),
    )
    assert metrics.total_devices == 2
    assert metrics.not_reported_count == 1
    folder = tmp_path / "Deployment_Information"
    snapshots = []
    for repaired in (False, True):
        _write_capture(folder, mode, grouped=True, repaired=repaired)
        assert loop.discover_deployment_info_folder(source) == folder.resolve()
        pressure = loop.parse_actionsite_pressure(folder)
        assert pressure == _expected(folder)
        workbook = tmp_path / f"loop-{repaired}.xlsx"
        document = tmp_path / f"loop-{repaired}.docx"
        loop.generate_excel(
            metrics, rows, fix_first, workbook, actionsite_pressure=pressure
        )
        # The generator's supported no-chart path avoids any browser requirement;
        # the actual Word tables and paragraphs are still written and read back.
        loop.generate_word(
            metrics, tmp_path / "not-rendered.png", document, actionsite_pressure=pressure
        )
        values = _workbook_values(workbook)
        assert values["Contributing Factors"][1:] == [
            tuple(row) for row in loop.actionsite_pressure_rows(pressure)
        ]
        doc = Document(document)
        tables = [
            [tuple(cell.text for cell in row.cells) for row in t.rows] for t in doc.tables
        ]
        assert any(row[0] == "ActionSite size" for table in tables for row in table)
        snapshots.append((values, [p.text for p in doc.paragraphs], tables))
    assert snapshots[0] == snapshots[1]
