# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
"""Content-preserving boundaries for the remaining Design D-item repairs."""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from Lib.Design.report_finishing import finish_report_body
from Lib.Design.report_tables import normalize_docx_tables


def test_cover_metadata_moves_clear_without_changing_values():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    values = [["Prepared for", "Report version"], ["Customer Name", "1.0"]]
    for row, words in zip(table.rows, values, strict=True):
        for cell, value in zip(row.cells, words, strict=True):
            cell.text = value
    finish_report_body(document.element.body)
    assert [[cell.text for cell in row.cells] for row in table.rows] == values
    assert [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid] == [2000, 2400]
    assert table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")) == "4400"


def test_method_subpoints_retain_every_word_and_end_their_keep_chain():
    document = Document()
    parent = document.add_paragraph(
        "Identify expensive content. There are two ways. (1) Use the profiler."
    )
    second = document.add_paragraph("(2) Use the performance analysis.")
    second.runs[0].bold = True
    document.add_paragraph("NOTE Existing source guidance.")
    before = " ".join(p.text for p in document.paragraphs)
    finish_report_body(document.element.body)
    assert " ".join(p.text for p in document.paragraphs) == before
    assert document.paragraphs[1].text.startswith("(1) ")
    assert document.paragraphs[2].text.startswith("(2) ")
    assert document.paragraphs[1].paragraph_format.left_indent.pt == 28
    assert document.paragraphs[2].paragraph_format.left_indent.pt == 28
    assert all(run.bold is False for run in document.paragraphs[2].runs)
    assert parent.runs[0].bold
    assert document.paragraphs[1].paragraph_format.keep_with_next
    assert document.paragraphs[2].paragraph_format.keep_with_next is False
    first = etree.tostring(document.element.body)
    finish_report_body(document.element.body)
    assert etree.tostring(document.element.body) == first


def test_existing_callout_icons_use_plain_portable_symbols_only():
    document = Document()
    for icon, label in (("\U0001f4dd", "NOTE"), ("\U0001f4a1", "TIP")):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(icon + " ")
        run.font.name = "Segoe UI Emoji"
        paragraph.add_run(label + " Existing words.")
    finish_report_body(document.element.body)
    assert [p.text for p in document.paragraphs] == [
        "[i] NOTE Existing words.",
        "[?] TIP Existing words.",
    ]
    assert all(p.runs[0].font.name == "Aptos" for p in document.paragraphs)
    assert not document.element.body.xpath('.//w:rFonts[@w:ascii="Segoe UI Emoji"]')


def test_release_mini_compaction_retains_type_size_and_wording():
    document = Document()
    document.add_heading("Release Intelligence Mini Report", level=1)
    heading = document.add_heading("CVE Security Fixes", level=2)
    document.add_paragraph("No usable CVE data rows were provided.")
    size = heading.style.font.size
    finish_report_body(document.element.body)
    assert heading.paragraph_format.space_before.pt == 8
    assert heading.style.font.size == size
    assert document.paragraphs[-1].text == "No usable CVE data rows were provided."


def test_final_hierarchy_pass_retains_methodology_spacing(tmp_path):
    from Lib.Design.report_layout import finalize_report_document

    document = Document()
    document.add_heading("Methodology", level=1)
    document.add_heading("Data Collection", level=2)
    document.add_paragraph("Existing methodology words.")
    path = tmp_path / "full.docx"
    document.save(path)
    finalize_report_document(path)
    checked = Document(path)
    assert checked.paragraphs[1].paragraph_format.space_before.pt == 4
    assert checked.paragraphs[2].text == "Existing methodology words."


def test_capacity_comparison_aligns_both_value_columns_the_same(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=5)
    table._tbl.tblPr.insert(0, OxmlElement("w:tblStyle"))
    table._tbl.tblPr[0].set(qn("w:val"), "BigFixContent")
    for cell, value in zip(
        table.rows[0].cells,
        ["Resource", "Required", "Observed", "Short by", "Basis"],
        strict=True,
    ):
        cell.text = value
    for values in (
        ["CPU cores", "2 cores", "4 cores", "Meets", "Captured"],
        ["Storage", "20 GB", "100 GB across C:", "Meets", "Captured"],
    ):
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = value
    path = tmp_path / "comparison.docx"
    document.save(path)
    normalize_docx_tables(path)
    result = Document(path)
    for row in result.tables[0].rows:
        assert (
            row.cells[1].paragraphs[0].alignment == row.cells[2].paragraphs[0].alignment
        )
