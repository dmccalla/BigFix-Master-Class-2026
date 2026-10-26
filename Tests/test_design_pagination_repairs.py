# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
"""Regression boundaries for the scoped Design pagination repairs."""

import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from Lib.Design.report_pagination import apply_report_pagination
from Lib.Design.report_tables import normalize_docx_tables


def table(document, headers, count):
    result = document.add_table(rows=1, cols=len(headers))
    style = OxmlElement("w:tblStyle")
    style.set(qn("w:val"), "BigFixContent")
    result._tbl.tblPr.insert(0, style)
    for cell, value in zip(result.rows[0].cells, headers, strict=True):
        cell.text = value
    for index in range(count):
        for cell in result.add_row().cells:
            cell.text = f"Captured value {index}"
    return result


def test_technical_evidence_has_one_boundary_and_retains_all_words():
    document = Document()
    document.add_paragraph("Operational Maturity was not assessed.")
    document.add_page_break()
    label = document.add_paragraph("SECTION 08")
    heading = document.add_heading("Technical Evidence", level=1)
    heading.paragraph_format.page_break_before = True
    apply_report_pagination(document)
    assert [p.text for p in document.paragraphs] == [
        "Operational Maturity was not assessed.",
        "SECTION 08",
        "Technical Evidence",
    ]
    assert label.paragraph_format.page_break_before
    assert heading.paragraph_format.page_break_before is False
    before = etree.tostring(document.element)
    apply_report_pagination(document)
    assert etree.tostring(document.element) == before


def test_reading_key_keeps_only_a_short_opening_group():
    document = Document()
    result = table(document, ["Infographic tile", "Value", "What it means"], 8)
    words = [cell.text for row in result.rows for cell in row.cells]
    apply_report_pagination(document)
    keeps = [
        row.cells[0].paragraphs[0].paragraph_format.keep_with_next
        for row in result.rows
    ]
    assert keeps[:3] == [True, True, False]
    assert all(value is None for value in keeps[3:])
    assert words == [cell.text for row in result.rows for cell in row.cells]


def test_findings_remove_only_empty_trailing_row_without_chaining_body():
    document = Document()
    result = table(
        document, ["Source ID", "Description", "Impact", "Suggested Remediation"], 6
    )
    result.add_row()
    apply_report_pagination(document)
    assert len(result.rows) == 7
    assert all(
        p.paragraph_format.keep_with_next is None
        for row in result.rows
        for cell in row.cells
        for p in cell.paragraphs
    )


def test_nested_query_alignment_changes_without_changing_the_query():
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    paragraph = cell.paragraphs[0]
    query = '/* Property Name */ names it of "Single Property"'
    paragraph.text = query
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    apply_report_pagination(document)
    assert paragraph.text == query
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_status_icon_does_not_prevent_observation_grouping():
    document = Document()
    detail = document.add_paragraph("Detail: Captured disk latency.")
    standard = document.add_paragraph("Standard: Existing threshold.")
    observation = document.add_paragraph("! Observation: Captured result.")
    apply_report_pagination(document)
    assert detail.paragraph_format.keep_with_next
    assert standard.paragraph_format.keep_with_next
    assert observation.paragraph_format.keep_together
    assert observation.paragraph_format.keep_with_next is False


def test_saved_mini_report_finalization_keeps_each_capacity_host_together(tmp_path):
    document = Document()
    label = document.add_paragraph("Captured host")
    table(document, ["Resource", "Required", "Observed", "Short by", "Basis"], 3)
    output = tmp_path / "mini.docx"
    document.save(output)
    normalize_docx_tables(output)
    checked = Document(output)
    assert checked.paragraphs[0].text == label.text
    assert checked.paragraphs[0].paragraph_format.keep_with_next
    flags = [
        row.cells[0].paragraphs[0].paragraph_format.keep_with_next
        for row in checked.tables[0].rows
    ]
    assert flags == [True, True, True, False]
    with zipfile.ZipFile(output) as package:
        first = package.read("word/document.xml")
    normalize_docx_tables(output)
    with zipfile.ZipFile(output) as package:
        assert package.read("word/document.xml") == first


def test_required_empty_paragraph_after_table_cannot_fill_a_new_page():
    document = Document()
    document.add_table(rows=2, cols=2)
    last = document.add_paragraph()
    apply_report_pagination(document)
    assert last.paragraph_format.line_spacing.pt == 1
    assert last.paragraph_format.space_after.pt == 0
    assert not last.text
