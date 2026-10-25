# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""Focused tests for the document stamp (Document_Stamp_Handoff_2026-09-07).

One helper formats the stamp line for every run document. These tests pin the
exact string, the blank-value dashes, the Word placement rule (one body
paragraph directly under the first heading), and the Excel placement rule
(row 1 cell A1 of the first sheet with everything else shifted down one row,
freeze panes, filters, merges, and row heights moving with the content).
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from Lib.Core.document_stamp import (  # noqa: E402
    add_stamp_arguments,
    document_stamp,
    is_stamp_text,
    stamp_docx_document,
    stamp_from_args,
    stamp_worksheet,
    strip_stamp_paragraphs_xml,
)

EXPECTED_FULL = "Customer: BigFix Inc · Assessment dates: 07-01-2026 to 07-04-2026"


# ---------------------------------------------------------------------------
# The stamp string
# ---------------------------------------------------------------------------


def test_stamp_exact_string():
    assert document_stamp("BigFix Inc", "07/01/2026", "07/04/2026") == EXPECTED_FULL


def test_stamp_separator_is_middle_dot_with_spaces():
    stamp = document_stamp("BigFix Inc", "07/01/2026", "07/04/2026")
    assert " · " in stamp
    assert chr(8212) not in stamp


def test_stamp_preserves_dash_dates():
    assert document_stamp("BigFix Inc", "07-01-2026", "07-04-2026") == EXPECTED_FULL


def test_stamp_normalizes_iso_dates():
    assert document_stamp("BigFix Inc", "2026-07-01", "2026-07-04") == EXPECTED_FULL


def test_stamp_blank_customer_prints_dash():
    stamp = document_stamp("", "07/01/2026", "07/04/2026")
    assert stamp.startswith("Customer: - · ")


def test_stamp_blank_dates_print_dash():
    assert document_stamp("BigFix Inc") == "Customer: BigFix Inc · Assessment dates: -"


def test_stamp_one_missing_date_prints_dash():
    stamp = document_stamp("BigFix Inc", "07/01/2026", "")
    assert stamp.endswith("Assessment dates: -")


def test_stamp_all_blank_renders_no_stamp():
    assert document_stamp("") == ""


def test_stamp_none_values_treated_as_blank():
    assert document_stamp(None, None, None) == ""


# ---------------------------------------------------------------------------
# The argparse plumbing
# ---------------------------------------------------------------------------


def test_add_stamp_arguments_round_trip():
    parser = argparse.ArgumentParser()
    add_stamp_arguments(parser)
    args = parser.parse_args(
        [
            "--customer",
            "BigFix Inc",
            "--assessment-start-date",
            "07/01/2026",
            "--assessment-end-date",
            "07/04/2026",
        ]
    )
    assert stamp_from_args(args) == EXPECTED_FULL


def test_stamp_arguments_default_blank():
    parser = argparse.ArgumentParser()
    add_stamp_arguments(parser)
    args = parser.parse_args([])
    assert stamp_from_args(args) == ""


# ---------------------------------------------------------------------------
# Word placement
# ---------------------------------------------------------------------------


def _docx_document():
    from docx import Document

    return Document()


def test_docx_stamp_lands_directly_under_first_heading():
    document = _docx_document()
    document.add_heading("Title Of Document", level=1)
    document.add_paragraph("First body paragraph.")
    stamp = document_stamp("BigFix Inc", "07/01/2026", "07/04/2026")
    stamp_docx_document(document, stamp)
    texts = [p.text for p in document.paragraphs]
    assert texts.index(stamp) == texts.index("Title Of Document") + 1


def test_docx_stamp_uses_body_style_without_direct_formatting():
    document = _docx_document()
    document.add_heading("Title Of Document", level=1)
    document.add_paragraph("Body.")
    stamp = document_stamp("BigFix Inc", "07/01/2026", "07/04/2026")
    stamp_docx_document(document, stamp)
    paragraph = next(p for p in document.paragraphs if p.text == stamp)
    assert paragraph.style.name == "Normal"
    for run in paragraph.runs:
        assert run.bold is None
        assert run.font.color.rgb is None


def test_docx_stamp_becomes_first_paragraph_when_no_heading():
    document = _docx_document()
    document.add_paragraph("Only body text.")
    stamp = document_stamp("BigFix Inc", "07/01/2026", "07/04/2026")
    stamp_docx_document(document, stamp)
    assert document.paragraphs[0].text == stamp


def test_docx_stamp_after_title_style_heading():
    document = _docx_document()
    document.add_heading("Mini Report", 0)
    document.add_paragraph("Intro.")
    stamp = document_stamp("BigFix Inc", "07/01/2026", "07/04/2026")
    stamp_docx_document(document, stamp)
    texts = [p.text for p in document.paragraphs]
    assert texts.index(stamp) == texts.index("Mini Report") + 1


# ---------------------------------------------------------------------------
# Excel placement
# ---------------------------------------------------------------------------


def _sample_worksheet():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Header A", "Header B", "Header C"])
    ws.append(["a1", "b1", "c1"])
    ws.append(["a2", "b2", "c2"])
    ws.merge_cells("A1:B1")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:C3"
    ws.column_dimensions["A"].width = 42
    ws.row_dimensions[1].height = 30
    return wb, ws


def test_worksheet_stamp_in_a1_and_content_shifted():
    wb, ws = _sample_worksheet()
    stamp = document_stamp("BigFix Inc", "07/01/2026", "07/04/2026")
    stamp_worksheet(ws, stamp)
    assert ws["A1"].value == stamp
    assert ws["A2"].value == "Header A"
    assert ws["A3"].value == "a1"


def test_worksheet_stamp_moves_freeze_filter_merge_and_heights():
    wb, ws = _sample_worksheet()
    stamp_worksheet(ws, document_stamp("BigFix Inc", "07/01/2026", "07/04/2026"))
    assert ws.freeze_panes == "A3"
    assert ws.auto_filter.ref == "A2:C4"
    assert "A2:B2" in {str(r) for r in ws.merged_cells.ranges}
    assert ws.row_dimensions[2].height == 30
    assert 1 not in ws.row_dimensions or ws.row_dimensions[1].height is None


def test_worksheet_stamp_keeps_column_widths():
    wb, ws = _sample_worksheet()
    stamp_worksheet(ws, document_stamp("BigFix Inc", "07/01/2026", "07/04/2026"))
    assert ws.column_dimensions["A"].width == 42


def test_worksheet_stamp_moves_data_validation():
    from openpyxl.worksheet.datavalidation import DataValidation

    wb, ws = _sample_worksheet()
    validation = DataValidation(type="list", formula1='"Yes,No"')
    validation.add("B2:B3")
    ws.add_data_validation(validation)
    stamp_worksheet(ws, document_stamp("BigFix Inc", "07/01/2026", "07/04/2026"))
    refs = {
        str(ref) for dv in ws.data_validations.dataValidation for ref in dv.sqref.ranges
    }
    assert refs == {"B3:B4"}


# ---------------------------------------------------------------------------
# Injection stripping: the stamp never rides into the assembled document
# ---------------------------------------------------------------------------


def _para_xml(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def test_is_stamp_text_matches_only_the_stamp_shape():
    assert is_stamp_text(document_stamp("BigFix Inc", "07/01/2026", "07/04/2026"))
    assert is_stamp_text("Customer: - " + chr(183) + " Assessment dates: -")
    assert not is_stamp_text(
        "Customer: BigFix Inc.  |  Data Collected: Sun  |  Generated: now"
    )
    assert not is_stamp_text("Body text about a customer.")


def test_strip_stamp_paragraphs_xml_removes_only_the_stamp():
    stamp = document_stamp("BigFix Inc", "07/01/2026", "07/04/2026")
    body = (
        _para_xml("Module Title")
        + _para_xml(stamp)
        + _para_xml("Standalone module excerpt. Keep the rest.")
        + _para_xml("Body paragraph.")
    )
    stripped = strip_stamp_paragraphs_xml(body)
    assert "Assessment dates:" not in stripped
    assert "Module Title" in stripped
    assert "Standalone module excerpt" in stripped
    assert "Body paragraph." in stripped


def test_stub_front_matter_strip_handles_title_stamp_notice_order():
    from Lib.Themes.word_stub_utils import (
        STUB_NOTICE_TEXT,
        strip_stub_front_matter_xml,
    )

    stamp = document_stamp("BigFix Inc", "07/01/2026", "07/04/2026")
    body = (
        _para_xml("FillDB Performance Analysis")
        + _para_xml(stamp)
        + _para_xml(STUB_NOTICE_TEXT)
        + _para_xml("Body paragraph.")
    )
    stripped = strip_stub_front_matter_xml(body, titles=("FillDB Performance Analysis",))
    assert "Assessment dates:" not in stripped
    assert "Standalone module excerpt" not in stripped
    assert "FillDB Performance Analysis" not in stripped
    assert "Body paragraph." in stripped
