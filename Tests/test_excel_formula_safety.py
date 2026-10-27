# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.


"""Workbook cells built from collected data must never become formulas.

A computer name or property value that starts with = + - or @ is data from the
customer's deployment. These tests keep it stored as text.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd
import xlsxwriter

from Lib.Themes.workbook_theme import (
    append_openpyxl_row,
    finish_openpyxl_workbook,
    xlsxwriter_options,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIRS = ("App", "Build", "Get", "Inject", "Lib", "Utilities")

FORMULA_LIKE = [
    '=WEBSERVICE("https://example.invalid")',
    "+1+1",
    "-1+1",
    "@SUM(1+1)",
    "\tvalue",
]


def test_openpyxl_rows_keep_formula_like_strings_as_text(tmp_path: Path) -> None:
    output_path = tmp_path / "formula-like-values.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    append_openpyxl_row(sheet, FORMULA_LIKE)
    workbook.save(output_path)

    loaded = openpyxl.load_workbook(output_path, data_only=False)
    cells = list(loaded.active[1])
    assert [cell.value for cell in cells] == FORMULA_LIKE
    assert all(cell.data_type == "s" for cell in cells)


def test_xlsxwriter_options_keep_formula_like_strings_as_text(tmp_path: Path) -> None:
    output_path = tmp_path / "xlsxwriter-formula-like-value.xlsx"
    workbook = xlsxwriter.Workbook(str(output_path), xlsxwriter_options())
    workbook.add_worksheet().write(0, 0, FORMULA_LIKE[0])
    workbook.close()

    loaded = openpyxl.load_workbook(output_path, data_only=False)
    cell = loaded.active["A1"]
    assert cell.value == FORMULA_LIKE[0]
    assert cell.data_type == "s"


def test_xlsxwriter_options_keep_caller_options() -> None:
    options = xlsxwriter_options({"nan_inf_to_errors": True})
    assert options == {"strings_to_formulas": False, "nan_inf_to_errors": True}


def test_loop_workbook_path_keeps_formula_like_strings_as_text(tmp_path: Path) -> None:
    """The pandas writer used for the loop workbook, then the post-write restyle."""
    output_path = tmp_path / "loop.xlsx"
    frame = pd.DataFrame({"ComputerName": FORMULA_LIKE, "Result": FORMULA_LIKE})

    with pd.ExcelWriter(
        output_path,
        engine="xlsxwriter",
        engine_kwargs={"options": xlsxwriter_options()},
    ) as writer:
        frame.to_excel(writer, sheet_name="Source Data", index=False)
    finish_openpyxl_workbook(output_path)

    loaded = openpyxl.load_workbook(output_path, data_only=False)
    rows = list(loaded["Source Data"].iter_rows(min_row=2))
    assert [row[0].value for row in rows] == FORMULA_LIKE
    assert all(cell.data_type != "f" for row in rows for cell in row)


def _shipped_sources():
    for directory in SOURCE_DIRS:
        base = ROOT / directory
        if base.is_dir():
            yield from base.rglob("*.py")


def test_workbook_writers_use_the_safe_options() -> None:
    """Every xlsxwriter workbook and pandas Excel writer in shipped code is safe."""
    offenders: list[str] = []
    for path in _shipped_sources():
        text = path.read_text(encoding="utf-8", errors="replace")
        for match in re.finditer(r"xlsxwriter\.Workbook\(|ExcelWriter\(", text):
            window = text[match.start() : match.start() + 250]
            if "xlsxwriter_options" not in window:
                line = text.count("\n", 0, match.start()) + 1
                offenders.append(f"{path.relative_to(ROOT)}:{line}")
    assert not offenders, "Workbook writers must pass xlsxwriter_options(): " + ", ".join(
        offenders
    )
