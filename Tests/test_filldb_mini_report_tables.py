# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""The FillDB Mini Report keeps objects and capture statistics in their own tables."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "Get" / "get_filldb_performance.py"

CSV = """FillDB Object,Count,Time (ms),Rate/s,Health Check

Fixlet results:,1021932,439062,2328,093%: Good

Action results:,88953,49436,1799,360%: Great

Parallel DB Update:,0,0,0,----: Not enough data.


Statistic,Value

Lines processed:,4237218

Duration (s):,68.54

Throughput (lines/s):,61821.1

FillDB start time:,2025-11-16 02:56:11-05:00

FillDB end time:,2025-11-17 15:57:07-05:00

FillDB duration (hh:mm:ss):,"1 day, 13:00:56"

FillDB parsing threads:,3

FillDB database threads:,3
"""


def _build(tmp_path: Path) -> list[list[list[str]]]:
    source = tmp_path / "FillDBPerf_analysis.csv"
    source.write_text(CSV, encoding="utf-8")
    output = tmp_path / "filldb.docx"
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--input",
            str(source),
            "--word-output",
            str(output),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        stdin=subprocess.DEVNULL,
        timeout=300,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    document = Document(str(output))
    return [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]


def test_object_table_holds_only_objects(tmp_path: Path) -> None:
    objects, _summary = _build(tmp_path)

    names = [row[0] for row in objects[1:]]
    assert names == ["Fixlet Results", "Action Results", "Parallel DB Update"]
    assert "Statistic" not in names


def test_object_counts_have_thousands_separators(tmp_path: Path) -> None:
    objects, _summary = _build(tmp_path)

    assert objects[1][:4] == ["Fixlet Results", "1,021,932", "439,062", "2,328"]
    assert objects[3][1:4] == ["0", "0", "0"]
    # The analyzer's own Health Check text is left as the analyzer wrote it.
    assert objects[1][4] == "093%: Good"


def test_capture_summary_holds_the_analyzer_statistics(tmp_path: Path) -> None:
    _objects, summary = _build(tmp_path)

    rows = {row[0]: row[1] for row in summary[1:]}
    assert rows["Lines Processed"] == "4,237,218"
    assert rows["Duration (s)"] == "68.54"
    assert rows["Throughput (lines/s)"] == "61,821.1"
    assert rows["FillDB Duration (hh:mm:ss)"] == "1 day, 13:00:56"
    assert rows["FillDB Parsing Threads"] == "3"
    assert rows["FillDB Database Threads"] == "3"


def test_score_text_names_no_summary_scorecard(tmp_path: Path) -> None:
    _build(tmp_path)
    paragraphs = [p.text for p in Document(str(tmp_path / "filldb.docx")).paragraphs]

    assert "This score covers the FillDB capture only." in paragraphs
    assert not any("Summary Scorecard" in text for text in paragraphs)
