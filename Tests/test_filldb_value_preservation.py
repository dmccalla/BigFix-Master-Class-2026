# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""Protect analyzer evidence from precision loss throughout FillDB reporting."""

from pathlib import Path

import get_filldb_performance as filldb
from docx import Document

from Lib.Design import html_charts


def test_analyzer_values_survive_csv_word_chart_and_narrative(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "precision.csv"
    source.write_text(
        "FillDB Object,Count,Time (ms),Rate/s,Health Check\n"
        "Parallel DB Update (Short Batch):,9007199254740993,164.000,0.123456789012345678900,093%: Good\n"
        "Parallel Parsing:,00007,0.00000000000000012300,1.2300e-10,----: Not enough data.\n"
        "Statistic,Value\n"
        "Lines processed:,9007199254740993\n"
        "Duration (s):,0.0012300\n"
        "Throughput (lines/s):,123456789.012345678900\n"
        "FillDB start time:,2026-10-27 17:31:33.123400-04:00\n"
        "FillDB end time:,2026-10-27 17:32:41.123400-04:00\n"
        "FillDB duration (hh:mm:ss):,0:01:08.000000\n",
        encoding="utf-8",
    )
    rows = filldb.parse_filldb_csv(source)
    metadata = filldb.build_capture_metadata(
        filldb.find_stat_text(rows, "FillDB start time"),
        filldb.find_stat_text(rows, "FillDB end time"),
        filldb.find_stat_text(rows, "FillDB duration (hh:mm:ss)"),
    )
    output = tmp_path / "precision.docx"
    filldb.write_word_output(rows, output, metadata=metadata)
    document = Document(output)
    for source_row, word_row in zip(rows[:2], document.tables[0].rows[1:]):
        assert [c.text for c in word_row.cells[1:]] == [
            source_row[key] for key in filldb.CSV_COLUMNS[1:]
        ]
    summary = {
        row.cells[1].text: row.cells[2].text for row in document.tables[1].rows[1:]
    }
    for row in rows[3:]:
        label = row["FillDB Object"].removesuffix(":")
        displayed_label = next(
            key for key in summary if key.casefold() == label.casefold()
        )
        assert summary[displayed_label] == row["Count"]
    assert summary["Collection Start"] == "2026-10-27 17:31:33.123400-04:00"
    assert summary["Capture Duration"] == "0:01:08.000000"
    text = " ".join(p.text for p in document.paragraphs)
    assert "at 164.000 ms." in text
    assert "(00007 rows)" in text
    assert "processed 9007199254740993 lines" in text
    captured = {}
    monkeypatch.setattr(
        html_charts, "render_chart", lambda _kind, data, _path: captured.update(data)
    )
    filldb.generate_filldb_chart(rows, tmp_path / "not-written.png", metadata)
    assert [b["mins"] for b in captured["bars"]] == [
        "164.000 ms",
        "0.00000000000000012300 ms",
    ]
    assert "9007199254740993 lines processed" in captured["meta"]
    assert "capture 0:01:08.000000" in captured["meta"]


def test_missing_or_zero_line_stat_never_uses_object_counts(
    tmp_path: Path, monkeypatch
) -> None:
    objects = [
        {
            "FillDB Object": "Parallel Parsing:",
            "Count": "700",
            "Time (ms)": "164",
            "Rate/s": "1",
            "Health Check": "100%: Good",
        }
    ]
    captured = {}
    monkeypatch.setattr(
        html_charts, "render_chart", lambda _kind, data, _path: captured.update(data)
    )
    for statistics, expected in [
        ([], "Lines processed: Not available"),
        ([{"FillDB Object": "Lines processed:", "Count": "0"}], "0 lines processed"),
    ]:
        filldb.generate_filldb_chart(
            objects + statistics, tmp_path / "not-written.png", {}
        )
        assert captured["meta"] == expected
