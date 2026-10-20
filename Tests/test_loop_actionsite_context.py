# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

from pathlib import Path

import get_loop_health as loop
import pandas as pd
from docx import Document
from docx.oxml.ns import qn

PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR"
    b"\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde"
    b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfeA\xe2%\xb5"
    b"\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _write_png(path: Path) -> Path:
    path.write_bytes(PNG_1X1)
    return path


def _sample_loop_metrics() -> loop.ReportMetrics:
    return loop.ReportMetrics(
        total_devices=10,
        score_total_devices=10,
        overall_status="Healthy",
        environment_score=95,
        score_driver="test driver",
        score_breakdown="test breakdown",
        report_start_text="7/1/2026",
        report_end_text="7/2/2026",
        report_span_text="2 days",
        generated_at_text="7/2/2026 17:32",
        avg_mean_seconds=300,
        avg_median_seconds=240,
        avg_max_seconds=600,
        slowest_avg_seconds=900,
        peak_max_seconds=1200,
        avg_healthy=8,
        avg_warning=2,
        avg_critical=0,
        avg_healthy_pct=0.8,
        avg_warning_pct=0.2,
        avg_critical_pct=0.0,
        avg_over_120=0,
        avg_over_120_pct=0.0,
        not_reported_count=0,
        error_count=0,
        not_reported_pct=0.0,
        error_pct=0.0,
        total_issue_count=0,
        total_issue_pct=0.0,
        healthy_threshold_seconds=1800,
        critical_threshold_seconds=3600,
        fix_first_count=0,
    )


def test_loop_chart_names_the_source_id_for_its_raised_finding(
    monkeypatch, tmp_path: Path
) -> None:
    metrics = _sample_loop_metrics()
    metrics.overall_status = "WARNING"
    metrics.environment_score = 79
    captured: dict = {}

    def capture_chart(kind, data, output):
        captured.update({"kind": kind, "data": data, "output": output})

    monkeypatch.setattr("Lib.Design.html_charts.render_chart", capture_chart)
    output = tmp_path / "loop.png"

    loop.generate_png(metrics, output)

    assert captured["kind"] == "loop"
    assert captured["output"] == output
    assert captured["data"]["findings"] == [
        {
            "id": "CLT-01",
            "text": "High. The client loop health score is below 90.",
        }
    ]


def test_loop_word_explains_the_score_under_loop_health(
    tmp_path: Path,
) -> None:
    word_output = tmp_path / "loop.docx"
    png_output = _write_png(tmp_path / "loop.png")

    loop.generate_word(_sample_loop_metrics(), png_output, word_output)

    doc = Document(str(word_output))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    score_heading = next(
        p for p in doc.paragraphs if p.text.strip() == "How the Score Is Calculated"
    )
    assert score_heading.style.name == "Heading 4"
    assert "Client Loop Context" not in paragraphs
    assert paragraphs.index("Loop Health") < paragraphs.index(
        "How the Score Is Calculated"
    )
    assert "Score Method: test driver" not in paragraphs
    assert "Score Breakdown: test breakdown" not in paragraphs
    assert (
        "Each device earns points for its group. Healthy earns 100, Good 85, "
        "Warning 67 and Critical 34. The score is the average across all "
        "devices." in paragraphs
    )
    assert not any("Section 04" in text for text in paragraphs)
    assert not any(
        table.cell(0, 0).text.startswith("SLOWEST EVALUATION, AVG")
        for table in doc.tables
    )


def test_loop_mini_report_includes_only_raised_automated_findings(
    tmp_path: Path,
) -> None:
    warning_metrics = _sample_loop_metrics()
    warning_metrics.overall_status = "WARNING"
    word_output = tmp_path / "warning-loop.docx"

    loop.generate_word(
        warning_metrics,
        _write_png(tmp_path / "warning-loop.png"),
        word_output,
        include_automated_findings=True,
    )

    doc = Document(str(word_output))
    headings = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    findings_table = next(
        table for table in doc.tables if table.cell(0, 0).text == "Finding ID"
    )

    assert headings.index("Automated Findings") < headings.index("Loop Health")
    assert [cell.text for cell in findings_table.rows[1].cells] == [
        "CLT-01",
        "High.\nThe client loop health score is below 90. See Loop Health.",
    ]

    healthy_output = tmp_path / "healthy-loop.docx"
    loop.generate_word(
        _sample_loop_metrics(),
        _write_png(tmp_path / "healthy-loop.png"),
        healthy_output,
        include_automated_findings=True,
    )
    healthy_headings = [
        paragraph.text
        for paragraph in Document(str(healthy_output)).paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    assert "Automated Findings" not in healthy_headings


def test_parse_actionsite_pressure_from_deployment_logs(tmp_path: Path) -> None:
    deploy = tmp_path / "Deployment_Information"
    deploy.mkdir()
    (deploy / "General_Info_BFROOT1_06-14-2026_2018.log").write_text(
        "Actionsite Size: 6931 KB\n",
        encoding="utf-8",
    )
    (deploy / "Operator_Info_BFROOT1_06-14-2026_2018.log").write_text(
        "Number of Master Operators: 3\n",
        encoding="utf-8",
    )
    (deploy / "Action_Info_BFROOT1_06-14-2026_2018.log").write_text(
        "Number of open actions issued by master operators: 42\n",
        encoding="utf-8",
    )
    (deploy / "Content_Info_BFROOT1_06-30-2026_1428").write_text(
        "\n".join(
            [
                "Number of External Sites: 57",
                "Number of Custom Sites: 23",
                "Number of Automatic Computer Groups: 6",
                "Number of Manual Computer Groups: 5",
                "Number of Fixlets: 5701 / 34130",
                "Number of Tasks: 1008 / 2144",
                "Number of Analyses: 121 / 6702",
                "Number of Active Analyses: 140",
                "Number of Properties: 10342",
                "Number of Active Properties: 877",
                "Number of Custom Properties: 34",
                "Number of Master Operator Fixlets: 15",
                "Number of Master Operator Tasks: 30",
                "Total Number of Master Operator Fixlets and Tasks: 45",
                "Total Number of Master Operator Analyses: 10",
            ]
        ),
        encoding="utf-8",
    )

    pressure = loop.parse_actionsite_pressure(deploy)

    assert pressure is not None
    assert pressure.actionsite_size_kb == 6931
    assert pressure.master_operator_count == 3
    assert pressure.master_open_actions == 42
    assert pressure.external_site_count == 57
    assert pressure.custom_site_count == 23
    assert pressure.fixlets_relevant_count == 5701
    assert pressure.fixlets_total_count == 34130
    assert pressure.tasks_relevant_count == 1008
    assert pressure.tasks_total_count == 2144
    assert pressure.analyses_relevant_count == 121
    assert pressure.analyses_total_count == 6702
    assert pressure.active_analyses_count == 140
    assert pressure.properties_count == 10342
    assert pressure.active_properties_count == 877
    assert pressure.custom_properties_count == 34
    assert loop.master_operator_custom_content_total(pressure) == 55
    assert loop.actionsite_pressure_rows(pressure) == [
        [
            "ActionSite size",
            "6.8 MB (6,931 KB)",
            "Healthy",
            "Healthy at 10 MB or below. Soft warning over 10 MB. Hard ceiling at 25 MB.",
        ],
        [
            "Subscribed sites",
            "80 total (57 external, 23 custom)",
            "Healthy",
            "Healthy at 100 or fewer. Review when over 100.",
        ],
        [
            "Computer groups",
            "11 total (6 automatic, 5 manual)",
            "Healthy",
            "Healthy at 100 or fewer. Review when over 100.",
        ],
        [
            "Fixlets and tasks",
            "6,709 / 36,274 relevant/total (5,701 / 34,130 fixlets, 1,008 / 2,144 tasks)",
            "Healthy",
            "Healthy at 80,000 total or fewer. Review when over 80,000.",
        ],
        [
            "Analyses",
            "121 / 6,702 relevant/total. 140 active.",
            "Healthy",
            "Healthy at 10,000 total or fewer and 500 active or fewer. Review if either is over threshold.",
        ],
        [
            "Properties",
            "10,342 total (877 active, 34 custom)",
            "Needs Attention",
            "Healthy at 10,000 total or fewer and 500 active or fewer. Review if either is over threshold.",
        ],
        [
            "Master operators",
            "3",
            "Healthy",
            "Healthy at 5 or fewer. Review when over 5.",
        ],
        [
            "Open actions issued by master operators",
            "42",
            "Needs Attention",
            "Healthy at 25 or fewer. Review when over 25.",
        ],
        [
            "Master-operator fixlets/tasks/analyses",
            "55 total (15 fixlets, 30 tasks, 10 analyses)",
            "Needs Attention",
            "Healthy below 50 combined. Review at 50 or more.",
        ],
    ]


def test_parse_actionsite_pressure_reads_log_only_fields_when_xml_exists(
    tmp_path: Path,
) -> None:
    deploy = tmp_path / "Deployment_Information"
    deploy.mkdir()
    (deploy / "General_Info_BFROOT1_08-10-2026_1200.xml").write_text(
        "\n".join(
            [
                '<?xml version="1.0" encoding="UTF-8"?>',
                "<BESAPI>",
                "<Query><Result>",
                '<Answer type="string">Number of Computers: 100</Answer>',
                "</Result></Query>",
                "</BESAPI>",
            ]
        ),
        encoding="utf-8",
    )
    (deploy / "General_Info_BFROOT1_08-10-2026_1200.log").write_text(
        "Actionsite Size: 19046 KB\n",
        encoding="utf-8",
    )

    pressure = loop.parse_actionsite_pressure(deploy)

    assert pressure is not None
    assert pressure.actionsite_size_kb == 19046
    assert loop.actionsite_pressure_rows(pressure)[0] == [
        "ActionSite size",
        "18.6 MB (19,046 KB)",
        "Needs Attention",
        "Healthy at 10 MB or below. Soft warning over 10 MB. Hard ceiling at 25 MB.",
    ]


def test_loop_word_orders_diagnostic_sections_from_design_response(
    tmp_path: Path,
) -> None:
    pressure = loop.ActionSitePressure(
        deployment_info_folder=tmp_path,
        actionsite_size_kb=12_000,
        master_operator_count=6,
    )
    offenders = [
        loop.LoopOffenderSummary(
            rank=1,
            occurrences=15,
            total_client_processing_seconds=630,
            max_client_processing_seconds=0,
            avg_client_processing_seconds=42,
            result="Background Evaluation of Custom Site: Resolve",
        )
    ]
    udp_summary = loop.UdpNotificationSummary(
        source_file=tmp_path / "udp-notification-results.csv",
        total_devices=4,
        receiving_count=1,
        not_receiving_count=2,
        unknown_count=1,
        relay_rows=[["relay-a", "2 / 3", "66.7%"]],
        subnet_rows=[["10.10.1.0/24", "2 / 3", "relay-a"]],
    )
    word_output = tmp_path / "loop.docx"

    loop.generate_word(
        _sample_loop_metrics(),
        tmp_path / "missing.png",
        word_output,
        actionsite_pressure=pressure,
        loop_offenders=offenders,
        udp_summary=udp_summary,
    )

    headings = [
        p.text
        for p in Document(str(word_output)).paragraphs
        if p.style.name.startswith("Heading")
    ]

    assert "ActionSite Pressure" not in headings
    assert "Client Loop Penalties" not in headings
    order = [
        "Reading the Infographic",
        "Loop Health",
        "How the Score Is Calculated",
        "Reporting Issues",
        "Loop Cost",
        "Top 10 Most Expensive Evaluations",
        "Contributing Factors",
        "Notification and Polling",
        "UDP Notification Coverage",
        "UDP Notification Issues by Relay",
        "Top Subnets Not Receiving UDP Pings",
    ]
    assert [h for h in headings if h in order] == order

    doc = Document(str(word_output))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    assert any(
        p.endswith(
            "reported no usable loop data. They count as Critical in the table and "
            "in the score. When more than 5% of devices have no usable data, CLT-01 "
            "gives it as a reason."
        )
        for p in paragraphs
    )
    assert not any("visibility problem" in p for p in paragraphs)

    offender_tables = [
        table
        for table in doc.tables
        if [cell.text for cell in table.rows[0].cells]
        == [
            "Rank",
            "Average Eval Time (Seconds)",
            "Occurrences",
            "Evaluation",
            "Interval",
        ]
    ]
    assert offender_tables
    assert [cell.text for cell in offender_tables[0].rows[1].cells] == [
        "1",
        "42",
        "15",
        "Background Evaluation of Custom Site: Resolve",
        "Not a property",
    ]


def test_context_strip_tiles_feed_chart_data_with_missing_udp_fallback() -> None:
    metrics = _sample_loop_metrics()
    offenders = [
        loop.LoopOffenderSummary(
            rank=1,
            occurrences=15,
            total_client_processing_seconds=630,
            max_client_processing_seconds=0,
            avg_client_processing_seconds=42,
            result="Background Evaluation of Custom Site: Resolve",
        )
    ]

    tiles = loop.context_strip_tiles(metrics, offenders, udp_summary=None)

    assert tiles == [
        {
            "label": "Devices reporting",
            "value": "10 / 10",
            "caption": "All scored devices have usable loop data",
            "muted": False,
        },
        {
            "label": "No UDP pings received",
            "value": "Not collected",
            "caption": "",
            "muted": True,
        },
        {
            "label": "Relay With Most UDP Issues",
            "value": "Not collected",
            "caption": "",
            "muted": True,
        },
    ]


def test_context_strip_tiles_use_total_machines_for_udp_and_relay_count() -> None:
    metrics = _sample_loop_metrics()
    udp_summary = loop.UdpNotificationSummary(
        source_file=Path("udp.csv"),
        total_devices=4,
        receiving_count=1,
        not_receiving_count=2,
        unknown_count=1,
        relay_rows=[["relay-a", "2 / 3", "66.7%"]],
        subnet_rows=[],
    )

    tiles = loop.context_strip_tiles(metrics, udp_summary=udp_summary)

    assert tiles[1]["value"] == "2 / 10"
    assert tiles[1]["caption"] == "Devices that received no UDP pings."
    assert tiles[2]["value"] == "2 / 3"
    assert tiles[2]["caption"] == "relay-a"


def _udp_two_missing() -> loop.UdpNotificationSummary:
    return loop.UdpNotificationSummary(
        source_file=Path("udp.csv"),
        total_devices=4,
        receiving_count=1,
        not_receiving_count=2,
        unknown_count=1,
        relay_rows=[["relay-a", "2 / 3", "66.7%"]],
        subnet_rows=[],
    )


def test_infographic_key_uses_infographic_values() -> None:
    offenders = [
        loop.LoopOffenderSummary(
            rank=1,
            occurrences=15,
            total_client_processing_seconds=630,
            max_client_processing_seconds=0,
            avg_client_processing_seconds=42,
            result="Background Evaluation of Custom Site: Resolve",
        )
    ]

    rows = loop.client_loop_infographic_key_rows(
        _sample_loop_metrics(), offenders, _udp_two_missing()
    )

    assert [row[0] for row in rows] == [
        "Health score",
        "Loop distribution",
        "Average loop time",
        "Slowest evaluation",
        "Severe loops over 120 min",
        "Devices reporting",
        "No UDP pings received",
        "Relay With Most UDP Issues",
        "Findings raised",
    ]
    values = {row[0]: row[1] for row in rows}
    assert values["Health score"] == "95 / 100"
    # Built with the old two thresholds, so there is no Good group.
    assert values["Loop distribution"] == "8 · 2 · 0 · 0"
    assert values["Average loop time"] == "5m 00s"
    assert values["Slowest evaluation"] == loop.format_loop_offender_strip_seconds(42)
    assert values["Severe loops over 120 min"] == "0"
    assert values["Devices reporting"] == "10 / 10"
    assert values["No UDP pings received"] == "2 / 10"
    assert values["Relay With Most UDP Issues"] == "2 / 3 on relay-a"
    assert values["Findings raised"] == "None raised"


def test_infographic_key_names_clt_01_when_loop_health_is_low() -> None:
    metrics = _sample_loop_metrics()
    metrics.overall_status = "WARNING"

    rows = loop.client_loop_infographic_key_rows(metrics)

    assert rows[-1][1] == "CLT-01"
    assert dict((row[0], row[1]) for row in rows)["Slowest evaluation"] == (
        "Not available"
    )
    assert dict((row[0], row[1]) for row in rows)["No UDP pings received"] == (
        "Not collected"
    )


def test_loop_word_places_infographic_key_directly_after_infographic(
    tmp_path: Path,
) -> None:
    word_output = tmp_path / "loop.docx"

    loop.generate_word(
        _sample_loop_metrics(),
        _write_png(tmp_path / "loop.png"),
        word_output,
        udp_summary=_udp_two_missing(),
    )

    doc = Document(str(word_output))
    body_elements = list(doc.element.body)

    def element_text(element) -> str:
        return "".join(node.text or "" for node in element.iter(qn("w:t")))

    infographic_index = next(
        index
        for index, element in enumerate(body_elements)
        if element.find(f".//{qn('w:drawing')}") is not None
    )
    key_heading_index = next(
        index
        for index, element in enumerate(body_elements)
        if element_text(element) == "Reading the Infographic"
    )
    key_table_index = next(
        index
        for index, element in enumerate(body_elements)
        if element_text(element).startswith("Infographic tileValueWhat it means")
    )
    loop_health_index = next(
        index
        for index, element in enumerate(body_elements)
        if element_text(element) == "Loop Health"
    )

    assert infographic_index < key_heading_index < key_table_index < loop_health_index
    assert [cell.text for cell in doc.tables[0].rows[0].cells] == [
        "Infographic tile",
        "Value",
        "What it means",
    ]


def test_actionsite_pressure_rows_use_spine_statuses_for_thresholds(
    tmp_path: Path,
) -> None:
    pressure = loop.ActionSitePressure(
        deployment_info_folder=tmp_path,
        actionsite_size_kb=11 * 1024,
        master_operator_count=6,
        master_open_actions=26,
    )

    rows = loop.actionsite_pressure_rows(pressure)

    assert [row[2] for row in rows] == [
        "Needs Attention",
        "Not Assessed",
        "Not Assessed",
        "Not Assessed",
        "Not Assessed",
        "Not Assessed",
        "Needs Attention",
        "Needs Attention",
        "Not Assessed",
    ]
    assert rows[-1][1] == "Not detected"


def test_actionsite_pressure_rows_flag_content_info_thresholds(tmp_path: Path) -> None:
    pressure = loop.ActionSitePressure(
        deployment_info_folder=tmp_path,
        external_site_count=80,
        custom_site_count=21,
        automatic_computer_group_count=51,
        manual_computer_group_count=50,
        fixlets_relevant_count=12_000,
        fixlets_total_count=70_000,
        tasks_relevant_count=2_000,
        tasks_total_count=10_001,
        analyses_relevant_count=500,
        analyses_total_count=10_001,
        active_analyses_count=500,
        properties_count=10_000,
        active_properties_count=501,
        custom_properties_count=25,
    )

    rows = loop.actionsite_pressure_rows(pressure)
    statuses = {row[0]: row[2] for row in rows}

    assert statuses["Subscribed sites"] == "Needs Attention"
    assert statuses["Computer groups"] == "Needs Attention"
    assert statuses["Fixlets and tasks"] == "Needs Attention"
    assert statuses["Analyses"] == "Needs Attention"
    assert statuses["Properties"] == "Needs Attention"


def test_aggregate_loop_offenders_rolls_up_detail_rows() -> None:
    df = pd.DataFrame(
        [
            {
                "ClientProcessingSeconds": 63.334,
                "Result": "Background Evaluation of Custom Site: Resolve",
            },
            {
                "ClientProcessingSeconds": 20.0,
                "Result": "Background Evaluation of Custom Site: Resolve",
            },
            {
                "ClientProcessingSeconds": 71.213,
                "Result": "Background Evaluation of External Site: Universal Checklist for Windows Server",
            },
        ]
    )

    offenders = loop.aggregate_loop_offenders(df, top_count=10)

    assert [item.result for item in offenders] == [
        "Background Evaluation of External Site: Universal Checklist for Windows Server",
        "Background Evaluation of Custom Site: Resolve",
    ]
    assert offenders[0].rank == 1
    assert offenders[0].occurrences == 1
    assert offenders[0].total_client_processing_seconds == 71.213
    assert round(offenders[1].avg_client_processing_seconds, 3) == 41.667


def test_load_loop_offenders_accepts_raw_evaluation_time_text(tmp_path: Path) -> None:
    timing_file = tmp_path / "property-evaluation-times.txt"
    timing_file.write_text(
        "\n".join(
            [
                "Count Percentage Result",
                "1 12.50% 16.319: CustomSite_Resolve.1055:Background Evaluation",
                "1 12.50% 20.060: Patching Support.344:Evaluate Property 1",
                "1 12.50% 27.486: CustomSite_Resolve.1055:Background Evaluation",
                "1 12.50% 71.213: Universal Checklist for Windows Server.214134858:Background Evaluation",
            ]
        ),
        encoding="utf-8",
    )

    offenders = loop.load_loop_offenders(timing_file, top_count=10)

    assert [item.result for item in offenders] == [
        "Background Evaluation of External Site: Universal Checklist for Windows Server",
        "Background Evaluation of Custom Site: Resolve",
        "Evaluate Property 1 of External Site: Patching Support",
    ]
    assert offenders[0].total_client_processing_seconds == 71.213
    assert offenders[1].occurrences == 2
    assert round(offenders[1].total_client_processing_seconds, 3) == 43.805


def test_aggregate_loop_offenders_accepts_prebuilt_top_offender_csv_shape() -> None:
    df = pd.DataFrame(
        [
            {
                "Rank": 1,
                "Occurrences": 9,
                "TotalClientProcessingSeconds": 180.5,
                "MaxClientProcessingSeconds": 42.0,
                "AvgClientProcessingSeconds": 20.055,
                "Result": "Background Evaluation of Custom Site: Resolve",
            },
            {
                "Rank": 2,
                "Occurrences": 2,
                "TotalClientProcessingSeconds": 122.0,
                "MaxClientProcessingSeconds": 61.0,
                "AvgClientProcessingSeconds": 61.0,
                "Result": "Evaluate Property 1 of External Site: Patching Support",
            },
        ]
    )

    offenders = loop.aggregate_loop_offenders(df, top_count=10)

    assert [item.result for item in offenders] == [
        "Evaluate Property 1 of External Site: Patching Support",
        "Background Evaluation of Custom Site: Resolve",
    ]
    assert offenders[0].occurrences == 2
    assert offenders[0].avg_client_processing_seconds == 61.0


def test_aggregate_loop_offenders_accepts_simplified_top10_csv_shape() -> None:
    df = pd.DataFrame(
        [
            {
                "Occurrences": 33,
                "Status": "OK",
                "Result": "Evaluate Property 1 of External Site: Patching Support",
            },
            {
                "Occurrences": 15,
                "Status": "OK",
                "Result": "Background Evaluation of Custom Site: Resolve",
            },
        ]
    )

    offenders = loop.aggregate_loop_offenders(df, top_count=10)

    assert [item.result for item in offenders] == [
        "Evaluate Property 1 of External Site: Patching Support",
        "Background Evaluation of Custom Site: Resolve",
    ]
    assert offenders[0].rank == 1
    assert offenders[0].occurrences == 33


def test_load_udp_notification_summary_rolls_up_relay_and_subnet(
    tmp_path: Path,
) -> None:
    udp_file = tmp_path / "udp-notification-results.csv"
    udp_file.write_text(
        "\n".join(
            [
                "Computer Name,IPv4 Addresses,Relay,Platform Health Check - Receiving UDP Pings",
                "client-1,10.10.1.5|192.168.1.5,relay-a,True",
                "client-2,10.10.1.6,relay-a,False",
                "client-3,10.10.2.7,relay-a,False",
                "client-4,10.10.1.8,relay-b,False",
                "client-5,,relay-b,Not Reported",
            ]
        ),
        encoding="utf-8",
    )

    summary = loop.load_udp_notification_summary(udp_file)

    assert summary is not None
    assert summary.total_devices == 5
    assert summary.receiving_count == 1
    assert summary.not_receiving_count == 3
    assert summary.unknown_count == 1
    assert summary.relay_rows == [
        ["relay-a", "2 / 3", "66.7%"],
        ["relay-b", "1 / 1", "100.0%"],
    ]
    assert summary.subnet_rows == [
        ["10.10.1.0/24", "2 / 3", "relay-a; relay-b"],
        ["10.10.2.0/24", "1 / 1", "relay-a"],
    ]


def test_udp_notification_notes_explain_missing_rollup_data(
    tmp_path: Path,
) -> None:
    udp_file = tmp_path / "udp-notification-results.csv"
    udp_file.write_text(
        "\n".join(
            [
                "Computer Name,IPv4 Addresses,Relay,Platform Health Check - Receiving UDP Pings",
                "client-1,10.20.1.5,relay-a,False",
                "client-2,No IPv4 Address,relay-a,False",
                "client-3,10.20.2.5,Not Reported,False",
                "client-4,10.20.2.6,relay-b,True",
            ]
        ),
        encoding="utf-8",
    )

    summary = loop.load_udp_notification_summary(udp_file)

    assert summary is not None
    assert summary.not_receiving_count == 3
    assert summary.not_receiving_missing_relay_count == 1
    assert summary.not_receiving_missing_subnet_count == 1
    assert summary.not_receiving_missing_rollup_data_count == 2
    assert summary.relay_rows == [["relay-a", "2 / 2", "100.0%"]]
    assert summary.subnet_rows == [
        ["10.20.1.0/24", "1 / 1", "relay-a"],
        ["10.20.2.0/24", "1 / 2", "Not reported"],
    ]

    notes = loop.udp_notification_data_notes(summary)
    assert len(notes) == 1
    assert (
        "2 client(s) not receiving UDP pings did not report a relay and/or IPv4 address"
        in notes[0]
    )

    word_output = tmp_path / "loop.docx"
    png_output = _write_png(tmp_path / "loop.png")
    loop.generate_word(
        _sample_loop_metrics(),
        png_output,
        word_output,
        udp_summary=summary,
    )

    paragraphs = [p.text for p in Document(str(word_output)).paragraphs if p.text.strip()]
    assert any(
        not p.startswith("NOTE") and "did not report a relay and/or IPv4 address" in p
        for p in paragraphs
    )


def test_discovers_loop_offenders_csv_near_loop_input(tmp_path: Path) -> None:
    input_file = tmp_path / "Results" / "Analyses" / "Agent_Performance_Counters"
    input_file.mkdir(parents=True)
    loop_csv = input_file / "Agent Performance Counters - Basic.csv"
    loop_csv.write_text(
        "Computer Name,Maximum Cycle Seconds,Average Cycle Seconds\nclient-1,120,60\n",
        encoding="utf-8",
    )
    offenders_csv = tmp_path / "Results" / "property-query-results.csv"
    offenders_csv.write_text(
        "\n".join(
            [
                "Rank,Occurrences,TotalClientProcessingSeconds,MaxClientProcessingSeconds,AvgClientProcessingSeconds,Result",
                "1,2,83.334,63.334,41.667,Background Evaluation of Custom Site: Resolve",
            ]
        ),
        encoding="utf-8",
    )

    assert loop.discover_loop_offenders_file(loop_csv) == offenders_csv.resolve()


def test_discovers_simplified_top10_csv_near_client_loop_input(tmp_path: Path) -> None:
    client_loop_dir = tmp_path / "Results" / "Client_Loop"
    client_loop_dir.mkdir(parents=True)
    loop_csv = client_loop_dir / "Client_Loop_07-02-2026_1732.csv"
    loop_csv.write_text(
        "Computer Name,Maximum Cycle Minutes,Average Cycle Minutes\nclient-1,16,4.3\n",
        encoding="utf-8",
    )
    offenders_csv = client_loop_dir / "top10-query-results.csv"
    offenders_csv.write_text(
        "\n".join(
            [
                "Occurrences,Status,Result",
                "15,OK,Background Evaluation of Custom Site: Resolve",
            ]
        ),
        encoding="utf-8",
    )

    assert loop.discover_loop_offenders_file(loop_csv) == offenders_csv.resolve()


def test_discovers_udp_notification_csv_near_client_loop_input(
    tmp_path: Path,
) -> None:
    client_loop_dir = tmp_path / "Results" / "Client_Loop"
    client_loop_dir.mkdir(parents=True)
    loop_csv = client_loop_dir / "Client_Loop_07-02-2026_1732.csv"
    loop_csv.write_text(
        "Computer Name,Maximum Cycle Minutes,Average Cycle Minutes\nclient-1,16,4.3\n",
        encoding="utf-8",
    )
    udp_csv = client_loop_dir / "udp-notification-results.csv"
    udp_csv.write_text(
        "\n".join(
            [
                "Computer Name,IPv4 Addresses,Relay,Platform Health Check - Receiving UDP Pings",
                "client-1,10.10.1.5,relay-a,True",
            ]
        ),
        encoding="utf-8",
    )

    assert loop.discover_udp_notifications_file(loop_csv) == udp_csv.resolve()
