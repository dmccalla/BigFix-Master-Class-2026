# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""11.0.7 agent multicore configuration coverage in the client performance path."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import get_loop_health as loop
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook


def _sample_rows() -> list[loop.ProcessedRow]:
    df = pd.DataFrame(
        {
            "Computer Name": ["scaled-client", "default-client", "scaled-client-2"],
            "Maximum Cycle Minutes": [20, 40, 80],
            "Average Cycle Minutes": [10, 30, 60],
            "Platform Health Check - Agent Multicore Scaling Enabled": [
                "1",
                "Not Set",
                "1",
            ],
            "Platform Health Check - Agent Total Physical CPU Core Count": [
                "8",
                "4",
                "8",
            ],
        }
    )
    df = loop.normalize_loop_time_columns(
        df, "Maximum Cycle Minutes", "Average Cycle Minutes"
    )
    processed, not_reported, errors = loop.process_rows(
        df,
        "Computer Name",
        "Maximum Cycle Minutes",
        "Average Cycle Minutes",
    )
    assert not_reported == 0
    assert errors == 0
    return processed


def _sample_metrics(
    processed: list[loop.ProcessedRow],
) -> tuple[loop.ReportMetrics, list[loop.ProcessedRow]]:
    return loop.calculate_metrics(
        processed=processed,
        total_source_rows=len(processed),
        not_reported_count=0,
        error_count=0,
        healthy_threshold_seconds=1800,
        critical_threshold_seconds=3600,
        report_start=date(2026, 9, 1),
        report_end=date(2026, 9, 2),
        generated_at=datetime(2026, 9, 30, 10, 30, 0),
    )


def test_agent_multicore_settings_are_reported_without_assessment() -> None:
    processed = _sample_rows()
    summary = loop.analyze_agent_multicore(processed)

    assert summary is not None
    assert summary.total_devices == 3
    assert summary.settings_rows == [
        ["On", "2 / 3"],
        ["Not Set", "1 / 3"],
    ]

    source = loop.processed_to_dataframe(processed)
    assert list(
        source[["Client Multicore Scaling", "Client Total Physical CPU Core Count"]].iloc[
            0
        ]
    ) == ["On", "8"]


def test_agent_multicore_workbook_and_word_sections_are_written(tmp_path: Path) -> None:
    processed = _sample_rows()
    summary = loop.analyze_agent_multicore(processed)
    assert summary is not None
    metrics, fix_first = _sample_metrics(processed)

    workbook = tmp_path / "loop.xlsx"
    loop.generate_excel(
        metrics,
        processed,
        fix_first,
        workbook,
        agent_multicore_summary=summary,
    )
    wb = load_workbook(workbook)
    assert "Client Multicore Settings" in wb.sheetnames
    settings = list(wb["Client Multicore Settings"].values)
    assert settings[0] == (
        "Scale Work Time With CPU Cores",
        "Devices / Total",
    )
    assert settings[1] == ("On", "2 / 3")

    word_output = tmp_path / "loop.docx"
    loop.generate_word(
        metrics,
        tmp_path / "missing.png",
        word_output,
        agent_multicore_summary=summary,
    )
    doc = Document(str(word_output))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Client Multicore Settings" in headings
    assert any("configuration context only" in p.text for p in doc.paragraphs)
    assert any("earlier than BigFix 11.0.7" in p.text for p in doc.paragraphs)
    multicore_table = next(
        table
        for table in doc.tables
        if table.rows[0].cells[0].text == "Scale Work Time With CPU Cores"
    )
    assert (
        multicore_table.rows[1].cells[1].paragraphs[0].alignment
        == WD_ALIGN_PARAGRAPH.CENTER
    )
