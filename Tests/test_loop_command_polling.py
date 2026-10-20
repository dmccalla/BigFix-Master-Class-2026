# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import get_loop_health as loop
import pandas as pd
from docx import Document
from openpyxl import load_workbook

_HEADER = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<BESAPI xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
    'xsi:noNamespaceSchemaLocation="BESAPI.xsd">\n'
)


def _doc(answers: list[str]) -> str:
    body = "".join(f'\t\t\t<Answer type="string">{a}</Answer>\n' for a in answers)
    return (
        f"{_HEADER}\t<Query>\n\t\t<Result>\n{body}\t\t</Result>\n\t</Query>\n</BESAPI>\n"
    )


def _sample_polling_rows() -> list[loop.ProcessedRow]:
    df = pd.DataFrame(
        {
            "Computer Name": ["internal-ok", "dmz-off", "soho-low"],
            "Maximum Cycle Minutes": [20, 80, 120],
            "Average Cycle Minutes": [10, 75, 90],
            "Platform Health Check - Command Polling Enabled": ["1", "0", "1"],
            "Platform Health Check - Command Polling Minutes": [
                "20",
                "Not Enabled",
                "60",
            ],
        }
    )
    df = loop.normalize_loop_time_columns(
        df, "Maximum Cycle Minutes", "Average Cycle Minutes"
    )
    enabled_col, interval_col = loop.auto_detect_command_polling_columns(df)

    processed, not_reported, errors = loop.process_rows(
        df,
        "Computer Name",
        "Maximum Cycle Minutes",
        "Average Cycle Minutes",
        udp_data={
            "internal-ok": {
                "receiving": "Yes",
                "relay": "relay-core",
                "subnet": "10.10.1.0/24",
            },
            "dmz-off": {
                "receiving": "No",
                "relay": "dmz-relay-01",
                "subnet": "192.168.10.0/24",
            },
            "soho-low": {
                "receiving": "No",
                "relay": "relay-edge-01",
                "subnet": "192.168.20.0/24",
            },
        },
        command_polling_enabled_column=enabled_col,
        command_polling_interval_column=interval_col,
    )

    assert not_reported == 0
    assert errors == 0
    return processed


def _sample_metrics(
    processed: list[loop.ProcessedRow],
) -> tuple[loop.ReportMetrics, list[loop.ProcessedRow]]:
    return loop.calculate_metrics(
        processed=processed,
        total_source_rows=len(processed),
        not_reported_count=0,
        error_count=0,
        healthy_threshold_seconds=1800,
        critical_threshold_seconds=3600,
        report_start=date(2026, 9, 1),
        report_end=date(2026, 9, 2),
        generated_at=datetime(2026, 9, 14, 10, 30, 0),
    )


def test_command_polling_analysis_groups_by_relay_and_subnet() -> None:
    processed = _sample_polling_rows()
    summary = loop.analyze_command_polling(processed)

    assert summary is not None
    assert summary.enabled_count == 2
    assert summary.disabled_count == 1
    assert summary.not_set_count == 0
    assert summary.interval_below_loop_count == 1
    assert summary.no_udp_polling_gap_count == 1

    dmz_row = next(row for row in summary.relay_rows if row[0] == "dmz-relay-01")
    assert dmz_row[8] == "1"

    soho_row = next(row for row in summary.subnet_rows if row[0] == "192.168.20.0/24")
    assert soho_row[7] == "1"

    slow_client = next(row for row in processed if row.computer_name == "soho-low")
    internal_client = next(row for row in processed if row.computer_name == "internal-ok")
    assert internal_client.command_polling_interval_seconds == 1200
    assert slow_client.command_polling_interval_seconds == 3600
    assert slow_client.command_polling_review == (
        "Current polling interval is below the average loop time."
    )


def test_command_polling_workbook_tabs_are_written(tmp_path: Path) -> None:
    processed = _sample_polling_rows()
    summary = loop.analyze_command_polling(processed)
    assert summary is not None

    metrics, fix_first = _sample_metrics(processed)
    workbook = tmp_path / "loop.xlsx"

    loop.generate_excel(
        metrics,
        processed,
        fix_first,
        workbook,
        command_polling_summary=summary,
    )

    wb = load_workbook(workbook)
    assert "Polling Summary" in wb.sheetnames
    assert "Polling Context" not in wb.sheetnames
    assert "Polling by Relay" in wb.sheetnames
    assert "Polling by Subnet" in wb.sheetnames
    assert "Polling Exceptions" not in wb.sheetnames
    relay_headers = [cell.value for cell in wb["Polling by Relay"][1]]
    assert relay_headers == [
        "Relay",
        "Total Devices",
        "Enabled",
        "Disabled",
        "NotSet",
        "Unknown",
        "EnabledMissingInterval",
        "IntervalBelowLoop",
        "NoUdpAndPollingOff",
    ]
    subnet_headers = [cell.value for cell in wb["Polling by Subnet"][1]]
    assert subnet_headers[1] == "Total Devices"
    source_headers = [cell.value for cell in wb["Source Data"][1]]
    assert "Command Polling" in source_headers
    assert "Command Polling Minutes" in source_headers
    assert "Suggested Command Polling Interval Seconds" in source_headers
    assert "Suggested Command Polling Interval" not in source_headers
    suggested_idx = source_headers.index("Suggested Command Polling Interval Seconds")
    current_idx = source_headers.index("Command Polling Minutes")
    source_rows = {
        row[0]: row for row in wb["Source Data"].iter_rows(min_row=2, values_only=True)
    }
    assert source_rows["dmz-off"][suggested_idx] == 5100
    assert source_rows["internal-ok"][current_idx] == "20 min"


def test_command_polling_word_section_is_written(tmp_path: Path) -> None:
    processed = _sample_polling_rows()
    summary = loop.analyze_command_polling(processed)
    assert summary is not None
    metrics, _fix_first = _sample_metrics(processed)
    word_output = tmp_path / "loop.docx"

    loop.generate_word(
        metrics,
        tmp_path / "missing.png",
        word_output,
        command_polling_summary=summary,
    )

    doc = Document(str(word_output))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Notification and Polling" in headings
    assert "Command Polling" in headings
    assert any(
        [cell.text for cell in table.rows[0].cells]
        == ["Status", "Devices / Total", "% of Total"]
        for table in doc.tables
    )
    assert any("devices that reported polling data." in p.text for p in doc.paragraphs)
    assert not any("generated with this module" in p.text for p in doc.paragraphs)
    assert not any(
        "Context Signal" in [cell.text for cell in table.rows[0].cells]
        for table in doc.tables
    )
    assert not any(
        "Actual Intervals" in [cell.text for cell in table.rows[0].cells]
        for table in doc.tables
    )
    assert "Command Polling Exceptions" not in headings
    # The Mini Report keeps 6 polling columns. Excel keeps all 9.
    six = [
        "Total Devices",
        "Enabled",
        "Missing Interval",
        "Interval Below Loop",
        "No UDP + Polling Off",
    ]
    header_rows = [[cell.text for cell in table.rows[0].cells] for table in doc.tables]
    assert ["Relay", *six] in header_rows
    assert ["Subnet", *six] in header_rows
    assert not any(
        "Not Set" in row for row in header_rows if row[0] in {"Relay", "Subnet"}
    )
    # CP-23: the lead states the three check rules.
    assert any(
        p.text.startswith("Interval Below Average Loop counts devices with polling")
        for p in doc.paragraphs
    )
    # CP-24: no Top 10 data, so the title paragraph says so.
    assert any(
        "This capture did not include the Top 10 evaluation data. This Mini Report "
        "has no Top 10 Most Expensive Evaluations section." in p.text
        for p in doc.paragraphs
    )
