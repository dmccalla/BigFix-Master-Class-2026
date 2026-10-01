# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""CPU idle client-setting coverage in the Agent Performance report path."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import get_loop_health as loop
import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook


def _sample_rows() -> list[loop.ProcessedRow]:
    df = pd.DataFrame(
        {
            "Computer Name": ["configured-client", "default-client"],
            "Maximum Cycle Minutes": [20, 40],
            "Average Cycle Minutes": [10, 30],
            "Platform Health Check - Agent CPU Work Idle": ["25", "Not Set"],
            "Platform Health Check - Agent CPU Sleep Idle": ["400", "Not Set"],
        }
    )
    df = loop.normalize_loop_time_columns(
        df, "Maximum Cycle Minutes", "Average Cycle Minutes"
    )
    processed, not_reported, errors = loop.process_rows(
        df,
        "Computer Name",
        "Maximum Cycle Minutes",
        "Average Cycle Minutes",
    )
    assert not_reported == 0
    assert errors == 0
    return processed


def _sample_metrics(
    processed: list[loop.ProcessedRow],
) -> tuple[loop.ReportMetrics, list[loop.ProcessedRow]]:
    return loop.calculate_metrics(
        processed=processed,
        total_source_rows=len(processed),
        not_reported_count=0,
        error_count=0,
        healthy_threshold_seconds=1800,
        critical_threshold_seconds=3600,
        report_start=date(2026, 9, 1),
        report_end=date(2026, 9, 2),
        generated_at=datetime(2026, 9, 30, 12, 0, 0),
    )


def test_cpu_idle_settings_are_reported_without_assessment() -> None:
    processed = _sample_rows()
    summary = loop.analyze_agent_cpu_idle(processed)

    assert summary is not None
    assert summary.total_devices == 2
    assert summary.settings_rows == [
        ["25", "400", "5.9%", "1"],
        ["Not Set", "Not Set", "Not Set", "1"],
    ]

    source = loop.processed_to_dataframe(processed)
    assert list(
        source[
            [
                "Agent CPU Work Idle",
                "Agent CPU Sleep Idle",
                "Agent CPU Idle Configured Duty Cycle (%)",
            ]
        ].iloc[0]
    ) == ["25", "400", "5.9%"]


def test_configured_cpu_duty_cycle_preserves_unreported_states() -> None:
    assert loop.configured_cpu_duty_cycle("20 ms", "60 ms") == "25.0%"
    assert loop.configured_cpu_duty_cycle("25", "400") == "5.9%"
    assert loop.configured_cpu_duty_cycle("Not Set", "Not Set") == "Not Set"
    assert (
        loop.configured_cpu_duty_cycle("Not Applicable", "Not Applicable")
        == "Not Applicable"
    )
    assert loop.configured_cpu_duty_cycle("0", "0") == "Unknown"


def test_cpu_idle_settings_are_written_to_excel_and_word(tmp_path: Path) -> None:
    processed = _sample_rows()
    summary = loop.analyze_agent_cpu_idle(processed)
    assert summary is not None
    metrics, fix_first = _sample_metrics(processed)

    workbook = tmp_path / "loop.xlsx"
    loop.generate_excel(
        metrics,
        processed,
        fix_first,
        workbook,
        agent_cpu_idle_summary=summary,
    )
    wb = load_workbook(workbook)
    assert "Agent CPU Idle Settings" in wb.sheetnames
    assert list(wb["Agent CPU Idle Settings"].values)[0] == (
        "Work Idle (ms)",
        "Sleep Idle (ms)",
        "Configured Duty Cycle (%)",
        "Devices",
    )

    word_output = tmp_path / "loop.docx"
    loop.generate_word(
        metrics,
        tmp_path / "missing.png",
        word_output,
        agent_cpu_idle_summary=summary,
    )
    doc = Document(str(word_output))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Agent CPU Idle Settings" in headings
    assert any("configuration context only" in p.text for p in doc.paragraphs)
    assert any(
        "not a measurement of observed CPU usage" in p.text for p in doc.paragraphs
    )


def test_headerless_fourteen_field_loop_capture_maps_cpu_idle_columns(
    tmp_path: Path,
) -> None:
    xml = tmp_path / "Client_Loop_10-01-2026_1200.xml"
    xml.write_text(
        """<?xml version="1.0" encoding="UTF-8"?>
<BESAPI><Query><Result>
<Answer type="string">CLIENT1,42,17,Mon, 28 Sep 2026 10:00:00 -0400,True,1,3600,1,20,60,1,25,400,8</Answer>
<Answer type="string">CLIENT2, Example Office,21,15.0,Mon, 28 Sep 2026 10:00:00 -0400,False,0,Not Enabled,0,20,60,Not Applicable,10,480,4</Answer>
</Result></Query></BESAPI>
""",
        encoding="utf-8",
    )

    df = loop.load_loop_input(xml)

    assert list(df.columns[-3:]) == [
        "Platform Health Check - Agent CPU Work Idle",
        "Platform Health Check - Agent CPU Sleep Idle",
        "Platform Health Check - Agent Total Physical CPU Core Count",
    ]
    assert (
        str(df.iloc[1]["Platform Health Check - Agent Multicore Scaling Enabled"])
        == "Not Applicable"
    )
    assert str(df.iloc[1]["Platform Health Check - Agent CPU Work Idle"]) == "10"
    assert str(df.iloc[1]["Platform Health Check - Agent CPU Sleep Idle"]) == "480"
    assert (
        str(df.iloc[1]["Platform Health Check - Agent Total Physical CPU Core Count"])
        == "4"
    )

    df = loop.normalize_loop_time_columns(
        df, "Maximum Cycle Minutes", "Average Cycle Minutes"
    )
    processed, not_reported, errors = loop.process_rows(
        df,
        "Computer Name",
        "Maximum Cycle Minutes",
        "Average Cycle Minutes",
    )
    assert not_reported == 0
    assert errors == 0
    summary = loop.analyze_agent_multicore(processed)
    assert summary is not None
    assert ["Not Applicable", "4", "1"] in summary.settings_rows


def test_headerless_loop_capture_requires_all_collection_fields(tmp_path: Path) -> None:
    xml = tmp_path / "Client_Loop_10-01-2026_1210.xml"
    xml.write_text(
        """<?xml version="1.0" encoding="UTF-8"?>
<BESAPI><Query><Result>
<Answer type="string">CLIENT1,42,17,Mon, 28 Sep 2026 10:00:00 -0400,True,1,3600,1,20,60,1</Answer>
</Result></Query></BESAPI>
""",
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="all 14 collection fields"):
        loop.load_loop_input(xml)
