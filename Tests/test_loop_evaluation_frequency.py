# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""Evaluation Frequency: active property intervals in the Client Performance report."""

from __future__ import annotations

from pathlib import Path

import get_loop_health as loop
import pytest
from docx import Document

PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR"
    b"\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde"
    b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfeA\xe2%\xb5"
    b"\x00\x00\x00\x00IEND\xaeB`\x82"
)

SECTION = "//===== Active Properties by Evaluation Interval ================"
HEADERS = (
    "Column Headers: Property Name, Source Analysis Name, Source Analysis Site Name, "
    "Site ID, Analysis ID, Property ID, Issued By, Evaluation Interval, Content Source"
)
APP_ANALYSIS = "Application Information (Windows)"
APP_PROPERTY = "Installed Applications - Windows"

# name, analysis, site, interval text, content source
PROPERTIES = [
    (APP_PROPERTY, APP_ANALYSIS, "BES Support", "00:00:00", "External"),
    ("Next Patch Window", "Patch Window Status", "Staples", "00:00:00", "Custom"),
    ("BIOS", "Hardware Information (Windows)", "BES Support", "01:00:00", "External"),
    ("Daily Thing", "Daily Analysis", "BES Support", "1 day", "External"),
    ("Six Hour", "Six Analysis", "Staples", "06:00:00", "Custom"),
    ("Quick", "Quick Analysis", "Staples", "00:05:00", "Custom"),
    ("Unset", "Unset Analysis", "BES Support", "no-val", "External"),
]


def _info_lines(with_section: bool = True) -> list[str]:
    lines = [
        "PROPERTY INFORMATION",
        " ",
        "Number of Properties: 10",
        "//===== Custom BES Properties with frequent evaluations (sub 15 minutes) ====",
        "Column Headers: Property Name, Source Analysis Site Name, Property ID, "
        "Issued By, Evaluation Interval",
        " ",
        "Quick",
        "Staples",
        "1",
        "2",
        "3",
        "bigfix",
        "00:05:00",
    ]
    if not with_section:
        return lines
    lines += [SECTION, HEADERS, " "]
    for name, analysis, site, interval, source in PROPERTIES:
        lines += [name, analysis, site, "1", "2", "3", "bigfix", interval, source]
    return lines


def _write_log(folder: Path, with_section: bool = True) -> Path:
    path = folder / "Property_Info_BFROOT1_10-01-2026_1200.log"
    path.write_text("\n".join(_info_lines(with_section)), encoding="utf-8")
    return path


def _write_xml(folder: Path) -> Path:
    path = folder / "Property_Info_BFROOT1_10-01-2026_1200.xml"
    answers = [f'<Answer type="string">{line}</Answer>' for line in _info_lines()]
    path.write_text(
        '<?xml version="1.0" encoding="UTF-8"?><BESAPI><Query><Result>'
        + "".join(answers)
        + "</Result></Query></BESAPI>",
        encoding="utf-8",
    )
    return path


@pytest.mark.parametrize(
    ("text", "minutes"),
    [
        ("00:00:00", 0),
        ("00:05:00", 5),
        ("01:00:00", 60),
        ("06:00:00", 360),
        ("1 day", 1440),
        ("2 days", 2880),
        ("1 day, 06:00:00", 1800),
        ("no-val", 0),
        ("", 0),
        ("garbage", None),
    ],
)
def test_interval_text_converts_to_minutes(text: str, minutes: float | None) -> None:
    assert loop.parse_interval_minutes(text) == minutes


@pytest.mark.parametrize(
    ("minutes", "band"),
    [
        (None, "Unreadable"),
        (0, "Every report"),
        (5, "Under 15 minutes"),
        (15, "15 minutes to 1 hour"),
        (60, "1 hour to 6 hours"),
        (359, "1 hour to 6 hours"),
        (360, "6 hours or longer"),
        (1440, "6 hours or longer"),
    ],
)
def test_interval_bands(minutes: float | None, band: str) -> None:
    assert loop.interval_band(minutes) == band


def test_property_intervals_parse_from_log_and_xml_alike(tmp_path: Path) -> None:
    log_dir = tmp_path / "log"
    xml_dir = tmp_path / "xml"
    log_dir.mkdir()
    xml_dir.mkdir()
    _write_log(log_dir)
    _write_xml(xml_dir)

    from_log = loop.parse_property_intervals(log_dir)
    from_xml = loop.parse_property_intervals(xml_dir)

    assert from_log is not None
    assert from_xml is not None
    assert len(from_log.properties) == len(PROPERTIES)
    assert from_log.properties == from_xml.properties
    first = from_log.properties[0]
    assert (first.name, first.analysis, first.site) == (
        APP_PROPERTY,
        APP_ANALYSIS,
        "BES Support",
    )
    assert first.interval_minutes == 0
    assert first.content_source == "External"


def test_property_intervals_are_none_for_a_capture_without_the_section(
    tmp_path: Path,
) -> None:
    _write_log(tmp_path, with_section=False)

    assert loop.parse_property_intervals(tmp_path) is None
    assert loop.parse_property_intervals(None) is None


def test_evaluation_frequency_rows_count_bands_and_status(tmp_path: Path) -> None:
    _write_log(tmp_path)
    summary = loop.parse_property_intervals(tmp_path)

    rows = loop.evaluation_frequency_rows(summary)

    assert rows == [
        ["Every report", "3 / 7", "42.9%", "1", "2", "Needs Attention"],
        ["Under 15 minutes", "1 / 7", "14.3%", "1", "0", "Needs Attention"],
        ["15 minutes to 1 hour", "0 / 7", "0.0%", "0", "0", "Advisory"],
        ["1 hour to 6 hours", "1 / 7", "14.3%", "0", "1", "Advisory"],
        ["6 hours or longer", "2 / 7", "28.6%", "1", "1", "Healthy"],
    ]
    assert loop.frequent_property_count(summary) == 4
    assert loop.frequent_property_count(None) == 0


def _offender(avg: float, result: str) -> loop.LoopOffenderSummary:
    return loop.LoopOffenderSummary(
        rank=1,
        occurrences=5,
        total_client_processing_seconds=avg * 5,
        max_client_processing_seconds=avg,
        avg_client_processing_seconds=avg,
        result=result,
    )


def _analysis_result(analysis: str, name: str) -> str:
    return (
        f"Analysis Name={analysis} | Display Name={name} | Name={name} | "
        f"Source Name={name} | Source ID=1"
    )


def test_top10_interval_names_the_property_interval_or_why_there_is_none(
    tmp_path: Path,
) -> None:
    _write_log(tmp_path)
    summary = loop.parse_property_intervals(tmp_path)
    offenders = [
        _offender(3.38, _analysis_result("Hardware Information (Windows)", "BIOS")),
        _offender(68.38, _analysis_result(APP_ANALYSIS, APP_PROPERTY)),
        _offender(50.0, "Background Evaluation of External Site: Enterprise Security"),
        _offender(9.0, _analysis_result("Not Collected", "Missing Property")),
        _offender(7.0, _analysis_result("Unset Analysis", "Unset")),
    ]

    intervals = [loop.offender_interval(o.result, summary) for o in offenders]

    assert intervals == [
        "1 hour",
        "Every report",
        "Not a property",
        "Not found",
        "Every report",
    ]
    assert loop.offender_interval(offenders[0].result, None) == "Not collected"


def _healthy_metrics() -> loop.ReportMetrics:
    return loop.ReportMetrics(
        total_devices=10,
        score_total_devices=10,
        overall_status="Healthy",
        environment_score=95,
        score_driver="test driver",
        score_breakdown="test breakdown",
        report_start_text="7/1/2026",
        report_end_text="7/2/2026",
        report_span_text="2 days",
        generated_at_text="7/2/2026 17:32",
        avg_mean_seconds=720,
        avg_median_seconds=240,
        avg_max_seconds=600,
        slowest_avg_seconds=900,
        peak_max_seconds=1200,
        avg_healthy=8,
        avg_warning=2,
        avg_critical=0,
        avg_healthy_pct=0.8,
        avg_warning_pct=0.2,
        avg_critical_pct=0.0,
        avg_over_120=0,
        avg_over_120_pct=0.0,
        not_reported_count=0,
        error_count=0,
        not_reported_pct=0.0,
        error_pct=0.0,
        total_issue_count=0,
        total_issue_pct=0.0,
        healthy_threshold_seconds=1800,
        critical_threshold_seconds=3600,
        fix_first_count=0,
    )


def test_word_report_places_evaluation_frequency_inside_loop_cost(
    tmp_path: Path,
) -> None:
    _write_log(tmp_path)
    summary = loop.parse_property_intervals(tmp_path)
    pressure = loop.ActionSitePressure(
        deployment_info_folder=tmp_path, actionsite_size_kb=1000
    )
    offenders = [_offender(68.38, _analysis_result(APP_ANALYSIS, APP_PROPERTY))]
    word_output = tmp_path / "loop.docx"
    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)

    loop.generate_word(
        _healthy_metrics(),
        png,
        word_output,
        actionsite_pressure=pressure,
        loop_offenders=offenders,
        property_intervals=summary,
        include_automated_findings=True,
    )

    doc = Document(str(word_output))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    order = [
        "Automated Findings",
        "Loop Health",
        "Loop Cost",
        "Top 10 Most Expensive Evaluations",
        "Evaluation Frequency",
        "Contributing Factors",
    ]
    assert "Cost of Repeated Evaluations" not in headings
    assert [h for h in headings if h in order] == order

    findings = next(t for t in doc.tables if t.cell(0, 0).text == "Finding ID")
    assert [c.text for c in findings.rows[1].cells] == [
        "CLT-01",
        "Medium.\n4 active properties are set to every report or to under 15 minutes. "
        "See Evaluation Frequency.",
    ]
    assert not any(t.cell(0, 0).text == "Property" for t in doc.tables)
    paragraphs = [p.text for p in doc.paragraphs]
    assert not any("could not be read" in p for p in paragraphs)


def test_word_report_omits_evaluation_frequency_without_the_data(
    tmp_path: Path,
) -> None:
    word_output = tmp_path / "loop.docx"
    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)

    loop.generate_word(
        _healthy_metrics(), png, word_output, include_automated_findings=True
    )

    headings = [
        p.text
        for p in Document(str(word_output)).paragraphs
        if p.style.name.startswith("Heading")
    ]
    assert "Evaluation Frequency" not in headings
    assert "Automated Findings" not in headings


# --- 10-18 review wording -------------------------------------------------------------


def test_top10_cell_names_the_property_and_its_analysis() -> None:
    assert (
        loop.format_loop_offender_evaluation(_analysis_result(APP_ANALYSIS, APP_PROPERTY))
        == f"{APP_PROPERTY} (analysis: {APP_ANALYSIS})"
    )
    site = "Background Evaluation of External Site: Enterprise Security"
    assert loop.format_loop_offender_evaluation(site) == site


def test_clt01_names_the_score_that_raises_it() -> None:
    metrics = _healthy_metrics()
    metrics.overall_status = "WARNING"
    reasons = loop.loop_finding_reasons(metrics)
    assert loop.loop_finding_text(reasons) == (
        "High. The client loop health score is below 90."
    )
    metrics.overall_status = "CRITICAL"
    reasons = loop.loop_finding_reasons(metrics)
    assert loop.loop_finding_text(reasons) == (
        "Critical. The client loop health score is below 70."
    )


def test_findings_line_counts_reasons_of_the_one_finding(tmp_path: Path) -> None:
    _write_log(tmp_path)
    summary = loop.parse_property_intervals(tmp_path)
    metrics = _healthy_metrics()
    metrics.overall_status = "WARNING"
    both = loop.loop_finding_reasons(metrics, summary)
    assert loop.loop_findings_line(both) == "Raises 1 High finding (2 reasons)"
    score_only = loop.loop_finding_reasons(metrics)
    assert loop.loop_findings_line(score_only) == "Raises 1 High finding"
    frequency_only = loop.loop_finding_reasons(_healthy_metrics(), summary)
    assert loop.loop_findings_line(frequency_only) == "Raises 1 Medium finding"
    assert loop.loop_findings_line([]) == ""


def test_infographic_lists_one_clt01_and_names_no_data_as_critical(
    monkeypatch, tmp_path: Path
) -> None:
    _write_log(tmp_path)
    summary = loop.parse_property_intervals(tmp_path)
    metrics = _healthy_metrics()
    metrics.overall_status = "WARNING"
    captured: dict = {}

    def capture_chart(kind, data, output):
        captured.update({"kind": kind, "data": data})

    monkeypatch.setattr("Lib.Design.html_charts.render_chart", capture_chart)

    loop.generate_png(metrics, tmp_path / "loop.png", property_intervals=summary)

    data = captured["data"]
    assert data["findings"] == [
        {
            "id": "CLT-01",
            "text": "High. The client loop health score is below 90. "
            "4 active properties are set to every report or to under 15 minutes.",
        }
    ]
    assert data["findingsLine"] == "Raises 1 High finding (2 reasons)"
    labels = [item["label"] for item in data["legend"]]
    assert labels[-2:] == ["Critical over 60 min", "Critical no usable data"]
    assert data["legend"][-1]["hatch"] is True


def test_word_report_uses_the_approved_sentences(tmp_path: Path) -> None:
    _write_log(tmp_path)
    summary = loop.parse_property_intervals(tmp_path)
    offenders = [_offender(68.38, _analysis_result(APP_ANALYSIS, APP_PROPERTY))]
    word_output = tmp_path / "loop.docx"
    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)

    loop.generate_word(
        _healthy_metrics(),
        png,
        word_output,
        loop_offenders=offenders,
        property_intervals=summary,
        include_automated_findings=True,
    )

    doc = Document(str(word_output))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "A property set to every report runs on every loop. So does a property" in text
    assert "with no interval, because BigFix evaluates it every report." in text
    assert "uses the loop time" not in text

    key = next(t for t in doc.tables if t.cell(0, 0).text == "Infographic tile")
    rows = {row.cells[0].text: [c.text for c in row.cells] for row in key.rows}
    assert rows["Findings raised"][1] == "CLT-01"
    assert (
        rows["Findings raised"][2]
        == "CLT-01 is raised for any of three reasons. The health score is below "
        "90. An active property is set to every report or to under 15 minutes. "
        "More than 5% of devices have no usable loop data."
    )
    assert rows["Loop distribution"][2].endswith(
        "The hatched part of Critical is devices with no usable data."
    )

    top10 = next(t for t in doc.tables if t.cell(0, 0).text == "Rank")
    assert top10.rows[1].cells[3].text == f"{APP_PROPERTY} (analysis: {APP_ANALYSIS})"
