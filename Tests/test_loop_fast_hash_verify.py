# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""FastHashVerify coverage in the Agent Performance report path."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import get_loop_health as loop
import pandas as pd
from docx import Document
from openpyxl import load_workbook

_HEADER = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<BESAPI xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
    'xsi:noNamespaceSchemaLocation="BESAPI.xsd">\n'
)


def _doc(answers: list[str]) -> str:
    body = "".join(f'\t\t\t<Answer type="string">{a}</Answer>\n' for a in answers)
    return (
        f"{_HEADER}\t<Query>\n\t\t<Result>\n{body}\t\t</Result>\n\t</Query>\n</BESAPI>\n"
    )


def _sample_fast_hash_rows() -> list[loop.ProcessedRow]:
    df = pd.DataFrame(
        {
            "Computer Name": ["enabled-client", "disabled-client", "default-client"],
            "Maximum Cycle Minutes": [20, 40, 80],
            "Average Cycle Minutes": [10, 30, 60],
            "Platform Health Check - FastHashVerify Enabled": ["1", "0", "Not Set"],
            "Platform Health Check - FastHashVerify Work ms": ["20", "20", "Not Set"],
            "Platform Health Check - FastHashVerify Sleep ms": [
                "60",
                "60",
                "Not Set",
            ],
        }
    )
    df = loop.normalize_loop_time_columns(
        df, "Maximum Cycle Minutes", "Average Cycle Minutes"
    )
    processed, not_reported, errors = loop.process_rows(
        df,
        "Computer Name",
        "Maximum Cycle Minutes",
        "Average Cycle Minutes",
    )
    assert not_reported == 0
    assert errors == 0
    return processed


def _sample_metrics(
    processed: list[loop.ProcessedRow],
) -> tuple[loop.ReportMetrics, list[loop.ProcessedRow]]:
    return loop.calculate_metrics(
        processed=processed,
        total_source_rows=len(processed),
        not_reported_count=0,
        error_count=0,
        healthy_threshold_seconds=1800,
        critical_threshold_seconds=3600,
        report_start=date(2026, 9, 1),
        report_end=date(2026, 9, 2),
        generated_at=datetime(2026, 9, 28, 10, 30, 0),
    )


def test_fast_hash_verify_analysis_and_source_data() -> None:
    processed = _sample_fast_hash_rows()
    summary = loop.analyze_fast_hash_verify(processed)

    assert summary is not None
    assert summary.total_devices == 3
    assert summary.enabled_count == 1
    assert summary.disabled_count == 1
    assert summary.not_set_count == 1
    assert summary.settings_rows == [["20 ms", "60 ms", "2"], ["Not Set", "Not Set", "1"]]

    source = loop.processed_to_dataframe(processed)
    assert list(
        source[["FastHashVerify", "FastHashVerify Work", "FastHashVerify Sleep"]].iloc[0]
    ) == [
        "Enabled",
        "20 ms",
        "60 ms",
    ]


def test_fast_hash_verify_workbook_and_word_sections_are_written(tmp_path: Path) -> None:
    processed = _sample_fast_hash_rows()
    summary = loop.analyze_fast_hash_verify(processed)
    assert summary is not None
    metrics, fix_first = _sample_metrics(processed)

    workbook = tmp_path / "loop.xlsx"
    loop.generate_excel(
        metrics,
        processed,
        fix_first,
        workbook,
        fast_hash_verify_summary=summary,
    )
    wb = load_workbook(workbook)
    assert "FastHashVerify Summary" in wb.sheetnames
    assert "FastHashVerify Settings" in wb.sheetnames
    source_headers = [cell.value for cell in wb["Source Data"][1]]
    assert "FastHashVerify" in source_headers
    assert "FastHashVerify Work" in source_headers
    assert "FastHashVerify Sleep" in source_headers

    word_output = tmp_path / "loop.docx"
    loop.generate_word(
        metrics,
        tmp_path / "missing.png",
        word_output,
        fast_hash_verify_summary=summary,
    )
    doc = Document(str(word_output))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "FastHashVerify Review" in headings
    assert any("large patch files" in p.text for p in doc.paragraphs)
    assert any(
        "Individual client FastHashVerify values include" in p.text
        for p in doc.paragraphs
    )
    assert not any(
        "Consider the customer's download workload" in p.text for p in doc.paragraphs
    )
