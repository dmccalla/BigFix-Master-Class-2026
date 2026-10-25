# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""Cases the 10-19-2026 lab run exposed that the sample data did not carry.

A Top 10 row can be a Retrieved Property, which has an interval the capture
does not hold. The UDP BESAPI capture has no header row, and the lab
collector writes ", " between fields where the older sample wrote ",".
"""

from __future__ import annotations

from pathlib import Path

import get_loop_health as loop
import run_diagnostics

RETRIEVED = (
    "Retrieved Property: BES Relay Service Installed (actionsite.13:Evaluate Property 1)"
)
SITE = "Background Evaluation of External Site: BES Support"

# Answer lines as the lab collector wrote them on 10-19-2026.
LAB_UDP_ANSWERS = [
    "BFRELAY1, 192.168.4.8, BFROOT1.dmc.lab:52311, Not Reported",
    "BFPLUGINPORTAL, 192.168.4.7, BFROOT1.dmc.lab:52311, True",
    "BFWEBUI, 192.168.4.4, BFROOT1.dmc.lab:52311, True",
    "BFROOT1, 192.168.4.3, BES Root Server, True",
    "DM1WIN1103, 192.168.4.127, BFROOT1.dmc.lab:52311, False",
]
# Answer lines as the older sample capture wrote them.
SAMPLE_UDP_ANSWERS = [
    "LAB-CLT-00001,10.24.1.1,BFRELAY001.dmc.lab:52311,False",
    "LAB-CLT-00002,10.24.1.2,BFRELAY001.dmc.lab:52311,True",
]


def _write_udp_xml(folder: Path, answers: list[str]) -> Path:
    path = folder / "UDP_Notifications_10-19-2026_1441.xml"
    body = "".join(f'<Answer type="string">{line}</Answer>' for line in answers)
    path.write_text(
        '<?xml version="1.0" encoding="UTF-8"?><BESAPI><Query><Result>'
        + body
        + "</Result></Query></BESAPI>",
        encoding="utf-8",
    )
    return path


def _write_udp_csv(folder: Path, answers: list[str]) -> Path:
    path = folder / "udp-notification-results.csv"
    header = '"Computer Name","IPv4 Addresses","Relay","Platform Health Check - Receiving UDP Pings"'
    rows = [
        ",".join(f'"{field.strip()}"' for field in line.split(",")) for line in answers
    ]
    path.write_text("\n".join([header, *rows]) + "\n", encoding="utf-8")
    return path


def test_retrieved_property_interval_is_not_collected() -> None:
    assert loop.offender_interval(RETRIEVED, None) == "Not collected"
    assert loop.offender_interval(SITE, None) == "Not a property"


def test_top10_lead_names_the_retrieved_property_case(tmp_path: Path) -> None:
    from docx import Document
    from test_loop_reporting_coverage import PNG_1X1, _metrics

    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)
    offender = loop.LoopOffenderSummary(
        rank=1,
        occurrences=4,
        total_client_processing_seconds=59.7,
        max_client_processing_seconds=15.8,
        avg_client_processing_seconds=14.9,
        result=RETRIEVED,
    )
    loop.generate_word(
        _metrics(10, 0), png, tmp_path / "loop.docx", loop_offenders=[offender]
    )

    doc = Document(str(tmp_path / "loop.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "A retrieved property's interval is not in the capture." in text
    descriptions = " ".join(shape._inline.docPr.get('descr', '') for shape in doc.inline_shapes)
    assert "interval Not collected" in descriptions
    assert "BES Relay Service Installed (retrieved property)" in descriptions


def test_udp_capture_reads_like_the_collector_csv(tmp_path: Path) -> None:
    for answers in (LAB_UDP_ANSWERS, SAMPLE_UDP_ANSWERS):
        folder = tmp_path / str(len(answers))
        folder.mkdir()
        from_xml = loop.load_udp_notification_summary(_write_udp_xml(folder, answers))
        from_csv = loop.load_udp_notification_summary(_write_udp_csv(folder, answers))
        assert from_xml is not None and from_csv is not None
        for field in (
            "total_devices",
            "receiving_count",
            "not_receiving_count",
            "unknown_count",
        ):
            assert getattr(from_xml, field) == getattr(from_csv, field), field
        assert from_xml.relay_rows == from_csv.relay_rows


def test_lab_udp_capture_counts(tmp_path: Path) -> None:
    summary = loop.load_udp_notification_summary(
        _write_udp_xml(tmp_path, LAB_UDP_ANSWERS)
    )
    assert (
        summary.total_devices,
        summary.receiving_count,
        summary.not_receiving_count,
        summary.unknown_count,
    ) == (5, 3, 1, 1)


def test_udp_capture_folds_extra_fields_into_the_computer_name() -> None:
    assert loop.udp_capture_rows("odd, name, 10.0.0.1, relay-a, True") == [
        ["odd, name", "10.0.0.1", "relay-a", "True"]
    ]
    assert (
        loop.udp_capture_rows("Computer Name,IPv4 Addresses,Relay,Receiving UDP Pings")
        == []
    )
    assert loop.udp_capture_rows("a, b, c, maybe") == []


def test_headerless_udp_capture_is_recognized(tmp_path: Path) -> None:
    path = _write_udp_xml(tmp_path, LAB_UDP_ANSWERS)
    assert loop._is_udp_notifications_file(path)
    assert run_diagnostics._looks_like_udp_notification_file(path)


def test_infographic_header_names_a_span_only_when_dates_were_given(
    monkeypatch, tmp_path: Path
) -> None:
    from test_loop_reporting_coverage import _metrics

    captured: dict = {}

    def capture_chart(kind, data, output):
        captured["meta"] = data["meta"]

    monkeypatch.setattr("Lib.Design.html_charts.render_chart", capture_chart)
    metrics = _metrics(10, 0)
    for span, expected in (
        ("Not provided", "10 devices scored"),
        ("Invalid range", "10 devices scored"),
        ("1 day", "10 devices scored &middot; 1 day collection"),
        ("6 days", "10 devices scored &middot; 6 days collection"),
    ):
        metrics.report_span_text = span
        loop.generate_png(metrics, tmp_path / "loop.png")
        assert captured["meta"] == expected


def test_activated_analyses_count_is_read(tmp_path: Path) -> None:
    # The Deployment Information task writes "Activated", as the lab capture shows.
    (tmp_path / "Content_Info_BFROOT1_10-19-2026_1444.log").write_text(
        "CONTENT INFORMATION\n"
        "Number of Analyses: 33 / 4404\n"
        "Number of Properties: 7261\n"
        "Number of Actived Properties: 265\n"
        "Number of Activated Analyses: 33\n",
        encoding="utf-8",
    )

    pressure = loop.parse_actionsite_pressure(tmp_path)

    assert pressure.active_analyses_count == 33


def test_rollup_note_counts_only_named_relays(tmp_path: Path) -> None:
    from docx import Document

    rows = [[f"relay{i:02d}", "1"] for i in range(11)] + [["Not reported", "77"]]
    document = Document()
    loop.add_capped_rollup_table(
        document,
        "Relays",
        "Top 10 Relays",
        ["Relay", "Devices"],
        rows,
        "relays",
    )

    text = "\n".join(p.text for p in document.paragraphs)
    assert (
        "Showing 10 of 11 relays. 77 devices did not report a relay."
    ) in text
