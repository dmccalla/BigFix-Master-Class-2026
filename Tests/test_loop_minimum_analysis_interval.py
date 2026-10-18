# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""_BESClient_Report_MinimumAnalysisInterval from Client_Settings_Info.

A device lists the setting only when it is set. A setting with no value uses
the 60 second default. The report names it only when a device sets it.
"""

from __future__ import annotations

from pathlib import Path

import get_loop_health as loop
import openpyxl
from test_loop_reporting_coverage import PNG_1X1, _metrics

SETTING = "_BESClient_Report_MinimumAnalysisInterval"


def _settings_lines(rows: list[tuple[str, str, str]]) -> list[str]:
    lines = [
        "CLIENT SETTINGS INFORMATION",
        " ",
        "//===== ROOT SERVER AND RELAY SETTINGS ==================================",
        "Column Headers: Setting Name, Setting Value, Number of Machines",
        " ",
        # A relay row with the same name must not count.
        f"{SETTING} ",
        " 900",
        "1",
        "//===== CLIENT SETTINGS ==================================",
        "Description: Unique client settings specified on clients",
        "Column Headers: Setting Name, Setting Value, Number of Machines",
        " ",
        "_BESClient_A ",
        " 1",
        "3",
    ]
    for name, value, count in rows:
        lines += [f"{name} ", f" {value}" if value else " ", count]
    lines += [
        "_BESClient_Z ",
        " 5",
        "2",
        "//===== ICMP CONTROLS ==================================",
    ]
    return lines


LAB_ROWS = [(SETTING, "600", "2"), (SETTING, "1800", "1")]


def _write_log(folder: Path, rows: list[tuple[str, str, str]]) -> None:
    path = folder / "Client_Settings_Info_BFROOT1_10-18-2026_1011.log"
    path.write_text("\n".join(_settings_lines(rows)) + "\n", encoding="utf-8")


def _write_xml(folder: Path, rows: list[tuple[str, str, str]]) -> None:
    path = folder / "Client_Settings_Info_BFROOT1_10-18-2026_1011.xml"
    answers = [f'<Answer type="string">{line}</Answer>' for line in _settings_lines(rows)]
    path.write_text(
        '<?xml version="1.0" encoding="UTF-8"?><BESAPI><Query><Result>'
        + "".join(answers)
        + "</Result></Query></BESAPI>",
        encoding="utf-8",
    )


def test_reads_the_client_rows_from_log_and_xml_alike(tmp_path: Path) -> None:
    log_dir = tmp_path / "log"
    xml_dir = tmp_path / "xml"
    log_dir.mkdir()
    xml_dir.mkdir()
    _write_log(log_dir, LAB_ROWS)
    _write_xml(xml_dir, LAB_ROWS)

    from_log = loop.parse_minimum_analysis_interval(log_dir)
    from_xml = loop.parse_minimum_analysis_interval(xml_dir)

    assert from_log == from_xml
    assert [(r.raw_value, r.seconds, r.devices) for r in from_log] == [
        ("1800", 1800, 1),
        ("600", 600, 2),
    ]


def test_a_setting_with_no_value_uses_sixty_seconds(tmp_path: Path) -> None:
    _write_log(tmp_path, [(SETTING, "", "4")])

    [row] = loop.parse_minimum_analysis_interval(tmp_path)

    assert (row.raw_value, row.seconds, row.devices) == ("", 60, 4)


def test_no_device_setting_it_gives_an_empty_list(tmp_path: Path) -> None:
    _write_log(tmp_path, [])

    assert loop.parse_minimum_analysis_interval(tmp_path) == []
    assert loop.parse_minimum_analysis_interval(tmp_path / "missing") is None


def test_paragraph_names_the_value_above_the_loop() -> None:
    rows = [
        loop.MinimumAnalysisInterval("1800", 1800, 1),
        loop.MinimumAnalysisInterval("600", 600, 2),
    ]

    assert loop.minimum_analysis_interval_paragraph(rows, 1519) == (
        f"3 devices set {SETTING}. On 1 device it is 30 minutes, longer than the "
        "25m 19s average loop time, so analyses there run at most once every 30 "
        "minutes. The other 2 devices use 10 minutes, which is shorter than the "
        "loop and changes nothing."
    )


def test_paragraph_when_every_value_is_below_the_loop() -> None:
    rows = [loop.MinimumAnalysisInterval("", 60, 4)]

    assert loop.minimum_analysis_interval_paragraph(rows, 1519) == (
        f"4 devices set {SETTING}. 4 devices use 60 seconds (no value set), "
        "which is shorter than the 25m 19s average loop time and changes nothing."
    )


def test_report_shows_the_paragraph_and_the_excel_sheet(tmp_path: Path) -> None:
    rows = [
        loop.MinimumAnalysisInterval("1800", 1800, 1),
        loop.MinimumAnalysisInterval("600", 600, 2),
    ]
    metrics = _metrics(10, 0)
    metrics.avg_mean_seconds = 1519
    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)
    summary = loop.PropertyIntervalSummary(source_file=tmp_path, properties=[])

    loop.generate_word(
        metrics,
        png,
        tmp_path / "loop.docx",
        property_intervals=summary,
        minimum_analysis_interval=rows,
    )
    loop.generate_excel(
        metrics,
        [],
        [],
        tmp_path / "loop.xlsx",
        property_intervals=summary,
        minimum_analysis_interval=rows,
    )

    from docx import Document

    text = "\n".join(p.text for p in Document(str(tmp_path / "loop.docx")).paragraphs)
    assert f"3 devices set {SETTING}." in text
    sheet = openpyxl.load_workbook(tmp_path / "loop.xlsx")["Min Analysis Interval"]
    assert list(sheet.iter_rows(values_only=True)) == [
        (
            "Setting",
            "Value",
            "Seconds Used",
            "Devices",
            "Effect at the Average Loop Time",
        ),
        (SETTING, "1800", 1800, 1, "Analyses run at most once every 30 minutes"),
        (SETTING, "600", 600, 2, "No change at the average loop time"),
    ]


def test_report_says_nothing_when_no_device_sets_it(tmp_path: Path) -> None:
    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)
    summary = loop.PropertyIntervalSummary(source_file=tmp_path, properties=[])

    loop.generate_word(
        _metrics(10, 0),
        png,
        tmp_path / "loop.docx",
        property_intervals=summary,
        minimum_analysis_interval=[],
    )

    from docx import Document

    text = "\n".join(p.text for p in Document(str(tmp_path / "loop.docx")).paragraphs)
    assert SETTING not in text
