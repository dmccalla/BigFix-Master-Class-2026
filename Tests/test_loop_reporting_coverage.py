# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""CLT-01 names missing loop data when more than 5% of scored devices have none.

The Devices Reporting tile is red only then.
"""

from __future__ import annotations

from pathlib import Path

import get_loop_health as loop
from docx import Document

PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR"
    b"\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde"
    b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfeA\xe2%\xb5"
    b"\x00\x00\x00\x00IEND\xaeB`\x82"
)
TEMPLATE = Path(__file__).resolve().parents[1] / "Assets" / "chart_template.html"


def _metrics(reporting: int, no_data: int) -> loop.ReportMetrics:
    total = reporting + no_data
    return loop.ReportMetrics(
        total_devices=reporting,
        score_total_devices=total,
        overall_status="Healthy",
        environment_score=95,
        score_driver="test driver",
        score_breakdown="test breakdown",
        report_start_text="7/1/2026",
        report_end_text="7/2/2026",
        report_span_text="2 days",
        generated_at_text="7/2/2026 17:32",
        avg_mean_seconds=720,
        avg_median_seconds=240,
        avg_max_seconds=600,
        slowest_avg_seconds=900,
        peak_max_seconds=1200,
        avg_healthy=reporting,
        avg_warning=0,
        avg_critical=no_data,
        avg_healthy_pct=reporting / total,
        avg_warning_pct=0.0,
        avg_critical_pct=no_data / total,
        avg_over_120=0,
        avg_over_120_pct=0.0,
        not_reported_count=no_data,
        error_count=0,
        not_reported_pct=no_data / total,
        error_pct=0.0,
        total_issue_count=no_data,
        total_issue_pct=no_data / total,
        healthy_threshold_seconds=1800,
        critical_threshold_seconds=3600,
        fix_first_count=0,
    )


def test_missing_data_is_a_reason_only_above_five_percent() -> None:
    assert not loop.no_usable_data_raised(_metrics(100, 0))
    assert not loop.no_usable_data_raised(_metrics(95, 5))
    assert loop.no_usable_data_raised(_metrics(94, 6))
    # The 10-06 sample: 60 of 2,537 is 2.4%.
    assert not loop.no_usable_data_raised(_metrics(2477, 60))


def test_missing_data_reason_names_the_devices_and_the_share() -> None:
    [reason] = loop.loop_finding_reasons(_metrics(94, 6))
    assert reason.impact == "Medium"
    assert reason.text == "6 of 100 devices (6.0%) have no usable loop data"
    assert reason.section == "Loop Health"
    [single] = loop.loop_finding_reasons(_metrics(10, 1))
    assert single.text == "1 of 11 devices (9.1%) has no usable loop data"


def test_one_finding_takes_the_highest_impact_of_its_reasons() -> None:
    metrics = _metrics(94, 6)
    metrics.overall_status = "CRITICAL"
    reasons = loop.loop_finding_reasons(metrics)
    assert [r.impact for r in reasons] == ["Critical", "Medium"]
    assert loop.loop_finding_impact(reasons) == "Critical"
    assert loop.loop_findings_line(reasons) == "Raises 1 Critical finding (2 reasons)"


def test_reporting_tile_is_red_only_when_missing_data_is_a_reason() -> None:
    below = loop.context_strip_tiles(_metrics(2477, 60))[0]
    assert below["label"] == "Devices reporting"
    assert "color" not in below

    above = loop.context_strip_tiles(_metrics(94, 6))[0]
    assert above["color"] == "#C2410C"


def test_template_does_not_turn_partial_reporting_red_by_itself() -> None:
    text = TEMPLATE.read_text(encoding="utf-8")
    start = text.index("if(/devices reporting/i.test(label)){")
    branch = text[start : text.index("return '#187A43';", start)]
    assert "if(c.color) return c.color;" in branch
    assert "#C2410C" not in branch


def test_infographic_lists_missing_data_under_clt01(monkeypatch, tmp_path: Path) -> None:
    captured: dict = {}

    def capture_chart(kind, data, output):
        captured.update({"kind": kind, "data": data})

    monkeypatch.setattr("Lib.Design.html_charts.render_chart", capture_chart)
    loop.generate_png(_metrics(94, 6), tmp_path / "loop.png")

    data = captured["data"]
    assert data["findings"] == [
        {
            "id": "CLT-01",
            "text": "Medium. 6 of 100 devices (6.0%) have no usable loop data.",
        }
    ]
    assert data["findingsLine"] == "Raises 1 Medium finding"
    assert data["context"][0]["color"] == "#C2410C"


def test_no_finding_below_the_limit(monkeypatch, tmp_path: Path) -> None:
    captured: dict = {}

    def capture_chart(kind, data, output):
        captured.update({"data": data})

    monkeypatch.setattr("Lib.Design.html_charts.render_chart", capture_chart)
    loop.generate_png(_metrics(2477, 60), tmp_path / "loop.png")

    assert captured["data"]["findings"] == []
    assert "color" not in captured["data"]["context"][0]


def test_word_report_lists_every_reason_in_one_row(tmp_path: Path) -> None:
    word_output = tmp_path / "loop.docx"
    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)
    metrics = _metrics(94, 6)
    metrics.overall_status = "WARNING"

    loop.generate_word(metrics, png, word_output, include_automated_findings=True)

    doc = Document(str(word_output))
    findings = next(t for t in doc.tables if t.cell(0, 0).text == "Finding ID")
    assert len(findings.rows) == 2
    assert [c.text for c in findings.rows[1].cells] == [
        "CLT-01",
        "High.\n"
        "The client loop health score is below 90. See Loop Health.\n"
        "6 of 100 devices (6.0%) have no usable loop data. See Loop Health.",
    ]
    assert any(t.cell(0, 0).text == "Infographic tile" for t in doc.tables)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Reporting Issues" not in headings
    assert not any(r.cells[0].text == "Not Reported" for t in doc.tables for r in t.rows)
    # The CP-12 sentence follows the slowest-device sentence under Loop Health.
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    slowest = next(
        i for i, p in enumerate(paragraphs) if p.startswith("The slowest device averages")
    )
    assert paragraphs[slowest + 1].startswith(
        "6 devices (6.0%) reported no usable loop data."
    )
    assert paragraphs.index("How the Score Is Calculated") < paragraphs.index(
        "Loop Health"
    )
