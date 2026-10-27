# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""Retrieved properties in Evaluation Frequency and the Top 10, 10-22-2026.

Collect Deployment Information captures every property that belongs to no
analysis as Reserved, Default or Custom. A reserved property is built into every
Client and its interval cannot be changed, so Evaluation Frequency and CLT-01
leave it out. A default property counts as External and a custom one as Custom.
A capture made before the section keeps the earlier wording.
"""

from __future__ import annotations

from pathlib import Path

import get_loop_health as loop
import openpyxl
import pytest
from docx import Document
from test_loop_evaluation_frequency import (
    HEADERS,
    PNG_1X1,
    SECTION,
    _healthy_metrics,
    _offender,
)

SAMPLE = Path(__file__).resolve().parents[1] / "SampleData" / "BigFix Master Class 2026"
SAMPLE_TOP10 = SAMPLE / "Client_Loop" / "top10-query-results.csv"

RETRIEVED_SECTION = "//===== Retrieved Properties by Evaluation Interval ================"
RETRIEVED_HEADERS = (
    "Column Headers: Property Name, Site ID, Item ID, Property Index, "
    "Evaluation Interval, Property Type"
)
SITE_ID = "2299764711"

# name, analysis, site, interval text, content source
ANALYSIS = [
    ("QA Every Report", "QA Analysis", "BES Support", "00:00:00", "External"),
    ("QA Daily", "QA Analysis", "QA Custom Site", "1 day", "Custom"),
]
# name, item ID, property index, interval text, type
RETRIEVED = [
    ("BES Relay Service Installed", "13", "1", "no-val", "Reserved"),
    ("Computer Name", "4", "1", "no-val", "Reserved"),
    ("RAM", "34", "1", "no-val", "Default"),
    ("Subnet Address", "38", "1", "01:00:00", "Default"),
    ("Patch Window (Local)", "1821", "1", "00:05:00", "Custom"),
    ("Asset Tag", "1820", "1", "1 day", "Custom"),
]

RESERVED_RESULT = (
    "Retrieved Property: BES Relay Service Installed (actionsite.13:Evaluate Property 1)"
)
DEFAULT_RESULT = "Retrieved Property: RAM (actionsite.34:Evaluate Property 1)"
CUSTOM_RESULT = (
    "Retrieved Property: Patch Window (Local) (actionsite.1821:Evaluate Property 1)"
)
MISSING_RESULT = "Retrieved Property: Old Property (actionsite.99:Evaluate Property 1)"

APPROVED_LEAD = (
    "This table covers the 2 active analysis properties and the 4 retrieved "
    "properties. Each row has an evaluation interval that sets how often the "
    "Client re-evaluates it. A property set to every report runs on every loop. "
    "So does a property with no interval, because BigFix evaluates it every "
    "report. Reserved properties are the properties BigFix builds into every "
    "Client, such as Computer Name, OS and IP Address. The Client evaluates them "
    "every report and their interval cannot be changed, so this table and CLT-01 "
    "leave them out. This Mini Report rates 6 hours or longer as Healthy."
)
APPROVED_SOURCE_NOTE = (
    "Evaluation Frequency labels analysis properties from custom, operator and "
    "master sites as Custom, and those from external sites as External. It labels "
    "custom retrieved properties as Custom and BigFix default retrieved properties "
    "as External. External property intervals may be outside customer control. "
    "Review changes with the content owner."
)
EARLIER_SOURCE_NOTE = (
    "Evaluation Frequency labels analysis properties from custom, operator and "
    "master sites as Custom, and those from external sites as External. External "
    "property intervals may be outside customer control. Review changes with the "
    "content owner."
)
APPROVED_FLAG_NOTE = "The property totals and the Evaluation Frequency chart use different definitions of “custom,” so their counts may not match."


def _lines(with_retrieved: bool = True) -> list[str]:
    lines = ["PROPERTY INFORMATION", " ", SECTION, HEADERS, " "]
    for name, analysis, site, interval, source in ANALYSIS:
        lines += [name, analysis, site, "1", "2", "3", "bigfix", interval, source]
    if with_retrieved:
        lines += [" ", RETRIEVED_SECTION, RETRIEVED_HEADERS, " "]
        for name, item, index, interval, kind in RETRIEVED:
            lines += [name, SITE_ID, item, index, interval, kind]
    return lines


def _summary(folder: Path, with_retrieved: bool = True) -> loop.PropertyIntervalSummary:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "Property_Info_BFROOT1_10-22-2026_1702.log"
    path.write_text("\n".join(_lines(with_retrieved)) + "\n", encoding="utf-8")
    summary = loop.parse_property_intervals(folder)
    assert summary is not None
    return summary


def _word(tmp_path: Path, summary: loop.PropertyIntervalSummary) -> Document:
    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)
    loop.generate_word(
        _healthy_metrics(),
        png,
        tmp_path / "loop.docx",
        loop_offenders=[_offender(22.6, RESERVED_RESULT)],
        property_intervals=summary,
    )
    return Document(str(tmp_path / "loop.docx"))


def test_log_and_xml_captures_read_the_same_retrieved_properties(
    tmp_path: Path,
) -> None:
    from_log = _summary(tmp_path / "log")
    xml_dir = tmp_path / "xml"
    xml_dir.mkdir()
    answers = "".join(f'<Answer type="string">{line}</Answer>' for line in _lines())
    (xml_dir / "Property_Info_BFROOT1_10-22-2026_1702.xml").write_text(
        '<?xml version="1.0" encoding="UTF-8"?><BESAPI><Query><Result>'
        + answers
        + "</Result></Query></BESAPI>",
        encoding="utf-8",
    )
    from_xml = loop.parse_property_intervals(xml_dir)

    assert from_xml is not None
    assert from_log.retrieved == from_xml.retrieved
    assert [
        (p.name, p.item_id, p.index, p.property_type) for p in from_log.retrieved
    ] == [(name, item, index, kind) for name, item, index, _, kind in RETRIEVED]
    # The analysis section ends where the retrieved section starts.
    assert [p.name for p in from_log.properties] == [row[0] for row in ANALYSIS]


def test_an_older_capture_has_no_retrieved_properties(tmp_path: Path) -> None:
    summary = _summary(tmp_path, with_retrieved=False)

    assert summary.retrieved is None
    assert summary.counted == summary.properties


def test_reserved_properties_are_listed_but_not_counted(tmp_path: Path) -> None:
    summary = _summary(tmp_path)

    assert [p.content_source for p in summary.retrieved] == [
        "Reserved",
        "Reserved",
        "External",
        "External",
        "Custom",
        "Custom",
    ]
    assert [p.name for p in summary.counted_retrieved] == [
        "RAM",
        "Subnet Address",
        "Patch Window (Local)",
        "Asset Tag",
    ]
    assert len(summary.counted) == 6


def test_evaluation_frequency_counts_default_as_external_and_custom_as_custom(
    tmp_path: Path,
) -> None:
    summary = _summary(tmp_path)
    counts = loop.property_interval_band_counts(summary)

    assert counts["Every report"] == {"total": 2, "custom": 0, "external": 2}
    assert counts["Under 15 minutes"] == {"total": 1, "custom": 1, "external": 0}
    assert counts["15 minutes to 1 hour"] == {"total": 0, "custom": 0, "external": 0}
    assert counts["1 hour to 6 hours"] == {"total": 1, "custom": 0, "external": 1}
    assert counts["6 hours or longer"] == {"total": 2, "custom": 2, "external": 0}
    rows = loop.evaluation_frequency_rows(summary)
    assert [row[1] for row in rows] == [
        loop.count_with_denominator(count, 6) for count in (2, 1, 0, 1, 2)
    ]


def test_clt01_counts_retrieved_properties_but_not_reserved_ones(
    tmp_path: Path,
) -> None:
    summary = _summary(tmp_path)

    assert loop.frequent_property_count(summary) == 3
    reasons = loop.loop_finding_reasons(_healthy_metrics(), summary)
    assert [r.text for r in reasons if r.section == "Evaluation Frequency"] == [
        "3 properties are set to every report or to under 15 minutes"
    ]


def test_top10_names_reserved_and_retrieved_properties(tmp_path: Path) -> None:
    summary = _summary(tmp_path / "new")
    older = _summary(tmp_path / "old", with_retrieved=False)
    label = loop.format_loop_offender_evaluation

    assert label(RESERVED_RESULT, summary) == (
        "BES Relay Service Installed (reserved property)"
    )
    assert label(DEFAULT_RESULT, summary) == "RAM (retrieved property)"
    assert label(CUSTOM_RESULT, summary) == "Patch Window (Local) (retrieved property)"
    assert label(MISSING_RESULT, summary) == "Old Property (retrieved property)"
    # Without the section a reserved property cannot be told apart.
    for capture in (older, None):
        assert label(RESERVED_RESULT, capture) == (
            "BES Relay Service Installed (retrieved property)"
        )


def test_top10_interval_comes_from_the_capture(tmp_path: Path) -> None:
    summary = _summary(tmp_path / "new")
    older = _summary(tmp_path / "old", with_retrieved=False)

    assert loop.offender_interval(RESERVED_RESULT, summary) == "Every report"
    assert loop.offender_interval(DEFAULT_RESULT, summary) == "Every report"
    assert loop.offender_interval(CUSTOM_RESULT, summary) == "5 minutes"
    assert loop.offender_interval(MISSING_RESULT, summary) == "Not found"
    assert loop.offender_interval(RESERVED_RESULT, older) == "Not collected"
    assert loop.offender_interval(RESERVED_RESULT, None) == "Not collected"


def test_word_report_explains_reserved_properties(tmp_path: Path) -> None:
    doc = _word(tmp_path, _summary(tmp_path))

    paragraphs = [p.text for p in doc.paragraphs]
    assert any(p.startswith("This graphic counts ") for p in paragraphs)
    all_text = " ".join(doc.element.body.xpath(".//w:t/text()"))
    assert (
        "Reserved properties are the properties BigFix builds into every Client"
        in all_text
    )
    assert "their interval cannot be changed" in all_text
    lead = next(p for p in paragraphs if p.startswith("1 of these 1 is"))
    assert lead.endswith(
        "Interval is how often a property is set to run. A site evaluation is not "
        "a property and has no interval."
    )
    descriptions = " ".join(s._inline.docPr.get("descr", "") for s in doc.inline_shapes)
    assert "BES Relay Service Installed (reserved property)" in descriptions
    assert "interval Every report" in descriptions
    assert "actionsite.13:Evaluate" not in descriptions


def test_word_report_keeps_the_earlier_wording_for_an_older_capture(
    tmp_path: Path,
) -> None:
    doc = _word(tmp_path, _summary(tmp_path, with_retrieved=False))

    paragraphs = [p.text for p in doc.paragraphs]
    assert any(p.startswith("This graphic counts ") for p in paragraphs)
    assert "Reserved properties" not in " ".join(doc.element.body.xpath(".//w:t/text()"))
    lead = next(p for p in paragraphs if p.startswith("1 of these 1 is"))
    assert lead.endswith("A retrieved property's interval is not in the capture.")
    descriptions = " ".join(s._inline.docPr.get("descr", "") for s in doc.inline_shapes)
    assert "BES Relay Service Installed (retrieved property)" in descriptions
    assert "interval Not collected" in descriptions


@pytest.mark.parametrize("with_retrieved", [True, False])
def test_both_notes_name_retrieved_properties_only_when_captured(
    tmp_path: Path, with_retrieved: bool
) -> None:
    summary = _summary(tmp_path, with_retrieved)
    pressure = loop.ActionSitePressure(
        deployment_info_folder=tmp_path, actionsite_size_kb=1000
    )
    png = tmp_path / "loop.png"
    png.write_bytes(PNG_1X1)
    loop.generate_word(
        _healthy_metrics(),
        png,
        tmp_path / "loop.docx",
        actionsite_pressure=pressure,
        property_intervals=summary,
    )
    loop.generate_excel(
        _healthy_metrics(),
        [],
        [],
        tmp_path / "loop.xlsx",
        actionsite_pressure=pressure,
        property_intervals=summary,
    )

    paragraphs = [p.text for p in Document(str(tmp_path / "loop.docx")).paragraphs]
    sheet = openpyxl.load_workbook(tmp_path / "loop.xlsx")["Data Notes"]
    excel_notes = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    source_note = APPROVED_SOURCE_NOTE if with_retrieved else EARLIER_SOURCE_NOTE
    flag_note = APPROVED_FLAG_NOTE if with_retrieved else loop.PROPERTY_SOURCE_NOTE
    assert source_note in paragraphs
    assert flag_note in " ".join(
        Document(str(tmp_path / "loop.docx")).element.body.xpath(".//w:t/text()")
    )
    assert excel_notes == [flag_note]


@pytest.mark.parametrize(
    ("retrieved", "opening"),
    [
        (
            RETRIEVED[2:3],
            "This table counts 3 properties: the 2 active analysis properties and the 1 "
            "retrieved property. Each row",
        ),
        (
            RETRIEVED[:2],
            "This table counts 2 properties: the 2 active analysis properties and no retrieved "
            "properties. Each row",
        ),
    ],
)
def test_lead_counts_read_naturally(retrieved: list[tuple], opening: str) -> None:
    summary = loop.PropertyIntervalSummary(
        source_file=Path("Property_Info.log"),
        properties=[
            loop.PropertyInterval(name, analysis, site, "bigfix", 0.0, source)
            for name, analysis, site, _, source in ANALYSIS
        ],
        retrieved=[
            loop.RetrievedProperty(name, SITE_ID, item, index, 0.0, kind)
            for name, item, index, _, kind in retrieved
        ],
    )

    assert loop.evaluation_frequency_lead(summary).startswith(opening)


def test_excel_lists_every_retrieved_property(tmp_path: Path) -> None:
    loop.generate_excel(
        _healthy_metrics(),
        [],
        [],
        tmp_path / "loop.xlsx",
        property_intervals=_summary(tmp_path),
    )

    sheet = openpyxl.load_workbook(tmp_path / "loop.xlsx")["Active Properties"]
    rows = [[value or "" for value in row] for row in sheet.iter_rows(values_only=True)]
    assert rows[0] == [
        "Property",
        "Analysis",
        "Site",
        "Issued By",
        "Interval",
        "Interval Band",
        "Content Source",
    ]
    assert [row[0] for row in rows[1:]] == [row[0] for row in ANALYSIS] + [
        row[0] for row in RETRIEVED
    ]
    assert rows[3] == [
        "BES Relay Service Installed",
        "",
        "",
        "",
        "Every report",
        "Every report",
        "Reserved",
    ]
    assert [row[6] for row in rows[3:]] == [
        "Reserved",
        "Reserved",
        "External",
        "External",
        "Custom",
        "Custom",
    ]


def test_sample_capture_carries_every_retrieved_property_type() -> None:
    deployment = SAMPLE / "Deployment_Information"
    if not deployment.is_dir() or not SAMPLE_TOP10.is_file():
        pytest.skip(f"no sample data under {SAMPLE}")
    summary = loop.parse_property_intervals(deployment)

    assert summary is not None and summary.retrieved is not None
    assert {p.property_type for p in summary.retrieved} == {
        "Reserved",
        "Default",
        "Custom",
    }
    results = [
        item.result
        for item in loop.load_loop_offenders(SAMPLE_TOP10)
        if item.result.startswith("Retrieved Property:")
    ]
    assert [loop.format_loop_offender_evaluation(r, summary) for r in results] == [
        "BES Relay Service Installed (reserved property)"
    ]
    assert [loop.offender_interval(r, summary) for r in results] == ["Every report"]
