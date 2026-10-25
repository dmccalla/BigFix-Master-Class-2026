# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""Relay and subnet rollups: every row in the data, top rows only in the Mini Report."""

from pathlib import Path

import get_loop_health as loop
from docx import Document
from docx.text.paragraph import Paragraph

LIMIT = loop.MINI_REPORT_ROLLUP_LIMIT


def _metrics() -> loop.ReportMetrics:
    return loop.ReportMetrics(
        total_devices=10,
        score_total_devices=10,
        overall_status="Healthy",
        environment_score=95,
        score_driver="test driver",
        score_breakdown="test breakdown",
        report_start_text="7/1/2026",
        report_end_text="7/2/2026",
        report_span_text="2 days",
        generated_at_text="7/2/2026 17:32",
        avg_mean_seconds=300,
        avg_median_seconds=240,
        avg_max_seconds=600,
        slowest_avg_seconds=900,
        peak_max_seconds=1200,
        avg_healthy=10,
        avg_warning=0,
        avg_critical=0,
        avg_healthy_pct=1.0,
        avg_warning_pct=0.0,
        avg_critical_pct=0.0,
        avg_over_120=0,
        avg_over_120_pct=0.0,
        not_reported_count=0,
        error_count=0,
        not_reported_pct=0.0,
        error_pct=0.0,
        total_issue_count=0,
        total_issue_pct=0.0,
        healthy_threshold_seconds=1800,
        critical_threshold_seconds=3600,
        fix_first_count=0,
    )


def _udp_summary(tmp_path: Path, groups: int) -> loop.UdpNotificationSummary:
    return loop.UdpNotificationSummary(
        source_file=tmp_path / "udp-notification-results.csv",
        total_devices=groups,
        receiving_count=0,
        not_receiving_count=groups,
        unknown_count=0,
        relay_rows=[[f"relay-{i:02d}", "1 / 1", "100.0%"] for i in range(groups)],
        subnet_rows=[
            [f"10.0.{i}.0/24", "1 / 1", f"relay-{i:02d}"] for i in range(groups)
        ],
    )


def _polling_summary(groups: int) -> loop.CommandPollingSummary:
    def row(name: str) -> list[str]:
        return [name, "1", "1", "0", "0", "0", "0", "0", "0"]

    return loop.CommandPollingSummary(
        total_devices=groups,
        enabled_count=groups,
        disabled_count=0,
        not_set_count=0,
        unknown_count=0,
        enabled_missing_interval_count=0,
        interval_below_loop_count=0,
        no_udp_polling_gap_count=0,
        relay_rows=[row(f"relay-{i:02d}") for i in range(groups)],
        subnet_rows=[row(f"10.0.{i}.0/24") for i in range(groups)],
    )


def _report(tmp_path: Path, groups: int) -> Document:
    output = tmp_path / "loop.docx"
    loop.generate_word(
        _metrics(),
        tmp_path / "missing.png",
        output,
        udp_summary=_udp_summary(tmp_path, groups),
        command_polling_summary=_polling_summary(groups),
    )
    return Document(str(output))


def _table_after(doc: Document, heading: str):
    """Return the first table that follows the named heading."""
    body = doc.element.body
    seen = False
    tables = iter(doc.tables)
    for el in body.iterchildren():
        if el.tag.endswith("}p") and Paragraph(el, doc).text.strip() == heading:
            seen = True
        elif el.tag.endswith("}tbl"):
            table = next(tables)
            if seen:
                return table
    raise AssertionError(f"no table after {heading!r}")


def test_udp_summary_keeps_every_relay_and_subnet(tmp_path: Path) -> None:
    lines = [
        "Computer Name,IPv4 Addresses,Relay,Platform Health Check - Receiving UDP Pings"
    ]
    lines += [f"client-{i},10.0.{i}.5,relay-{i:02d},False" for i in range(LIMIT + 5)]
    udp_file = tmp_path / "udp-notification-results.csv"
    udp_file.write_text("\n".join(lines), encoding="utf-8")

    summary = loop.load_udp_notification_summary(udp_file)

    assert summary is not None
    assert len(summary.relay_rows) == LIMIT + 5
    assert len(summary.subnet_rows) == LIMIT + 5


def test_report_caps_rollups_and_says_so(tmp_path: Path) -> None:
    doc = _report(tmp_path, LIMIT + 5)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    notes = [p.text for p in doc.paragraphs if p.text.startswith("Showing")]

    capped = {
        f"UDP Notification Issues by Relay (Top {LIMIT})": "UDP Issues by Relay",
        f"Top {LIMIT} Subnets Not Receiving UDP Pings": "UDP Issues by Subnet",
        f"Command Polling by Relay (Top {LIMIT})": "Polling by Relay",
        f"Command Polling by Subnet (Top {LIMIT})": "Polling by Subnet",
    }
    for heading in capped:
        assert heading in headings
        assert len(_table_after(doc, heading).rows) == LIMIT + 1  # header + top rows
    for noun in ("relays", "subnets"):
        expected = f"Showing {LIMIT} of {LIMIT + 5:,} {noun}."
        assert sum(n.startswith(expected) for n in notes) == 2
    assert not any("Excel" in n or "workbook" in n.lower() for n in notes)


def test_report_keeps_short_rollups_whole(tmp_path: Path) -> None:
    doc = _report(tmp_path, 3)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]

    for heading in (
        "UDP Notification Issues by Relay",
        "Top Subnets Not Receiving UDP Pings",
        "Command Polling by Relay",
        "Command Polling by Subnet",
    ):
        assert heading in headings
        assert len(_table_after(doc, heading).rows) == 4
    assert not [p.text for p in doc.paragraphs if p.text.startswith("Showing")]
