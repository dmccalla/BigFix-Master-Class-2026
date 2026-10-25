# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
"""Evidence and score boundaries for the reviewed mini-report presentation."""


import get_filldb_performance as filldb
import get_loop_health as loop
import get_perfmon_charts as perfmon
import pytest
from docx import Document
from PIL import Image
from test_loop_evaluation_frequency import _healthy_metrics, _offender
from test_loop_retrieved_properties import _summary

from Lib.Design import html_charts


def _capture_details(monkeypatch):
    jobs = []

    def render(batch):
        for kind, data, path in batch:
            jobs.append((kind, data, path))
            Image.new("RGB", (30, 10), "white").save(path)

    monkeypatch.setattr(html_charts, "render_charts", render)
    return jobs


def test_frequency_uses_counted_population_and_reconciles_clt01(monkeypatch, tmp_path):
    jobs = _capture_details(monkeypatch)
    png = tmp_path / "summary.png"
    Image.new("RGB", (30, 10), "white").save(png)
    summary = _summary(tmp_path / "capture")
    loop.generate_word(
        _healthy_metrics(), png, tmp_path / "report.docx", property_intervals=summary
    )
    data = next(data for _, data, _ in jobs if data["kind"] == "frequency")
    assert data["total"] == 6
    assert sum(r["count"] for r in data["rows"]) == data["total"]
    assert max(r["count"] for r in data["rows"]) < data["total"]
    assert data["frequent"] == 3
    assert [r["status"] for r in data["rows"][:2]] == [
        "Needs Attention",
        "Needs Attention",
    ]
    assert all(r["custom"] + r["external"] == r["count"] for r in data["rows"])


def test_property_name_hyphen_survives_chart_and_alt_text(monkeypatch, tmp_path):
    jobs = _capture_details(monkeypatch)
    png = tmp_path / "summary.png"
    Image.new("RGB", (30, 10), "white").save(png)
    name = "Installed Applications - Windows"
    offender = _offender(
        9.2, "Analysis Name=Application Information (Windows) | Name=" + name
    )
    output = tmp_path / "report.docx"
    loop.generate_word(_healthy_metrics(), png, output, loop_offenders=[offender])
    data = next(data for _, data, _ in jobs if data["kind"] == "costs")
    assert name in data["rows"][0]["name"]
    assert data["rows"][0]["property"]
    assert name in " ".join(
        s._inline.docPr.get("descr", "") for s in Document(output).inline_shapes
    )


def test_udp_denominator_includes_records_without_relay(tmp_path):
    source = tmp_path / "udp.csv"
    source.write_text(
        "Computer Name,IPv4 Address,Relay,Receiving UDP Pings\nclient-a,10.0.0.1,relay-a,No\nclient-b,10.0.0.2,,Yes\n",
        encoding="utf-8",
    )
    summary = loop.load_udp_notification_summary(source)
    assert summary.total_devices == 2
    assert summary.receiving_count == 1 and summary.not_receiving_count == 1
    tiles = loop.context_strip_tiles(_healthy_metrics(), udp_summary=summary)
    tile = next(t for t in tiles if t["label"] == "No UDP pings received")
    assert tile["value"] == "1 / 2"
    assert "UDP capture" in tile["caption"]
    assert "with a relay recorded" not in tile["caption"]


def test_storage_operations_reference_is_not_a_pass_limit():
    threshold = perfmon.get_threshold("PhysicalDisk_Total_Disk_Transfers_sec")
    assert threshold.get("critical") is None
    assert "Context counter" in threshold["label"]
    assert "no pass/fail limit" in threshold["label"]
    assert "It is scored as Pass" in threshold["label"]


def _filldb_rows():
    rows = []
    for name in [
        "Parallel DB Update (Short Batch)",
        "Parallel DB Update",
        "Batch Rate",
        "Parallel Parsing",
        "computer administrators",
    ]:
        rows.append(
            {
                "FillDB Object": name + ":",
                "Count": "10",
                "Time (ms)": "1000",
                "Rate/s": "10",
                "Health Check": "Great",
            }
        )
    rows[1].update(
        Count="0", **{"Time (ms)": "0", "Health Check": "----: Not enough data."}
    )
    rows[-1]["Health Check"] = "004%: Poor"
    return rows


def test_filldb_discloses_poor_context_object_without_changing_score(
    monkeypatch, tmp_path
):
    payload = []
    monkeypatch.setattr(
        html_charts, "render_chart", lambda kind, data, path: payload.append(data)
    )
    filldb.generate_filldb_chart(_filldb_rows(), tmp_path / "filldb.png", {})
    data = payload[0]
    assert data["score"] == 100
    assert "3 were assessed and all met the standard" in data["scopeNote"]
    assert "Parallel DB Update had no data" in data["scopeNote"]
    assert (
        "Computer Administrators is rated Poor but is not a guide check"
        in data["scopeNote"]
    )
    assert len(data["bars"]) == 5
    batch = next(r for r in data["bars"] if r["label"].startswith("Batch Rate"))
    assert "same totals as Parallel DB Update (Short Batch)" in batch["label"]
    assert "Do not add them" in data["workloadNote"]


def test_filldb_does_not_claim_different_batch_totals_match(monkeypatch, tmp_path):
    payload = []
    monkeypatch.setattr(
        html_charts, "render_chart", lambda kind, data, path: payload.append(data)
    )
    rows = _filldb_rows()
    rows[2]["Time (ms)"] = "1100"
    filldb.generate_filldb_chart(rows, tmp_path / "filldb.png", {})
    assert all("same totals" not in r["label"] for r in payload[0]["bars"])
    assert "Batch Rate repeats" not in payload[0]["workloadNote"]


@pytest.mark.parametrize("scope", [None, "PerfMon Capture 2"])
def test_every_disk_plot_is_preserved_in_each_capture(tmp_path, scope):
    import hashlib
    import json

    from Lib.Design.perfmon_report_layout import finalize_perfmon_layout
    document = Document()
    document.add_heading("PerfMon Capture 1", level=2)
    first = tmp_path / "total.png"
    second = tmp_path / "disk.png"
    Image.new("RGB", (30, 10), "red").save(first)
    Image.new("RGB", (30, 10), "blue").save(second)
    document.add_picture(str(first))
    document.add_heading("PerfMon Capture 2", level=2)
    document.add_picture(str(first))
    document.add_picture(str(second))
    output = tmp_path / "two_captures.docx"
    document.save(output)
    rows = []
    for path,counter in [(first,"PhysicalDisk_Total_Current_Disk_Queue_Length"),(second,"PhysicalDisk_0_C_Current_Disk_Queue_Length")]:
        rows.append({"counter":counter,"png_sha256":hashlib.sha256(path.read_bytes()).hexdigest(),"raw_series_digest":"same","plot_series_digest":"same","status":"OK","samples":3,"maximum":1,"section":"Storage Queue Depth"})
    (tmp_path / "perfmon_plot_manifest.json").write_text(json.dumps(rows),encoding="utf-8")
    finalize_perfmon_layout(output,tmp_path,section_title=scope)
    result=Document(output)
    assert len(result.inline_shapes)==3
    validation=json.loads((tmp_path / "perfmon_layout_validation.json").read_text())
    assert validation["plots_omitted"] == []
    assert validation["plots_preserved"] == 2
