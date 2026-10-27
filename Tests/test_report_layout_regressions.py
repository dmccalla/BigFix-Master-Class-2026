# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
"""Report structure regressions from the rendered preview review."""

import pytest
from docx import Document
from docx.oxml.ns import qn

from Lib.Design.report_layout import apply_report_hierarchy, finalize_report_document


def test_heading_sizes_descend_and_sections_use_spacing_without_rules():
    document = apply_report_hierarchy(Document())
    sizes = [document.styles[f"Heading {i}"].font.size.pt for i in range(1, 6)]
    assert sizes == [28, 24, 20, 14, 11]
    for level in (2, 3):
        style = document.styles[f"Heading {level}"]
        assert style.element.find(qn("w:pPr")).find(qn("w:pBdr")) is None
        assert style.paragraph_format.space_before.pt == 20
        assert style.paragraph_format.space_after.pt == 6
    family = document.styles["Heading 2"].font.name
    for level in range(1, 6):
        assert document.styles[f"Heading {level}"].font.name == family
    for level in (4, 5):
        assert (
            document.styles[f"Heading {level}"]
            .element.find(qn("w:pPr"))
            .find(qn("w:pBdr"))
            is None
        )


@pytest.mark.parametrize("level", range(1, 6))
def test_imported_character_style_does_not_toggle_heading_weight(level):
    document = Document()
    heading = document.add_heading("Imported heading", level=level)
    heading.runs[0].style = "Strong"
    body = document.add_paragraph("Body emphasis")
    body.runs[0].style = "Strong"

    apply_report_hierarchy(document)

    assert heading.text == "Imported heading"
    assert document.styles[f"Heading {level}"].font.bold
    assert not heading._p.xpath(".//w:rPr/w:rStyle")
    assert body.runs[0].style.name == "Strong"


def test_live_page_fields_survive_repeated_layout():
    document = apply_report_hierarchy(Document())
    apply_report_hierarchy(document)
    fields = document.sections[0].footer._element.xpath(".//w:fldSimple/@w:instr")
    assert fields == ["PAGE", "NUMPAGES"]
    assert not document.sections[0].different_first_page_header_footer


def test_existing_footer_is_preserved():
    document = Document()
    document.sections[0].footer.paragraphs[0].text = "Existing footer"
    apply_report_hierarchy(document)
    assert document.sections[0].footer.paragraphs[0].text == "Existing footer"


def test_about_section_label_stays_with_its_own_page(tmp_path):
    document = Document()
    document.add_paragraph("Previous section content")
    document.add_page_break()
    document.add_paragraph()
    document.add_paragraph("SECTION 13")
    document.add_heading("About BigFix", level=1)
    path = tmp_path / "report.docx"
    document.save(path)
    finalize_report_document(path)
    checked = Document(path)
    assert checked.paragraphs[1].paragraph_format.page_break_before
    assert checked.paragraphs[2].text == "About BigFix"


def test_module_intro_stays_with_its_first_table(tmp_path):
    document = Document()
    document.add_heading("Client Performance Analysis", level=2)
    document.add_paragraph("Capture introduction")
    document.add_table(rows=1, cols=1)
    document.add_heading("How the Score Is Calculated", level=3)
    path = tmp_path / "report.docx"
    document.save(path)
    finalize_report_document(path)
    checked = Document(path)
    assert checked.paragraphs[1].paragraph_format.keep_with_next
    assert checked.paragraphs[2].text == "How the Score Is Calculated"


def test_section_gaps_preserve_cover_spacing_and_explicit_page_boundaries():
    from docx.shared import Pt

    document = Document()
    cover_spacer = document.add_paragraph()
    cover_spacer.paragraph_format.line_spacing = Pt(10)
    document.add_heading("Introduction", level=1)
    document.add_paragraph()
    section = document.add_heading("Report Scope", level=2)
    section.paragraph_format.space_before = Pt(60)
    document.add_paragraph()
    introduction = document.add_paragraph("Section content")
    document.add_page_break()
    document.add_heading("Next section", level=2)
    chapter_label = document.add_paragraph("SECTION 12")
    document.add_heading("Methodology", level=1)

    apply_report_hierarchy(document)

    assert cover_spacer._p.getparent() is document.element.body
    assert cover_spacer.paragraph_format.line_spacing == Pt(10)
    assert section.paragraph_format.space_before is None
    assert section._p.getprevious() is document.paragraphs[1]._p
    assert "Section content" in document.paragraphs[3].text
    assert document.element.body.xpath('.//w:br[@w:type="page"]')
    assert chapter_label.paragraph_format.page_break_before
    assert introduction.paragraph_format.keep_together


def test_tagged_headings_own_gaps_in_imported_web_text():
    from docx.oxml import OxmlElement

    document = Document()
    document.add_heading("Introduction", level=1)
    previous = document.add_paragraph("Previous section text")
    spacer = document.add_paragraph()
    heading = document.add_heading("Action Latency", level=4)
    explanation = document.add_paragraph("The subsection explanation")
    for paragraph in (previous, explanation):
        spacing = paragraph._p.get_or_add_pPr().get_or_add_spacing()
        for side in ("before", "after"):
            spacing.set(qn("w:" + side + "Autospacing"), "1")
    block = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    block.append(content)
    heading._p.addprevious(block)
    content.append(heading._p)
    content.append(explanation._p)

    apply_report_hierarchy(document)

    assert spacer._p.getparent() is None
    assert previous.paragraph_format.space_after == 0
    assert explanation.paragraph_format.space_before == 0
    assert (
        explanation._p.get_or_add_pPr()
        .get_or_add_spacing()
        .get(qn("w:beforeAutospacing"))
        == "0"
    )
    assert content[0] is heading._p and explanation.text == "The subsection explanation"


def test_chapter_rule_reserves_space_before_an_immediate_note_table():
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("Previous chapter content")
    label = document.add_paragraph("SECTION 11")
    document.add_heading("Reference Data", level=1)
    rule = document.add_paragraph()
    drawing = OxmlElement("w:drawing")
    anchor = OxmlElement("wp:anchor")
    extent = OxmlElement("wp:extent")
    extent.set("cx", "6838950")
    extent.set("cy", "0")
    anchor.append(extent)
    drawing.append(anchor)
    rule.add_run()._r.append(drawing)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Existing chapter note"
    table_before = table._tbl.xml

    apply_report_hierarchy(document)
    apply_report_hierarchy(document)

    assert label.paragraph_format.page_break_before
    assert rule.paragraph_format.line_spacing == Pt(6)
    assert rule.paragraph_format.space_after == Pt(12)
    assert rule.paragraph_format.keep_with_next
    assert table._tbl.xml == table_before
