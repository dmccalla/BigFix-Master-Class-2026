# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
"""Portability, template ownership and approved Client content regressions."""

from pathlib import Path

import get_loop_health as loop
import pytest
from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from test_loop_evaluation_frequency import _healthy_metrics
from test_loop_retrieved_properties import _summary
from test_mini_report_design_fixes import _capture_details, _filldb_rows

from Lib.Design import approved_fonts, html_charts
from Lib.Design.report_layout import apply_report_hierarchy


def _text(document):
    return "\n".join("".join(p.xpath(".//w:t/text()")) for p in document.element.body.xpath(".//w:p"))


def test_optional_fonts_use_the_shipped_regular_and_bold(monkeypatch):
    monkeypatch.setattr(approved_fonts, "_local_body_fonts", lambda: {})
    css = approved_fonts.fontface_css(html_charts._FONT_FILES)
    assert css.count("@font-face") == 2
    assert "font-weight:400" in css and "font-weight:700" in css
    assert "base64," in css
    # Header faces are not required by supporting charts, which have no titles.
    for kind in ("clientdetail", "loop", "filldb"):
        assert css in html_charts._build_html(kind, {})


def test_hierarchy_preserves_body_defaults_and_existing_font():
    document = Document()
    document.styles["Heading 2"].font.name = "Aptos"
    defaults = document.styles.element.find(qn("w:docDefaults"))
    before = etree.tostring(defaults)
    apply_report_hierarchy(document)
    assert etree.tostring(defaults) == before
    assert document.styles["Heading 3"].font.name == "Aptos"
    assert document.styles["Heading 3"].font.size.pt == 20
    assert document.styles["Heading 4"].font.size.pt == 12


def test_bullet_definitions_are_ordered_and_paragraphs_have_list_references():
    from Lib.Design.report_layout import add_bullet

    document = Document()
    paragraph = add_bullet(document, "Bullet content is retained")
    numbering = document.part.numbering_part.element
    children = list(numbering)
    first_num = next(i for i, n in enumerate(children) if n.tag == qn("w:num"))
    assert all(
        i < first_num for i, n in enumerate(children) if n.tag == qn("w:abstractNum")
    )
    numpr = paragraph._p.find(qn("w:pPr")).find(qn("w:numPr"))
    number_id = numpr.find(qn("w:numId")).get(qn("w:val"))
    instance = next(
        n
        for n in children
        if n.tag == qn("w:num") and n.get(qn("w:numId")) == number_id
    )
    abstract_id = instance.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(
        n
        for n in children
        if n.tag == qn("w:abstractNum") and n.get(qn("w:abstractNumId")) == abstract_id
    )
    assert abstract.find(qn("w:multiLevelType")).get(qn("w:val")) == "singleLevel"
    level = abstract.find(qn("w:lvl"))
    assert level.find(qn("w:pStyle")).get(qn("w:val")) == "ReportBullet"
    assert level.find(qn("w:lvlText")).get(qn("w:val")) == "\u2022"


def test_standalone_style_templates_retain_the_body_font():
    templates = list(
        (Path(__file__).resolve().parents[1] / "Templates").glob(
            "*_Style_Template.docx"
        )
    )
    if not templates:
        pytest.skip("Standalone style templates belong to the standalone applications")
    for template in templates:
        document = Document(template)
        family = document.styles["Normal"].font.name
        for level in (1, 2, 3):
            assert document.styles[f"Heading {level}"].font.name == family
        defaults = document.styles.element.find(qn("w:docDefaults"))
        spacing = (
            defaults.find(qn("w:pPrDefault")).find(qn("w:pPr")).find(qn("w:spacing"))
        )
        assert spacing.get(qn("w:line")) == "276"
        table_style = next(s for s in document.styles if s.style_id == "BigFixContent")
        table_spacing = table_style.element.find(qn("w:pPr")).find(qn("w:spacing"))
        assert table_spacing.get(qn("w:line")) == "240"


@pytest.mark.parametrize("with_chart", [False, True])
def test_frequency_explanations_and_guidance_survive_both_paths(
    monkeypatch, tmp_path, with_chart
):
    _capture_details(monkeypatch)
    png = tmp_path / "summary.png"
    if with_chart:
        Image.new("RGB", (30, 10), "white").save(png)
    output = tmp_path / "report.docx"
    loop.generate_word(
        _healthy_metrics(),
        png,
        output,
        property_intervals=_summary(tmp_path / "capture"),
    )
    document = Document(output)
    text = _text(document)
    assert "active analysis properties" in text and "retrieved properties" in text
    assert "A property with no interval counts as Every report" in text
    assert "6 hours or longer is Healthy" in text
    assert "That count also includes external properties" in text
    assert "Reading the Infographic" in text
    assert "Reserved properties" in text
    assert "This Mini Report rates" not in text
    # The first content section follows the title without skipped outline levels.
    headings = [p for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings[0].style.name == "Heading 1"
    assert headings[1].text == "How the Score Is Calculated"
    assert headings[1].style.name == "Heading 2"
    assert document.styles["Heading 2"].font.size.pt == 20
    assert document.styles["Heading 3"].font.size.pt == 12
    # Guidance remains present even though it sits in a tagged Word block.
    guidance = etree.parse(
        str(Path(loop.__file__).parents[1] / "Lib/Design/client_guidance.xml")
    )
    for paragraph in guidance.xpath(
        "//w:p", namespaces={"w": qn("w:p").split("}")[0][1:]}
    ):
        content = "".join(
            paragraph.xpath(
                ".//w:t/text()", namespaces={"w": qn("w:p").split("}")[0][1:]}
            )
        )
        assert not content.strip() or content in text
    bullets = document.element.body.xpath(
        './/w:p[w:pPr/w:pStyle[@w:val="ReportBullet"]]'
    )
    assert len(bullets) == 10
    assert all(p.find(qn("w:pPr")).find(qn("w:numPr")) is not None for p in bullets)
    heading = document.element.body.xpath(
        './/w:sdt[w:sdtPr/w:tag[@w:val="ClientGuidanceStartingPoints"]]/w:sdtContent/w:p'
    )[0]
    assert heading.find(qn("w:pPr")).find(qn("w:pageBreakBefore")) is not None


def test_section_anchor_survives_a_display_heading_edit(monkeypatch, tmp_path):
    jobs = _capture_details(monkeypatch)
    original = DocumentClass.add_heading

    def renamed(self, text="", level=1):
        return original(
            self,
            "Loop health display wording" if text == "Loop Health" else text,
            level,
        )

    monkeypatch.setattr(DocumentClass, "add_heading", renamed)
    png = tmp_path / "summary.png"
    Image.new("RGB", (30, 10), "white").save(png)
    output = tmp_path / "report.docx"
    loop.generate_word(_healthy_metrics(), png, output)
    assert any(data["kind"] == "groups" for _, data, _ in jobs)
    assert "Loop health display wording" in _text(Document(output))


def test_content_chart_names_the_active_measure_and_retains_relevant_counts(
    monkeypatch, tmp_path
):
    jobs = _capture_details(monkeypatch)
    png = tmp_path / "summary.png"
    Image.new("RGB", (30, 10), "white").save(png)
    pressure = loop.ActionSitePressure(
        tmp_path,
        analyses_relevant_count=121,
        analyses_total_count=6702,
        active_analyses_count=140,
        properties_count=10342,
        active_properties_count=877,
        custom_properties_count=34,
    )
    output = tmp_path / "report.docx"
    loop.generate_word(_healthy_metrics(), png, output, actionsite_pressure=pressure)
    counts = next(data for _, data, _ in jobs if data["kind"] == "counts")
    rows = {r["label"]: r for r in counts["rows"]}
    assert rows["Active analyses"]["count"] == 140
    assert rows["Active properties"]["count"] == 877
    text = _text(Document(output))
    assert "121 relevant of 6,702 total. 140 active." in text
    assert "877 active, 34 with the BigFix custom flag" in text



def test_unchecked_vm_is_not_reported_as_a_detection():
    import inject_perfmon_charts as injector

    xml = injector.build_perfmon_hardware_context_xml(
        {
            "environment": {
                "Virtual Machine": False,
                "Virtual Machine Source": "Not supplied",
            }
        }
    )
    assert "Not supplied" in xml
    assert "No (not detected)" not in xml


def test_real_client_and_filldb_render_without_optional_fonts(monkeypatch, tmp_path):
    """Exercise the real browser renderer rather than replacing it with a stub."""
    import sys

    if sys.platform != "win32":
        pytest.skip("The desktop application's supported renderer uses Windows Edge")
    import get_filldb_performance as filldb

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "no-user-fonts"))
    monkeypatch.setenv("WINDIR", str(tmp_path / "no-system-fonts"))
    approved_fonts._local_body_fonts.cache_clear()
    try:
        assert approved_fonts._local_body_fonts() == {}
        png = tmp_path / "client.png"
        loop.generate_png(_healthy_metrics(), png)
        output = tmp_path / "client.docx"
        loop.generate_word(
            _healthy_metrics(),
            png,
            output,
            property_intervals=_summary(tmp_path / "capture"),
        )
        filldb_png = tmp_path / "filldb.png"
        filldb.generate_filldb_chart(_filldb_rows(), filldb_png, {})
        assert png.stat().st_size > 10000
        assert filldb_png.stat().st_size > 10000
        assert len(Document(output).inline_shapes) >= 3
    finally:
        approved_fonts._local_body_fonts.cache_clear()
