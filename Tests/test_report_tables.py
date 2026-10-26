# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""The table layout rule from Design's 10-20 change list, Section 2."""

from __future__ import annotations

import re
import zipfile

import get_loop_health as loop
import pytest
from docx import Document
from lxml import etree

from Lib.Design import report_tables as rt

W = rt.W
NS = f'xmlns:w="{rt.W_NS}"'


def _table(style: str | None, width: str, rows: list[list[str]]) -> etree._Element:
    style_xml = f'<w:tblStyle w:val="{style}"/>' if style else ""
    cols = len(rows[0])
    body = "".join(
        "<w:tr>"
        + "".join(f"<w:tc><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc>" for text in row)
        + "</w:tr>"
        for row in rows
    )
    grid = "".join('<w:gridCol w:w="2000"/>' for _ in range(cols))
    return etree.fromstring(
        f"<w:tbl {NS}><w:tblPr>{style_xml}{width}"
        '<w:jc w:val="center"/></w:tblPr>'
        f"<w:tblGrid>{grid}</w:tblGrid>{body}</w:tbl>"
    )


def _jc(element) -> str | None:
    found = element.find(f".//{W}jc")
    return None if found is None else found.get(W + "val")


@pytest.mark.parametrize(
    "value",
    [
        "14",
        "1,234",
        "-3",
        "84.6%",
        "11 / 13",
        "1,234 / 2,537",
        "12.3 MB",
        "2.05 ms",
        "8 GB",
        "00:12:34",
        "1:05",
        "104.261",
        "2 cores",
        "16 cores",
    ],
)
def test_number_cells(value: str) -> None:
    assert rt.is_number_cell(value)


@pytest.mark.parametrize(
    "value",
    [
        "Not Set",
        "CLT-01",
        "6.7 MB (6,851 KB)",
        "10-20-2026",
        "11.0.6.137",
        "4 readings over",
        "Healthy",
    ],
)
def test_word_cells(value: str) -> None:
    assert not rt.is_number_cell(value)


def test_scope_is_bigfixcontent_with_dxa_or_auto_width() -> None:
    rows = [["A", "B"], ["x", "1"]]
    assert rt.in_scope(_table("BigFixContent", '<w:tblW w:w="0" w:type="auto"/>', rows))
    assert rt.in_scope(
        _table("BigFixContent", '<w:tblW w:w="11500" w:type="dxa"/>', rows)
    )
    assert rt.in_scope(_table("BigFixContent", "", rows))
    assert not rt.in_scope(
        _table("BigFixContent", '<w:tblW w:w="5000" w:type="pct"/>', rows)
    )
    assert not rt.in_scope(_table(None, '<w:tblW w:w="10795" w:type="dxa"/>', rows))
    assert not rt.in_scope(
        _table("TableGridLight", '<w:tblW w:w="0" w:type="auto"/>', rows)
    )


def test_layout_rule() -> None:
    tbl = _table(
        "BigFixContent",
        '<w:tblW w:w="11500" w:type="dxa"/>',
        [
            ["Relay", "Devices / Total", "% of Total", "Work (ms)", "Status"],
            ["relay1", "11 / 13", "84.6%", "Not Set", "Healthy"],
            ["relay2", "2 / 13", "15.4%", "250", "Critical"],
            ["", "", "", "", ""],
        ],
    )
    assert rt.normalize_table(tbl)
    tbl_pr = tbl.find(W + "tblPr")
    assert tbl_pr.find(W + "tblW").get(W + "w") == "10800"
    assert tbl_pr.find(W + "tblW").get(W + "type") == "dxa"
    assert tbl_pr.find(W + "jc").get(W + "val") == "left"
    assert tbl_pr.find(W + "tblInd").get(W + "w") == "0"
    assert tbl_pr.find(W + "tblLayout").get(W + "type") == "fixed"
    order = [etree.QName(child).localname for child in tbl_pr]
    assert order == ["tblStyle", "tblW", "jc", "tblInd", "tblLayout"]

    grid = [int(col.get(W + "w")) for col in tbl.find(W + "tblGrid")]
    assert sum(grid) == 10800
    rows = tbl.findall(W + "tr")
    assert rows[0].find(f"{W}trPr/{W}tblHeader") is not None
    assert rows[0].find(f"{W}trPr/{W}jc").get(W + "val") == "left"
    assert rows[1].find(f"{W}trPr/{W}tblHeader") is None
    for tr in rows:
        widths = [
            int(tc.find(f"{W}tcPr/{W}tcW").get(W + "w")) for tc in tr.findall(W + "tc")
        ]
        assert widths == grid
        aligned = [_jc(tc) for tc in tr.findall(W + "tc")]
        # Number columns right, header included. Text columns left.
        assert aligned == ["left", "right", "right", "left", "left"]


def _center_header(tbl: etree._Element) -> etree._Element:
    for paragraph in tbl.find(W + "tr").iter(W + "p"):
        paragraph.insert(
            0, etree.fromstring(f'<w:pPr {NS}><w:jc w:val="center"/></w:pPr>')
        )
    return tbl


@pytest.mark.parametrize(
    "tbl",
    [
        _table(None, '<w:tblW w:w="10795" w:type="dxa"/>', [["A", "B"], ["x", "1"]]),
        _table("TableGridLight", "", [["A", "B"], ["x", "1"]]),
    ],
)
def test_tables_outside_the_rule_are_untouched(tbl: etree._Element) -> None:
    before = etree.tostring(tbl)
    assert not rt.normalize_table(tbl)
    assert etree.tostring(tbl) == before


def test_percent_width_data_table_gets_column_alignment_only() -> None:
    tbl = _table(
        "BigFixContent",
        '<w:tblW w:w="5000" w:type="pct"/>',
        [["Area", "Readings"], ["Storage", "59"]],
    )
    assert rt.normalize_table(tbl)
    tbl_pr = tbl.find(W + "tblPr")
    assert tbl_pr.find(W + "tblW").get(W + "w") == "5000"
    assert tbl_pr.find(W + "tblW").get(W + "type") == "pct"
    assert tbl_pr.find(W + "tblLayout") is None
    assert [int(col.get(W + "w")) for col in tbl.find(W + "tblGrid")] == [2000, 2000]
    for tr in tbl.findall(W + "tr"):
        assert [_jc(tc) for tc in tr.findall(W + "tc")] == ["left", "right"]


def test_nested_tables_are_left_alone() -> None:
    inner = '<w:tbl><w:tblPr><w:tblStyle w:val="BigFixContent"/></w:tblPr><w:tblGrid><w:gridCol w:w="500"/></w:tblGrid><w:tr><w:tc><w:p/></w:tc></w:tr></w:tbl>'
    body = etree.fromstring(
        f'<w:body {NS}><w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="9000"/></w:tblGrid>'
        f"<w:tr><w:tc>{inner}<w:p/></w:tc></w:tr></w:tbl></w:body>"
    )
    assert rt.normalize_report_tables(body) == 0


def test_note_under_a_table_drops_the_spacer_and_uses_4_and_12_pt() -> None:
    document = Document()
    rows = [[f"relay{i:02d}", "1"] for i in range(11)]
    loop.add_capped_rollup_table(
        document, "Relays", "Top 10 Relays", ["Relay", "Devices"], rows, "relays"
    )
    body = [el for el in document.element.body if el.tag != W + "sectPr"]
    assert body[-2].tag == W + "tbl"
    note = document.paragraphs[-1]
    assert note.text.startswith("Showing 10 of 11 relays.")
    assert note.paragraph_format.space_before.pt == 4
    assert note.paragraph_format.space_after.pt == 12


def test_note_after_a_paragraph_keeps_that_paragraph() -> None:
    document = Document()
    document.add_paragraph("Lead text.")
    loop.add_note_paragraph(document, "A note.")
    assert [p.text for p in document.paragraphs] == ["Lead text.", "A note."]
    assert document.paragraphs[-1].paragraph_format.space_before.pt == 4
    assert document.paragraphs[-1].paragraph_format.space_after.pt == 12


def test_code_built_columns_are_sized_from_their_text() -> None:
    """Design's method (Part 3B): every column gets its least width and the
    rest is shared by how much text each column holds."""
    tbl = _table(
        "BigFixContent",
        "",
        [
            ["Relay", "Total Devices", "Enabled", "Disabled", "Not Set", "Unknown"],
            ["BFRELAY017.dmc.lab", "100", "25", "0", "75", "0"],
        ],
    )
    rt.normalize_table(tbl)
    grid = [int(col.get(W + "w")) for col in tbl.find(W + "tblGrid")]
    assert sum(grid) == 10800
    # The relay name decides the widest column, and no column is under 1000.
    assert grid[0] == max(grid)
    assert grid[0] >= len("BFRELAY017.dmc.lab") * rt._CHAR_DXA + rt._CELL_MARGINS_DXA
    assert min(grid) >= rt._LEAST_COLUMN_DXA


def test_long_text_wraps_at_words_not_mid_word() -> None:
    long_text = "Background Evaluation of External Site: DISA STIG Checklist " * 3
    tbl = _table(
        "BigFixContent",
        "",
        [["Rank", "Occurrences", "Evaluation"], ["1", "1,458", long_text]],
    )
    rt.normalize_table(tbl)
    grid = [int(col.get(W + "w")) for col in tbl.find(W + "tblGrid")]
    assert sum(grid) == 10800
    header_word = len("Occurrences") * rt._CHAR_DXA * rt._HEADER_BOLD
    assert grid[1] >= header_word + rt._CELL_MARGINS_DXA - 10
    assert grid[2] > grid[0] + grid[1]


def test_template_tables_keep_their_proportions() -> None:
    tbl = _table(
        "BigFixContent",
        "",
        [["Resource Name", "Description"], ["BigFix Forum", "Community answers"]],
    )
    cols = tbl.find(W + "tblGrid").findall(W + "gridCol")
    cols[0].set(W + "w", "1000")
    cols[1].set(W + "w", "3000")
    rt.normalize_table(tbl)
    assert [int(col.get(W + "w")) for col in cols] == [2700, 8100]


def test_header_row_stays_with_the_first_body_row() -> None:
    tbl = _table("BigFixContent", "", [["A", "B"], ["x", "1"]])
    rt.normalize_table(tbl)
    header, body = tbl.findall(W + "tr")
    assert header.find(f"{W}trPr/{W}cantSplit").get(W + "val") == "1"
    assert all(p.find(f"{W}pPr/{W}keepNext") is not None for p in header.iter(W + "p"))
    assert all(p.find(f"{W}pPr/{W}keepNext") is None for p in body.iter(W + "p"))


def test_command_polling_tables_take_their_fixed_widths() -> None:
    headers = [
        "Relay",
        "Total Devices",
        "Enabled",
        "Missing Interval",
        "Interval Below Loop",
        "No UDP + Polling Off",
    ]
    name = "bfrelay-east-01.corp.example.com"
    tbl = _table("BigFixContent", "", [headers, [name, "100", "25", "0", "0", "6"]])
    rt.normalize_table(tbl)
    grid = [int(col.get(W + "w")) for col in tbl.find(W + "tblGrid")]
    assert grid == [3000, 1300, 1300, 1700, 1700, 1800]
    for header, width in zip(headers, grid, strict=True):
        longest = max(len(word) for word in header.split())
        assert width >= longest * rt._CHAR_DXA * rt._HEADER_BOLD + 226, header
    # The name is left as it is. Word breaks it where the column ends.
    assert name in "".join(tbl.itertext())


def test_empty_text_keeps_open_and_close_tags() -> None:
    """A self-closing <w:t/> would let a <w:t>...</w:t> text pattern in a later
    report step run into the markup after it."""
    tbl = _table("BigFixContent", "", [["A", "B"], ["x", "1"]])
    body = f'<w:body {NS}><w:p><w:r><w:t xml:space="preserve"></w:t></w:r></w:p>'
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f"<w:document {NS}>{body}{etree.tostring(tbl).decode()}</w:body></w:document>"
    )
    out = rt.normalize_document_xml(xml)
    assert '<w:t xml:space="preserve"></w:t>' in out
    assert "<w:t/>" not in out and 'preserve"/>' not in out


def test_document_control_version_table_takes_its_fixed_widths() -> None:
    tbl = _table(
        "BigFixContent",
        "",
        [["Version", "Date", "Notes"], ["1.0", "10-20-2026", "Initial Draft"]],
    )
    rt.normalize_table(tbl)
    assert [int(col.get(W + "w")) for col in tbl.find(W + "tblGrid")] == [
        1200,
        1800,
        7800,
    ]


def test_percent_width_data_table_keeps_header_with_body() -> None:
    tbl = _table(
        "BigFixContent",
        '<w:tblW w:w="5000" w:type="pct"/>',
        [["Finding ID", "Area"], ["PRF-INF-02", "Storage Latency"]],
    )
    rt.normalize_table(tbl)
    header, body = tbl.findall(W + "tr")
    assert header.find(f"{W}trPr/{W}cantSplit").get(W + "val") == "1"
    assert all(p.find(f"{W}pPr/{W}keepNext") is not None for p in header.iter(W + "p"))
    assert all(p.find(f"{W}pPr/{W}keepNext") is None for p in body.iter(W + "p"))


def test_measured_widths_follow_designs_method() -> None:
    """Part 3B by hand. Floors: ID 1000 (the minimum) and Description 1496.5
    (11 letters at 105 x 1.1 plus 226). Naturals: 1000 and 1591 (the 80th
    percentile body text, "Disk is full!", 13 characters at 105 plus 226).
    The spare 8303.5 is shared 150 : 244.5 and rounded to 10."""
    rows = [["ID", "Description"], ["A1", "Short text"], ["A2", "Disk is full!"]]
    tbl = _table("BigFixContent", "", rows)
    assert rt.measured_widths(tbl.findall(W + "tr"), 2) == [4160, 6640]


def test_check_results_header_takes_its_column_alignment() -> None:
    """Item 2 (Design, 10-20): a centered header never sits over left cells."""
    tbl = _center_header(
        _table(
            "BigFixContent",
            '<w:tblW w:w="5000" w:type="pct"/>',
            [["Check Name", "Results"], ["Relay health", "Passed"]],
        )
    )
    assert rt.normalize_table(tbl)
    header = tbl.find(W + "tr")
    assert [_jc(tc) for tc in header.findall(W + "tc")] == ["left", "left"]


def test_section_health_counts_stay_centered() -> None:
    tbl = _center_header(
        _table(
            "BigFixContent",
            '<w:tblW w:w="5000" w:type="pct"/>',
            [["Section", "Passed"], ["Server", "12"]],
        )
    )
    body = tbl.findall(W + "tr")[1]
    body.findall(W + "tc")[1].find(W + "p").insert(
        0, etree.fromstring(f'<w:pPr {NS}><w:jc w:val="center"/></w:pPr>')
    )
    rt.normalize_table(tbl)
    header = tbl.find(W + "tr")
    assert [_jc(tc) for tc in header.findall(W + "tc")] == ["left", "center"]
    assert _jc(body.findall(W + "tc")[1]) == "center"


def test_pre_aligned_percent_header_still_gets_compact_spacing() -> None:
    tbl = _center_header(
        _table(
            "BigFixContent",
            '<w:tblW w:w="5000" w:type="pct"/>',
            [["ID", "Component", "Status"], ["RSD-38", "Server", "Enabled"]],
        )
    )
    header, body = tbl.findall(W + "tr")
    original_body = etree.tostring(body)
    original_grid = etree.tostring(tbl.find(W + "tblGrid"))
    for paragraph in header.iter(W + "p"):
        rt._set(rt._child(paragraph.find(W + "pPr"), "spacing", rt._PPR_ORDER), after="240")
    rt.normalize_table(tbl)
    assert etree.tostring(body) == original_body
    assert etree.tostring(tbl.find(W + "tblGrid")) == original_grid
    assert header.find(f"{W}trPr/{W}tblHeader") is not None
    assert header.find(f"{W}trPr/{W}jc").get(W + "val") == "left"
    for paragraph in header.iter(W + "p"):
        assert paragraph.find(f"{W}pPr/{W}pStyle").get(W + "val") == rt.HEADER_PARAGRAPH_STYLE
        assert paragraph.find(f"{W}pPr/{W}spacing") is None
    before = etree.tostring(tbl)
    rt.normalize_table(tbl)
    assert etree.tostring(tbl) == before


def test_findings_source_id_cells_do_not_wrap() -> None:
    tbl = _table(
        "BigFixContent",
        "",
        [
            ["Source ID", "Description", "Impact", "Suggested Remediation"],
            ["PERF-13", "Components per baseline", "High", "Split the baseline."],
        ],
    )
    rt.normalize_table(tbl)
    grid = [int(col.get(W + "w")) for col in tbl.find(W + "tblGrid")]
    assert grid == [1350, 3800, 1100, 4550]
    header, body = tbl.findall(W + "tr")
    assert header.find(f"{W}tc/{W}tcPr/{W}noWrap") is None
    first, *rest = body.findall(W + "tc")
    assert first.find(f"{W}tcPr/{W}noWrap") is not None
    assert all(tc.find(f"{W}tcPr/{W}noWrap") is None for tc in rest)


def _out_of_order(parent, order: list[str]) -> bool:
    names = [etree.QName(child).localname for child in parent]
    ranks = [order.index(name) for name in names if name in order]
    return ranks != sorted(ranks)


def test_cell_and_table_properties_are_in_schema_order_in_the_file(tmp_path) -> None:
    """Item 1 (Design, 10-20): vAlign before shd or tcMar, and tblLook before
    tblCellMar, draw white lines in Word. Checked on a saved file."""
    report = tmp_path / "report.docx"
    document = Document()
    document.add_table(rows=2, cols=2)
    document.save(report)
    with zipfile.ZipFile(report) as package:
        members = {name: package.read(name) for name in package.namelist()}
    xml = members["word/document.xml"].decode("utf-8")
    xml = re.sub(
        r"<w:tblPr>.*?</w:tblPr>",
        '<w:tblPr><w:tblStyle w:val="BigFixContent"/><w:tblLook w:val="04A0"/>'
        '<w:tblCellMar><w:left w:w="80" w:type="dxa"/></w:tblCellMar></w:tblPr>',
        xml,
        count=1,
        flags=re.DOTALL,
    )
    xml = re.sub(
        r"<w:tcPr>(.*?)</w:tcPr>",
        r'<w:tcPr>\1<w:vAlign w:val="center"/><w:shd w:val="clear" w:fill="F0F2F5"/>'
        r'<w:tcMar><w:top w:w="40" w:type="dxa"/></w:tcMar></w:tcPr>',
        xml,
    )
    members["word/document.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(report, "w") as package:
        for name, data in members.items():
            package.writestr(name, data)

    rt.normalize_docx_tables(report)

    with zipfile.ZipFile(report) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    cells = list(root.iter(W + "tcPr"))
    assert cells
    assert all(tc_pr.find(W + "vAlign") is not None for tc_pr in cells)
    assert not any(_out_of_order(tc_pr, rt._TCPR_ORDER) for tc_pr in cells)
    assert not any(
        _out_of_order(tbl_pr, rt._TBLPR_ORDER) for tbl_pr in root.iter(W + "tblPr")
    )


def test_label_rows_merged_across_columns_do_not_decide_alignment() -> None:
    """Capacity Comparison: host label rows span every column. "2 cores" and
    "4 GB" make Required a number column, header and cells (Design, 10-21)."""
    tbl = _table(
        "BigFixContent",
        "",
        [["Resource", "Required", "Basis"], ["CPU cores", "2 cores", "Figure 5"]],
    )
    label = etree.fromstring(
        f'<w:tr {NS}><w:tc><w:tcPr><w:gridSpan w:val="3"/></w:tcPr>'
        "<w:p><w:r><w:t>BFROOT1, Root server</w:t></w:r></w:p></w:tc></w:tr>"
    )
    tbl.findall(W + "tr")[0].addnext(label)
    rt.normalize_table(tbl)
    header, _, body = tbl.findall(W + "tr")
    assert [_jc(tc) for tc in header.findall(W + "tc")] == ["left", "right", "left"]
    assert [_jc(tc) for tc in body.findall(W + "tc")] == ["left", "right", "left"]


def test_capacity_configuration_keeps_the_widths_selected_for_host_names():
    table = _table(
        "BigFixContent", '<w:tblW w:w="10800" w:type="dxa"/>',
        [["Host", "Roles", "Cores", "RAM", "Storage", "Platform", "Source"],
         ["BFPLUGINPORTAL", "Plugin Portal, BES Explorer", "4", "8 GB", "100 GB", "Virtual", "Capture"]],
    )
    selected = (2160, 2808, 936, 864, 1296, 1296, 1440)
    for column, width in zip(table.find(W + "tblGrid"), selected):
        column.set(W + "w", str(width))
    assert rt.normalize_table(table)
    assert tuple(int(column.get(W + "w")) for column in table.find(W + "tblGrid")) == selected
    for row in table.findall(W + "tr"):
        assert tuple(int(cell.find(W + "tcPr/" + W + "tcW").get(W + "w")) for cell in row.findall(W + "tc")) == selected
