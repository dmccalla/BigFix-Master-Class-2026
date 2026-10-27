# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

"""Resolved names and timings must survive raw XML copied later in the bundle."""

import os
from pathlib import Path

import get_loop_health as loop
import pytest
import run_diagnostics as builder


@pytest.fixture
def collection(tmp_path):
    anchor = tmp_path / "Client_Loop_10-27-2026_1808.xml"
    anchor.touch()
    raw = tmp_path / "Top10_10-27-2026_1808.xml"
    raw.write_text(
        "<BESAPI><Query><Result><Answer>18.878: actionsite.13:Evaluate Property 1</Answer>"
        "<Answer>11.516: actionsite.13:Evaluate Property 1</Answer>"
        "<Answer>8.000: actionsite.34:Evaluate Property 1</Answer>"
        "</Result></Query></BESAPI>",
        encoding="utf-8",
    )
    resolved = tmp_path / "top10-query-results.csv"
    resolved.write_text(
        "Occurrences,AvgClientProcessingSeconds,TotalClientProcessingSeconds,MaxClientProcessingSeconds,Status,Result\n"
        "2,15.197,30.394,18.878,OK,Retrieved Property: BES Relay Service Installed (actionsite.13:Evaluate Property 1)\n"
        "1,8,8,8,OK,Retrieved Property: RAM (actionsite.34:Evaluate Property 1)\n",
        encoding="utf-8",
    )
    os.utime(resolved, (1_000_000, 1_000_000))
    os.utime(raw, (2_000_000, 2_000_000))
    return anchor, raw, resolved


@pytest.mark.parametrize("route", ["app", "standalone"])
def test_resolved_names_win_over_newer_raw_xml(collection, route):
    anchor, _, resolved = collection
    selected = (
        builder._find_loop_evaluation_file([anchor.parent])
        if route == "app"
        else loop.discover_loop_offenders_file(anchor)
    )
    assert selected == resolved.resolve()
    rows = loop.load_loop_offenders(selected)
    assert (
        loop.format_loop_offender_evaluation(rows[0].result)
        == "BES Relay Service Installed (retrieved property)"
    )
    assert rows[0].occurrences == 2
    assert rows[0].avg_client_processing_seconds == 15.197
    assert rows[0].total_client_processing_seconds == 30.394
    assert loop.offender_kind(rows[0].result) == 2


def test_raw_fallback_does_not_merge_different_property_ids(collection):
    _, raw, resolved = collection
    resolved.unlink()
    rows = loop.load_loop_offenders(raw)
    assert [row.result for row in rows] == [
        "actionsite.13:Evaluate Property 1",
        "actionsite.34:Evaluate Property 1",
    ]
    assert [row.occurrences for row in rows] == [2, 1]
    assert rows[0].avg_client_processing_seconds == pytest.approx(15.197)
    assert rows[1].avg_client_processing_seconds == 8


def test_raw_fallback_uses_captured_property_name_by_id(collection):
    _, raw, _ = collection
    summary = loop.PropertyIntervalSummary(
        source_file=Path("Property_Info.xml"),
        properties=[],
        retrieved=[
            loop.RetrievedProperty(
                "BES Relay Service Installed", "1", "13", "1", 0, "Reserved"
            ),
            loop.RetrievedProperty("RAM", "1", "34", "1", 0, "Default"),
        ],
    )
    rows = loop.load_loop_offenders(raw, property_intervals=summary)
    assert [
        loop.format_loop_offender_evaluation(row.result, summary) for row in rows
    ] == ["BES Relay Service Installed (reserved property)", "RAM (retrieved property)"]
    assert [row.occurrences for row in rows] == [2, 1]
    assert loop.offender_interval(rows[0].result, summary) == "Every report"


@pytest.mark.skipif(
    os.environ.get("WORKBENCH_RUN_DELIVERY") != "1",
    reason="Set WORKBENCH_RUN_DELIVERY=1 for the real name-resolution report build",
)
def test_real_report_uses_resolved_evaluation_names(collection, tmp_path):
    import hashlib
    import json
    import shutil
    import subprocess
    import sys
    from zipfile import ZipFile

    from openpyxl import load_workbook
    from test_client_delivery import write_evidence

    _, raw, resolved = collection
    evidence = write_evidence(tmp_path / "evidence")
    shutil.copy2(resolved, evidence["offenders"])
    competing = evidence["capture"].parent / raw.name
    shutil.copy2(raw, competing)
    newer = evidence["offenders"].stat().st_mtime + 60
    os.utime(competing, (newer, newer))
    output = tmp_path / "report"
    output.mkdir()
    root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [
            sys.executable,
            str(root / "Get/get_loop_health.py"),
            "--input-file",
            str(evidence["capture"]),
            "--output-folder",
            str(output),
            "--deployment-info-folder",
            str(evidence["deployment"]),
            "--udp-notifications-file",
            str(evidence["udp"]),
        ],
        cwd=root,
        capture_output=True,
        text=True,
        stdin=subprocess.DEVNULL,
        timeout=180,
        check=False,
    )
    (output / "build.log").write_text(result.stdout + result.stderr, encoding="utf-8")
    assert result.returncode == 0, result.stdout + result.stderr
    document = next(output.glob("*.docx"))
    sidecars = list(document.with_suffix("").glob("*.chartdata.json"))
    costs = next(
        data
        for path in sidecars
        if (data := json.loads(path.read_text()))["kind"] == "costs"
    )
    expected = [
        "BES Relay Service Installed (retrieved property)",
        "RAM (retrieved property)",
    ]
    assert [row["name"] for row in costs["rows"]] == expected
    assert [row["seconds"] for row in costs["rows"]] == [15.197, 8]
    assert [row["occurrences"] for row in costs["rows"]] == [2, 1]
    workbook = load_workbook(next(output.glob("*.xlsx")), data_only=True)
    try:
        rows = list(workbook["Evaluation Offenders"].iter_rows(values_only=True))
        index = list(rows[0]).index("Result")
        assert [
            loop.format_loop_offender_evaluation(row[index]) for row in rows[1:]
        ] == expected
        assert [float(row[1]) for row in rows[1:]] == [15.197, 8]
        assert [row[2] for row in rows[1:]] == [2, 1]
    finally:
        workbook.close()
    with ZipFile(document) as bundle:
        embedded = {
            hashlib.sha256(bundle.read(name)).hexdigest()
            for name in bundle.namelist()
            if name.startswith("word/media/")
        }
        for path in sidecars:
            png = path.with_name(path.name.removesuffix(".chartdata.json") + ".png")
            assert hashlib.sha256(png.read_bytes()).hexdigest() in embedded
        assert b"Evaluate Property 1 of External Site: actionsite" not in bundle.read(
            "word/document.xml"
        )
