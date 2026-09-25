# Copyright (c) 2026 Duncan McCalla
# SPDX-License-Identifier: MIT
#
# Master Class Diagnostics Workbench
# See the LICENSE and NOTICE files in the project root for license text and attribution.

from docx import Document
from docx.shared import Inches

from Lib.Themes.word_stub_utils import (
    STANDALONE_NARROW_MARGIN_INCHES,
    STUB_NOTICE_TEXT,
    apply_standalone_narrow_margins,
    create_stub_document_from_template,
    set_picture_alt_text,
    strip_stub_front_matter_xml,
)

_TINY_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
    b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00\x00"
    b"\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfeA"
    b"\xe2!\xbc\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _para(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def test_strip_stub_front_matter_removes_title_and_notice() -> None:
    body_xml = "\n".join(
        [
            _para("Client Loop Health"),
            _para(STUB_NOTICE_TEXT),
            _para("Dynamic content remains."),
        ]
    )

    stripped = strip_stub_front_matter_xml(body_xml, titles=("Client Loop Health",))

    assert "Client Loop Health" not in stripped
    assert "Standalone module excerpt" not in stripped
    assert "Dynamic content remains." in stripped


def test_strip_stub_front_matter_keeps_image_first_content() -> None:
    image_para = "<w:p><w:r><w:drawing><wp:inline/></w:drawing></w:r></w:p>"
    body_xml = image_para + _para("Dynamic content remains.")

    stripped = strip_stub_front_matter_xml(body_xml, titles=("Client Loop Health",))

    assert stripped.startswith("<w:p>")
    assert "<w:drawing>" in stripped
    assert "Dynamic content remains." in stripped


def test_set_picture_alt_text_adds_docpr_and_picture_description(tmp_path) -> None:
    png = tmp_path / "chart.png"
    png.write_bytes(_TINY_PNG)
    document = Document()
    picture = document.add_picture(str(png))

    set_picture_alt_text(
        picture,
        " Example chart showing status and counts. ",
        "Example Chart",
    )

    xml = document._element.xml
    assert 'title="Example chart showing status and counts."' in xml
    assert 'descr="Example chart showing status and counts."' in xml
    assert 'name="Example Chart"' in xml


def test_apply_standalone_narrow_margins_sets_all_sections() -> None:
    document = Document()

    apply_standalone_narrow_margins(document)

    margin = Inches(STANDALONE_NARROW_MARGIN_INCHES)
    section = document.sections[0]
    assert section.top_margin == margin
    assert section.bottom_margin == margin
    assert section.left_margin == margin
    assert section.right_margin == margin


def test_template_stub_document_uses_narrow_margins(tmp_path) -> None:
    template = tmp_path / "template.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    document.save(template)

    stub = create_stub_document_from_template(template)

    margin = Inches(STANDALONE_NARROW_MARGIN_INCHES)
    section = stub.sections[0]
    assert section.top_margin == margin
    assert section.bottom_margin == margin
    assert section.left_margin == margin
    assert section.right_margin == margin
